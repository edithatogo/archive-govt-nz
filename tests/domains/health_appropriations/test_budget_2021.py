"""The 2021 source layout retains its own periods and definition coordinates."""

from __future__ import annotations

import hashlib
import re
from datetime import date
from importlib import import_module
from pathlib import Path
from typing import Any, cast

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook
from tests.domains.health_appropriations.test_budget_2022 import row

from archive_govt_nz.domains.health_appropriations import budget_revenue as revenue
from archive_govt_nz.domains.health_appropriations import source_operations
from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    dispatch_bronze,
)
from archive_govt_nz.domains.health_appropriations.adapter_registry import (
    AdapterContext,
    context_adapter_registrations,
)
from archive_govt_nz.domains.health_appropriations.budget import (
    normalize_budget_workbook,
)
from archive_govt_nz.domains.health_appropriations.budget_expenditure_2024 import (
    DEFINITIONS as EXPENSE_DEFINITIONS,
)
from archive_govt_nz.domains.health_appropriations.budget_expenditure_2024 import (
    FIELDS as EXPENSE_FIELDS,
)

OBSERVED = "2026-10-19T17:04:26Z"
OVERVIEWS = {
    "expenditure": (
        "This Expense workbook contains details of actual government expenditure "
        "for the years ended 30 June 2017, 2018, 2019 and 2020; estimated actual "
        "government expenditure for the year ending 30\xa0June 2021 and budgeted "
        "government expenditure for the year ending 30 June 2022 as published in "
        "Budget 2021."
    ),
    "revenue": (
        "The Revenue workbook contains details of  actual government Crown revenue "
        "and capital receipts for the years ended 30 June 2017, 2018, 2019 and 2020; "
        "estimated actual government Crown revenue and capital receipts for the year "
        "ending 30 June 2021 and budgeted government Crown revenue and capital "
        "receipts for the year ending 30 June 2022 as published in Budget 2021."
    ),
}


def fixture_source(
    root: Path, monkeypatch: pytest.MonkeyPatch, kind: str, *, edition: int = 2021
) -> tuple[Path, str]:
    assert edition in (2020, 2021)
    overview = re.sub(
        r"20[0-9]{2}",
        lambda match: str(int(match[0]) + edition - 2021),
        OVERVIEWS[kind],
    )
    book = Workbook()
    raw = book.active
    assert raw is not None
    raw.title = "Raw Data"
    raw.append(list(EXPENSE_FIELDS if kind == "expenditure" else revenue.FIELDS))
    for year, label in (
        (edition - 4, "Actuals"),
        (edition, "Estimated Actual"),
        (edition + 1, "Main Estimates"),
    ):
        raw.append(row(kind, year, label))
    notes = book.create_sheet("Explanation")
    if kind == "expenditure":
        for old, value in EXPENSE_DEFINITIONS.items():
            notes[f"B{int(old[1:]) + 2018 - edition}"] = value
        notes["B3"] = overview
    else:
        for old, value in revenue.DEFINITIONS.items():
            notes["B4" if old == "B3" else f"B{int(old[1:]) + 2019 - edition}"] = (
                overview if old == "B3" else value
            )
    intro = book.create_sheet("Intro")
    for cell in ("A10", "A13", "A14"):
        intro[cell] = "Synthetic notice; no permission inferred."
    path = root / "source.xlsx"
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    module = (
        import_module(
            f"archive_govt_nz.domains.health_appropriations.budget_expenditure_{edition}"
        )
        if kind == "expenditure"
        else revenue
    )
    monkeypatch.setattr(
        module,
        "SOURCE_SHA256" if kind == "expenditure" else f"SOURCE_SHA256_{edition}",
        digest,
    )
    return path, digest


def normalize(
    kind: str, path: Path, output: Path, args: dict[str, str]
) -> dict[str, Any]:
    operation = (
        normalize_budget_workbook
        if kind == "expenditure"
        else revenue.normalize_budget_revenue
    )
    return cast(
        "dict[str, Any]",
        operation(
            path,
            output,
            expected_sha256=args["expected_sha256"],
            source_vintage=args["source_vintage"],
            source_locator=args["source_locator"],
            observed_at=args["observed_at"],
        ),
    )


@pytest.mark.parametrize("kind", ["expenditure", "revenue"])
def test_2021_native_lineage_periods_and_all_routes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str
) -> None:
    path, digest = fixture_source(tmp_path, monkeypatch, kind)
    args = {
        "expected_sha256": digest,
        "source_vintage": "Budget-2021",
        "source_locator": f"https://www.treasury.govt.nz/sites/default/files/2021-12/b21-{kind}-data.xlsx",
        "observed_at": OBSERVED,
    }
    output = tmp_path / "output"
    receipt = normalize(kind, path, output, args)
    assert receipt["counts"] == {
        "input": 3,
        "normalized": 3,
        "rejected": 0,
        "blank": 0,
        "out_of_scope": 0,
    }
    name = "budget_facts.parquet" if kind == "expenditure" else "revenue_facts.parquet"
    facts = pq.read_table(output / name).to_pylist()
    assert [f["valid_time_end"] for f in facts] == [
        date(y, 6, 30) for y in (2017, 2021, 2022)
    ]
    assert all(
        f["unit"] == "$000"
        and f["currency"] is None
        and f["valid_time_start"] is None
        and f["rights_state"] == "not_evaluated"
        for f in facts
    )
    links = pq.read_table(output / "field_lineage.parquet").to_pylist()
    unit_cell = "B37" if kind == "expenditure" else "B35"
    intro_cell = "B3" if kind == "expenditure" else "B4"
    assert any(
        link["field"] == "unit"
        and link["source_coordinate"] == f"'Explanation'!{unit_cell}"
        and link["raw_value"].startswith("Amount $000:")
        for link in links
    )
    assert any(
        link["field"] == "source_vintage"
        and link["source_coordinate"] == f"'Explanation'!{intro_cell}"
        and link["raw_value"] == OVERVIEWS[kind]
        for link in links
    )
    context = AdapterContext(
        **{k: v for k, v in args.items() if k != "expected_sha256"}
    )
    registrations = context_adapter_registrations(
        cpi=context,
        population=context,
        qes=context,
        gdp=context,
        budget_expenditure=context if kind == "expenditure" else None,
        budget_revenue=context if kind == "revenue" else None,
    )
    common = dispatch_bronze(
        path.read_bytes(),
        source_sha256=digest,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        registrations=registrations,
    )
    assert list(common.output.records) == facts
    request = source_operations.SourceRequest(
        path,
        tmp_path / "repeat",
        f"budget-{kind}-2021/v1",
        digest,
        args["source_vintage"],
        args["source_locator"],
        OBSERVED,
    )
    result = source_operations.operate_source(request, dry_run=False)
    assert result["status"] != "failed"
    for filename in (
        name,
        "field_lineage.parquet",
        "row_dispositions.parquet",
        "MANIFEST.json",
    ):
        assert (request.output_dir / filename).read_bytes() == (
            output / filename
        ).read_bytes()
    for key, value in {
        "expected_sha256": "0" * 64,
        "source_locator": "https://www.treasury.govt.nz/unreviewed.xlsx",
        "source_vintage": "Budget-2022",
    }.items():
        with pytest.raises(ValueError, match="contract"):
            normalize(kind, path, tmp_path / key, args | {key: value})
        assert not (tmp_path / key).exists()


@pytest.mark.parametrize("kind", ["expenditure", "revenue"])
@pytest.mark.parametrize("change", ["definition", "row_limit", "period"])
def test_2021_drift_limits_and_unqualified_period(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str, change: str
) -> None:
    from openpyxl import load_workbook  # noqa: PLC0415 - fixture mutation only

    path, _ = fixture_source(tmp_path, monkeypatch, kind)
    book = load_workbook(path)
    if change == "definition":
        book["Explanation"]["B37" if kind == "expenditure" else "B35"] = (
            "Unsupported definition"
        )
    elif change == "row_limit":
        bound = 6425 if kind == "expenditure" else 1014
        book["Raw Data"].cell(bound + 2, 1, "Outside the reviewed source bound")
    else:
        book["Raw Data"].append(row(kind, 2023, "Main Estimates"))
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    module = (
        import_module(
            "archive_govt_nz.domains.health_appropriations.budget_expenditure_2021"
        )
        if kind == "expenditure"
        else revenue
    )
    monkeypatch.setattr(
        module,
        "SOURCE_SHA256" if kind == "expenditure" else "SOURCE_SHA256_2021",
        digest,
    )
    output = tmp_path / "output"
    args = {
        "expected_sha256": digest,
        "source_vintage": "Budget-2021",
        "source_locator": f"https://www.treasury.govt.nz/sites/default/files/2021-12/b21-{kind}-data.xlsx",
        "observed_at": OBSERVED,
    }
    if change == "period":
        receipt = normalize(kind, path, output, args)
        assert receipt["status"] == "partial"
        assert receipt["counts"]["rejected"] == 1
        assert receipt["counts"]["normalized"] == 3
    else:
        with pytest.raises(ValueError, match="contract"):
            normalize(kind, path, output, args)
        assert not output.exists()

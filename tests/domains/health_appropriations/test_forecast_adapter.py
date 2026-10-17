"""Reviewed forecast common adapters preserve native facts and every remainder."""

from __future__ import annotations

import hashlib
from dataclasses import asdict
from pathlib import Path
from typing import cast

import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook
from tests.domains.health_appropriations.test_forecast_successors import (
    PROFILES,
    _source,
)
from tests.domains.health_appropriations.test_reviewed_overview_adapters import _base

from archive_govt_nz.domains.health_appropriations import forecast
from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    dispatch_bronze,
)
from archive_govt_nz.domains.health_appropriations.adapter_registry import (
    AdapterContext,
    context_adapter_registrations,
)
from archive_govt_nz.domains.health_appropriations.forecast_adapter import (
    forecast_health_registration,
)

_OBSERVED = "2026-10-17T00:00:00Z"
_LOCATOR = "https://example.test/forecast.xlsx"


@pytest.mark.parametrize(("profile", "vintage", "health_row"), PROFILES)
def test_common_forecast_matches_native_rows_lineage_and_loss_accounting(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    profile: str,
    vintage: str,
    health_row: int,
) -> None:
    source = tmp_path / "source.xlsx"
    digest = _source(source, health_row)
    payload = source.read_bytes()
    output = tmp_path / "native"
    receipt = forecast.normalize_forecast_workbook(
        source,
        output,
        expected_sha256=digest,
        profile=profile,
        observed_at=_OBSERVED,
        source_vintage=vintage,
        source_locator=_LOCATOR,
    )
    facts = pq.read_table(output / "forecast_facts.parquet").to_pylist()
    lineage = pq.read_table(output / "field_lineage.parquet").to_pylist()
    cells = pq.read_table(output / "cell_dispositions.parquet").to_pylist()

    def forbidden(*_args: object, **_kwargs: object) -> None:
        pytest.fail("common extraction must not call the filesystem normalizer")

    monkeypatch.setattr(forecast, "normalize_forecast_workbook", forbidden)
    registration = forecast_health_registration(
        profile=profile,
        source_locator=_LOCATOR,
        source_vintage=vintage,
        observed_at=_OBSERVED,
    )
    result = dispatch_bronze(
        payload,
        source_sha256=digest,
        media_type=registration.media_type,
        registrations=(registration,),
    )
    assert result.selection.status == "selected"
    assert result.output.layout == profile
    assert list(result.output.records) == facts
    assert len(result.output.records) == 10
    assert [row["amount_type"] for row in result.output.records] == ["Actual"] * 5 + [
        "Forecast"
    ] * 5
    assert [asdict(row) for row in result.output.lineage] == [
        {
            key: row[key]
            for key in (
                "record_id",
                "field",
                "source_coordinate",
                "raw_value",
                "normalized_value",
                "rule",
            )
        }
        for row in lineage
    ]
    expected_losses = [
        {key: row[key] for key in ("source_coordinate", "disposition", "reason")}
        for row in cells
        if row["disposition"] != "normalized"
    ] + [
        {
            "source_coordinate": "worksheet:" + row["sheet"],
            "disposition": "preserved_only",
            "reason": row["reason"],
        }
        for row in cast("list[dict[str, str]]", receipt["excluded_sheets"])
    ]
    assert [asdict(row) for row in result.output.losses] == expected_losses
    assert source.read_bytes() == payload
    assert result == dispatch_bronze(
        payload,
        source_sha256=digest,
        media_type=registration.media_type,
        registrations=(registration,),
    )


@pytest.mark.parametrize(("profile", "vintage", "health_row"), PROFILES)
def test_common_forecast_retains_rejected_formula_without_claiming_operation_pass(
    tmp_path: Path,
    profile: str,
    vintage: str,
    health_row: int,
) -> None:
    source = tmp_path / "source.xlsx"
    _source(source, health_row)
    workbook = load_workbook(source)
    workbook["Core Crown Expense Tables"].cell(health_row, 6, "=1+1")
    workbook.save(source)
    workbook.close()
    payload = source.read_bytes()
    registration = forecast_health_registration(
        profile=profile,
        source_locator=_LOCATOR,
        source_vintage=vintage,
        observed_at=_OBSERVED,
    )
    result = dispatch_bronze(
        payload,
        source_sha256=hashlib.sha256(payload).hexdigest(),
        media_type=registration.media_type,
        registrations=(registration,),
    )
    assert (
        result.selection.status == "selected"
    )  # Selection is not successful normalization.
    assert len(result.output.records) == 9
    assert (
        sum(loss.reason == "formula_not_evaluated" for loss in result.output.losses)
        == 1
    )
    assert any(loss.disposition == "context" for loss in result.output.losses)
    receipt = forecast.normalize_forecast_workbook(
        source,
        tmp_path / "partial",
        expected_sha256=hashlib.sha256(payload).hexdigest(),
        profile=profile,
        observed_at=_OBSERVED,
        source_vintage=vintage,
        source_locator=_LOCATOR,
    )
    assert receipt["status"] == "partial"
    assert cast("dict[str, int]", receipt["counts"])["rejected"] == 1


@pytest.mark.parametrize(("profile", "vintage", "health_row"), PROFILES)
def test_common_forecast_rejects_wrong_vintage_and_hash(
    tmp_path: Path,
    profile: str,
    vintage: str,
    health_row: int,
) -> None:
    source = tmp_path / "source.xlsx"
    digest = _source(source, health_row)
    payload = source.read_bytes()
    registration = forecast_health_registration(
        profile=profile,
        source_locator=_LOCATOR,
        source_vintage=vintage + "-wrong",
        observed_at=_OBSERVED,
    )
    result = dispatch_bronze(
        payload,
        source_sha256=digest,
        media_type=registration.media_type,
        registrations=(registration,),
    )
    assert result.selection.status == "preserved_only"
    assert not result.output.records
    with pytest.raises(ValueError, match="source_hash_mismatch"):
        registration.adapter.extract(payload, source_sha256="0" * 64)


def test_registry_selects_only_the_requested_reviewed_forecast_edition(
    tmp_path: Path,
) -> None:
    registrations = context_adapter_registrations(
        **_base(),
        befu_2026=AdapterContext(_LOCATOR, "BEFU-2026", _OBSERVED),
        hyefu_2025=AdapterContext(_LOCATOR, "HYEFU-2025", _OBSERVED),
    )
    assert len(registrations) == 6
    for profile, vintage, row in PROFILES:
        source = tmp_path / (vintage + ".xlsx")
        digest = _source(source, row)
        result = dispatch_bronze(
            source.read_bytes(),
            source_sha256=digest,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            registrations=registrations,
        )
        assert result.selection.status == "selected"
        assert result.output.layout == profile
        assert len(result.selection.matched_adapter_ids) == 1


@pytest.mark.parametrize("profile", ["befu", "hyefu", "unknown/v1"])
def test_legacy_or_unknown_profile_is_not_registered_as_a_successor(
    profile: str,
) -> None:
    with pytest.raises(ValueError, match="unsupported_common_forecast_profile"):
        forecast_health_registration(
            profile=profile,
            source_locator=_LOCATOR,
            source_vintage="BEFU-2026",
            observed_at=_OBSERVED,
        )


@pytest.mark.parametrize("fault", ["hash", "size", "type", "profile"])
def test_shared_snapshot_boundary_rejects_before_workbook_parsing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fault: str
) -> None:
    source = tmp_path / "source.xlsx"
    digest = _source(source, 9)
    payload = source.read_bytes()

    def forbidden(*_args: object, **_kwargs: object) -> None:
        pytest.fail("invalid snapshot reached workbook parsing")

    monkeypatch.setattr(forecast, "inventory_workbook", forbidden)
    if fault == "size":
        monkeypatch.setattr(forecast, "_MAX_SOURCE_BYTES", len(payload) - 1)
    expected = {
        "hash": "source_hash_mismatch",
        "size": "source_byte_limit",
        "type": "bronze_bytes_required",
        "profile": "unsupported_forecast_profile",
    }
    with pytest.raises((TypeError, ValueError), match=expected[fault]):
        forecast.parse_forecast_snapshot(
            cast("bytes", bytearray(payload)) if fault == "type" else payload,
            expected_sha256="0" * 64 if fault == "hash" else digest,
            profile="unsupported/v1" if fault == "profile" else "befu-2026/v1",
            observed_at=_OBSERVED,
            source_vintage="BEFU-2026",
            source_locator=_LOCATOR,
        )

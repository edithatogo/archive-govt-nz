"""The pinned QES ordinary-hourly series is dispatchable from Bronze."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from tests.domains.health_appropriations.test_qes import fixture

from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    dispatch_bronze,
)
from archive_govt_nz.domains.health_appropriations.qes_adapter import (
    QesAdapter,
    qes_registration,
)


def test_dispatch_emits_nine_quarters_with_full_cell_accounting(
    tmp_path: Path,
) -> None:
    source = tmp_path / "qes.xlsx"
    fixture(source)
    bronze = source.read_bytes()
    result = dispatch_bronze(
        bronze,
        source_sha256=hashlib.sha256(bronze).hexdigest(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        registrations=(
            qes_registration(
                source_locator="https://www.stats.govt.nz/qes.xlsx",
                source_vintage="QES-2026-Q2",
                observed_at="2026-08-31T00:00:00Z",
            ),
        ),
    )
    assert result.selection.adapter_id == "stats-nz-qes-qemq-sasz9a"
    assert result.output.layout.endswith("/v1")
    assert len(result.output.records) == 9
    assert result.output.lineage
    assert result.output.losses
    assert all(row["rights_state"] == "not_evaluated" for row in result.output.records)
    assert result.output.dimension_links
    assert all(
        link.source_coordinate.startswith("'Table 8'!")
        for link in result.output.dimension_links
    )
    assert {link.kind for link in result.output.dimension_links} >= {
        "measure",
        "period",
        "unit",
    }


def test_qes_wrong_vintage_is_preserved_only(tmp_path: Path) -> None:
    source = tmp_path / "qes.xlsx"
    fixture(source)
    bronze = source.read_bytes()
    result = dispatch_bronze(
        bronze,
        source_sha256=hashlib.sha256(bronze).hexdigest(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        registrations=(
            qes_registration(
                source_locator="https://www.stats.govt.nz/qes.xlsx",
                source_vintage="QES-2025-Q2",
                observed_at="2026-08-31T00:00:00Z",
            ),
        ),
    )
    assert result.selection.status == "preserved_only"


def test_qes_adapter_hash_invalid_workbook_and_vintage(tmp_path: Path) -> None:
    source = tmp_path / "qes.xlsx"
    fixture(source)
    bronze = source.read_bytes()
    adapter = QesAdapter("source", "QES-2026-Q2", "now")
    with pytest.raises(ValueError, match="source_hash_mismatch"):
        adapter.extract(bronze, source_sha256="0" * 64)
    assert not adapter.matches_layout(b"not an xlsx")
    assert not QesAdapter("source", "wrong", "now").matches_layout(bronze)
    output = adapter.extract(
        b"invalid", source_sha256=hashlib.sha256(b"invalid").hexdigest()
    )
    assert output.losses[0].reason == "unsupported_qes_workbook_layout"


def test_qes_probe_rejects_empty_payload_and_missing_required_sheets(
    tmp_path: Path,
) -> None:
    adapter = QesAdapter("source", "QES-2026-Q2", "now")
    assert not adapter.matches_layout(b"")

    source = tmp_path / "missing-table.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Contents"
    book["Contents"]["A1"] = "Quarterly Employment Survey: June 2026 quarter"
    book.save(source)
    book.close()
    assert not adapter.matches_layout(source.read_bytes())

    source = tmp_path / "missing-contents.xlsx"
    fixture(source)
    book = load_workbook(source)
    del book["Contents"]
    book.save(source)
    book.close()
    assert not adapter.matches_layout(source.read_bytes())


def test_qes_probe_rejects_oversized_cell_text(tmp_path: Path) -> None:
    source = tmp_path / "oversized-cell.xlsx"
    fixture(source)
    book = load_workbook(source)
    book["Table 8"]["Z1"] = "x" * 4097
    book.save(source)
    book.close()
    assert not QesAdapter("source", "QES-2026-Q2", "now").matches_layout(
        source.read_bytes()
    )

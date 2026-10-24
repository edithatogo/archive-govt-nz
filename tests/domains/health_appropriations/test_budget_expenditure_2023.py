"""2023 expenditure requires its own source identity, periods and row bound."""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from archive_govt_nz.domains.health_appropriations import budget_expenditure_2023 as exp
from archive_govt_nz.domains.health_appropriations import (
    budget_expenditure_2024 as prior,
)
from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    dispatch_bronze,
)
from archive_govt_nz.domains.health_appropriations.adapter_registry import (
    AdapterContext,
    context_adapter_registrations,
)
from archive_govt_nz.domains.health_appropriations.budget import (
    normalize_budget_workbook,
)
from archive_govt_nz.domains.health_appropriations.budget_adapter import (
    budget_expenditure_registration,
)

ROW = [
    "Ministry of Health",
    "Health",
    101,
    100,
    "Health services",
    "Hospital care",
    "MCA",
    "Non-Departmental Output Expenses",
    None,
    "Health",
    1234.125,
    2019,
    "Actuals",
    "M",
    "Hospital scope",
    36,
    "Health",
]
OBSERVED = "2026-10-18T07:09:45.484302+00:00"


def fixture(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, rows: list
) -> tuple[Path, str]:
    book = Workbook()
    raw = book.active
    assert raw is not None
    raw.title = "Raw Data"
    raw.append(list(prior.FIELDS))
    for row in rows:
        raw.append(row)
    definitions = book.create_sheet("Explanation")
    for coordinate, value in exp.DEFINITIONS.items():
        definitions[coordinate] = value
    intro = book.create_sheet("Intro")
    for coordinate in ("A10", "A13", "A14"):
        intro[coordinate] = "Synthetic notice; no permission inferred."
    path = tmp_path / "source.xlsx"
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setattr(exp, "SOURCE_SHA256", digest)
    return path, digest


def context(digest: str) -> dict:
    return {
        "expected_sha256": digest,
        "source_locator": exp.SOURCE_LOCATOR,
        "source_vintage": "Budget-2023",
        "observed_at": OBSERVED,
    }


def test_2023_source_periods_and_competing_dispatch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rows = [
        ROW,
        [*ROW[:11], 2023, "Estimated Actual", *ROW[13:]],
        [*ROW[:11], 2024, "Main Estimates", "A", *ROW[14:]],
        [*ROW[:11], 2025, "Main Estimates", *ROW[13:]],
    ]
    path, digest = fixture(tmp_path, monkeypatch, rows)
    output = tmp_path / "output"
    receipt = exp.normalize(path, output, **context(digest))
    assert receipt["counts"]["normalized"] == 3
    assert receipt["counts"]["rejected"] == 1
    facts = pq.read_table(output / "budget_facts.parquet").to_pylist()
    assert [f["valid_time_end"] for f in facts] == [
        date(2019, 6, 30),
        date(2023, 6, 30),
        date(2024, 6, 30),
    ]
    assert all(
        f["transformation_id"] == "budget-2023-health-expenditure/v1" for f in facts
    )
    assert all(
        f["unit"] == "$000"
        and f["currency"] is None
        and f["valid_time_start"] is None
        and f["rights_state"] == "not_evaluated"
        for f in facts
    )
    args = {
        "source_locator": exp.SOURCE_LOCATOR,
        "source_vintage": "Budget-2023",
        "observed_at": OBSERVED,
    }
    common = dispatch_bronze(
        path.read_bytes(),
        source_sha256=digest,
        media_type=prior.MEDIA_TYPE,
        registrations=(
            budget_expenditure_registration(**args),
            prior.registration(**args),
            exp.registration(**args),
        ),
    )
    assert common.selection.adapter_id == "nz-budget-2023-health-expenditure"
    assert list(common.output.records) == facts
    overview_links = [
        r for r in common.output.lineage if r.source_coordinate == "'Explanation'!B3"
    ]
    assert len(overview_links) == 6
    assert {r.field for r in overview_links} == {"source_vintage", "amount_type"}
    assert not prior.source_binding("Budget-2023", exp.SOURCE_LOCATOR, digest)
    adapter_context = AdapterContext(**args)
    routed = dispatch_bronze(
        path.read_bytes(),
        source_sha256=digest,
        media_type=prior.MEDIA_TYPE,
        registrations=context_adapter_registrations(
            cpi=adapter_context,
            population=adapter_context,
            qes=adapter_context,
            gdp=adapter_context,
            budget_expenditure=adapter_context,
        ),
    )
    assert list(routed.output.records) == facts
    legacy = tmp_path / "legacy"
    normalize_budget_workbook(path, legacy, **context(digest))
    assert {p.name: p.read_bytes() for p in legacy.iterdir()} == {
        p.name: p.read_bytes() for p in output.iterdir()
    }


@pytest.mark.parametrize("extra", [0, 1])
def test_2023_exact_source_row_bound(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, extra: int
) -> None:
    other = [*ROW[:1], "Education", *ROW[2:]]
    path, digest = fixture(tmp_path, monkeypatch, [other] * (6695 + extra) + [ROW])
    output = tmp_path / "output"
    if extra:
        with pytest.raises(ValueError, match="contract"):
            exp.normalize(path, output, **context(digest))
        assert not output.exists()
    else:
        receipt = exp.normalize(path, output, **context(digest))
        assert receipt["counts"] == {
            "input": 6696,
            "normalized": 1,
            "out_of_scope": 6695,
            "blank": 0,
            "rejected": 0,
        }


@pytest.mark.parametrize("change", ["hash", "locator", "vintage"])
def test_2023_source_binding_is_checked_before_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, change: str
) -> None:
    path, digest = fixture(tmp_path, monkeypatch, [ROW])
    args = context(digest)
    args[
        {
            "hash": "expected_sha256",
            "locator": "source_locator",
            "vintage": "source_vintage",
        }[change]
    ] = {
        "hash": "0" * 64,
        "locator": prior.SOURCE_LOCATOR,
        "vintage": "Budget-2024",
    }[change]
    output = tmp_path / "output"
    with pytest.raises(ValueError, match="contract"):
        exp.normalize(path, output, **args)
    assert not output.exists()

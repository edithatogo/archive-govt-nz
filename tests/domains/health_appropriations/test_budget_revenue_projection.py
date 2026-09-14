"""Revenue canonical facts remain source-faithful and separate from spending."""

from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from dataclasses import replace
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from archive_govt_nz.domains.health_appropriations import (
    budget_canonical_export as budget_export,
)
from archive_govt_nz.domains.health_appropriations import budget_revenue as extraction
from archive_govt_nz.domains.health_appropriations import (
    budget_revenue_canonical_export as revenue_export,
)
from archive_govt_nz.domains.health_appropriations.budget_revenue_canonical_export import (
    export_budget_revenue,
)
from archive_govt_nz.domains.health_appropriations.budget_revenue_projection import (
    RULE,
    project_budget_revenue,
)
from archive_govt_nz.domains.health_appropriations.budget_revenue_reader import (
    _object,
    read_verified_budget_revenue,
)
from archive_govt_nz.domains.health_appropriations.canonical_consumer import (
    NOMINAL_REVENUE_SCHEMA,
    query_nominal_revenue,
)
from archive_govt_nz.domains.health_appropriations.local_provenance_reader import (
    CanonicalPackageInput,
    read_local_provenance,
    read_verified_canonical_tables,
)
from archive_govt_nz.schemas.health_recordsets import recordset_schema

_HEADERS = [
    "Department",
    "Vote",
    "App ID",
    "Description",
    "Revenue Type",
    "Amount $000",
    "Year",
    "Amount Type",
]
_ROW = [
    "Ministry of Health",
    "Health",
    397,
    "Hospital reimbursement",
    "Non-Tax Revenue",
    58746,
    2021,
    "Actuals",
]
_DEFINITIONS = {
    "B3": (
        "The Revenue workbook contains details of  actual government Crown revenue "
        "and capital receipts for the years ended 30 June 2021, 2022, 2023 and 2024; "
        "estimated actual government Crown revenue and capital receipts for the year "
        "ending 30 June 2025 and budgeted government Crown revenue and capital "
        "receipts for the year ending 30 June 2026 as published in Budget 2025."
    ),
    "B37": "Amount $000: Amount (in thousands) for the Year as reported in the Main Estimates.",
    "B38": "Year: Year ending that the Amount relates to (at 30 June) as reported in the Main Estimates.",
    "B39": 'Amount Type: "Actuals" - as audited for prior Years, "Estimated Actual" - for the Year immediately prior to the current Main Estimates, "Main Estimates" - for the Main Estimates Year.',
    "B43": "App ID: Number used to uniquely identify each Crown revenue or capital receipt line.",
}


def _source(tmp_path: Path) -> Path:
    book = Workbook()
    raw = book.active
    assert raw is not None
    raw.title = "Raw Data"
    raw.append(_HEADERS)
    raw.append(_ROW)
    raw.append([*_ROW[:4], "Capital Receipts", 0, 2026, "Main Estimates"])
    explanation = book.create_sheet("Explanation")
    for cell, text in _DEFINITIONS.items():
        explanation[cell] = text
    intro = book.create_sheet("Intro")
    for cell in ("A10", "A13", "A14"):
        intro[cell] = "Synthetic notice observation; no permission asserted."
    book.create_sheet("Pivot Trend by Vote")["A1"] = "retained"
    path = tmp_path / "source.xlsx"
    book.save(path)
    book.close()
    return path


def inputs(tmp_path: Path) -> dict[str, Any]:
    original = _source(tmp_path)
    root = tmp_path / "raw"
    receipt = extraction.normalize_budget_revenue(
        original,
        root,
        expected_sha256=hashlib.sha256(original.read_bytes()).hexdigest(),
        source_locator="data/raw/b25-revenue-data.xlsx",
        observed_at="2026-08-30T00:00:00Z",
    )
    return {
        "manifest": json.loads((root / "MANIFEST.json").read_text()),
        "manifest_sha256": hashlib.sha256(
            (root / "MANIFEST.json").read_bytes()
        ).hexdigest(),
        "facts": pq.read_table(root / "revenue_facts.parquet"),
        "lineage": pq.read_table(
            root / "field_lineage.parquet", schema=extraction.LINEAGE_SCHEMA
        ),
        "dispositions": pq.read_table(root / "row_dispositions.parquet"),
        "receipt": receipt,
        "root": root,
        "original": original,
    }


def test_reader_requires_exact_manifest_and_payload_pins(tmp_path: Path) -> None:
    subject = inputs(tmp_path)
    facts, lineage, dispositions, manifest = read_verified_budget_revenue(
        subject["root"], subject["manifest_sha256"]
    )
    assert facts.equals(subject["facts"])
    assert lineage.equals(subject["lineage"])
    assert dispositions.equals(subject["dispositions"])
    assert manifest == subject["manifest"]
    with pytest.raises(ValueError, match=r"^budget_revenue_package_contract$"):
        read_verified_budget_revenue(subject["root"], "a" * 64)


def test_reader_rejects_duplicate_manifest_keys() -> None:
    with pytest.raises(ValueError, match=r"^budget_revenue_package_contract$"):
        _object([("same", 1), ("same", 2)])


def test_local_revenue_export_is_deterministic_and_local_only(tmp_path: Path) -> None:
    subject = inputs(tmp_path)
    first = tmp_path / "first"
    plan = export_budget_revenue(
        subject["root"], subject["manifest_sha256"], subject["original"], first
    )
    assert plan["status"] == "planned"
    assert not first.exists()
    written = export_budget_revenue(
        subject["root"],
        subject["manifest_sha256"],
        subject["original"],
        first,
        dry_run=False,
    )
    second = tmp_path / "second"
    export_budget_revenue(
        subject["root"],
        subject["manifest_sha256"],
        subject["original"],
        second,
        dry_run=False,
    )
    assert written["status"] == "passed"
    assert {path.name for path in first.iterdir()} == {
        "revenue_fact.parquet",
        "field_lineage.parquet",
        "projection_receipt.json",
        "LOCAL_REVENUE.json",
    }
    assert {path.name: path.read_bytes() for path in first.iterdir()} == {
        path.name: path.read_bytes() for path in second.iterdir()
    }


def test_local_revenue_provenance_recomputes_verified_projection(
    tmp_path: Path,
) -> None:
    subject = inputs(tmp_path)
    output = tmp_path / "canonical"
    export_budget_revenue(
        subject["root"],
        subject["manifest_sha256"],
        subject["original"],
        output,
        dry_run=False,
    )
    package = CanonicalPackageInput(
        kind="revenue",
        root=output,
        marker_sha256=hashlib.sha256(
            (output / revenue_export.MARKER).read_bytes()
        ).hexdigest(),
        original=subject["original"],
        raw_root=subject["root"],
        raw_manifest_sha256=subject["manifest_sha256"],
    )
    tables, receipt = read_verified_canonical_tables(package)
    assert set(tables) == {"revenue_fact", "field_lineage"}
    assert receipt["kind"] == "revenue"
    inventory = read_local_provenance((package,))
    assert inventory["packages"][0]["kind"] == "revenue"
    assert len(inventory["inventory"]["products"]) == 2


def test_nominal_revenue_query_preserves_source_observations(tmp_path: Path) -> None:
    subject = inputs(tmp_path)
    output = tmp_path / "canonical"
    export_budget_revenue(
        subject["root"],
        subject["manifest_sha256"],
        subject["original"],
        output,
        dry_run=False,
    )
    package = CanonicalPackageInput(
        kind="revenue",
        root=output,
        marker_sha256=hashlib.sha256(
            (output / revenue_export.MARKER).read_bytes()
        ).hexdigest(),
        original=subject["original"],
        raw_root=subject["root"],
        raw_manifest_sha256=subject["manifest_sha256"],
    )
    table, receipt = query_nominal_revenue((package,))
    assert table.schema.equals(NOMINAL_REVENUE_SCHEMA, check_metadata=True)
    assert table.num_rows == 2
    assert {row["revenue_type"] for row in table.to_pylist()} == {
        "Capital Receipts",
        "Non-Tax Revenue",
    }
    assert receipt["aggregation"] == "none"
    assert receipt["netting"] == "prohibited"


def test_nominal_revenue_query_rejects_invalid_or_ambiguous_packages(
    tmp_path: Path,
) -> None:
    subject = inputs(tmp_path)
    output = tmp_path / "canonical"
    export_budget_revenue(
        subject["root"],
        subject["manifest_sha256"],
        subject["original"],
        output,
        dry_run=False,
    )
    package = CanonicalPackageInput(
        kind="revenue",
        root=output,
        marker_sha256=hashlib.sha256(
            (output / revenue_export.MARKER).read_bytes()
        ).hexdigest(),
        original=subject["original"],
        raw_root=subject["root"],
        raw_manifest_sha256=subject["manifest_sha256"],
    )
    for value in ((), (package, package), (replace(package, kind="budget"),)):
        with pytest.raises(ValueError, match=r"^canonical_consumer_invalid$"):
            query_nominal_revenue(value)


def test_local_revenue_export_rejects_an_invalid_input_before_writing(
    tmp_path: Path,
) -> None:
    subject = inputs(tmp_path)
    output = tmp_path / "output"
    with pytest.raises(ValueError, match=r"^budget_revenue_canonical_export_input$"):
        export_budget_revenue(subject["root"], "a" * 64, subject["original"], output)
    assert not output.exists()


def test_local_revenue_export_internal_guard_is_bounded() -> None:
    invalid = 0
    with pytest.raises(ValueError, match=r"^budget_revenue_canonical_export_contract$"):
        revenue_export._require(invalid)  # noqa: SLF001 - exercises local contract guard.


def test_local_revenue_export_records_a_bounded_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    subject = inputs(tmp_path)
    output = tmp_path / "output"

    def fail_readback(*_args: object) -> None:
        raise ValueError

    io = revenue_export._io  # noqa: SLF001 - exercises the exporter's write verification.
    monkeypatch.setattr(io, "_readback", fail_readback)
    with pytest.raises(ValueError, match=r"^budget_revenue_canonical_export_write$"):
        export_budget_revenue(
            subject["root"],
            subject["manifest_sha256"],
            subject["original"],
            output,
            dry_run=False,
        )
    assert json.loads((output / "FAILURE.json").read_text()) == {
        "schema_version": revenue_export.SCHEMA,
        "status": "failed",
    }


def test_local_revenue_export_supports_an_unpinned_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    subject = inputs(tmp_path)
    output = tmp_path / "output"
    io = revenue_export._io  # noqa: SLF001 - exercises the cross-platform writer path.

    def unpinned(path: Path) -> budget_export._PinnedDirectory:
        status = path.stat()
        return budget_export._PinnedDirectory(  # noqa: SLF001
            path, (status.st_dev, status.st_ino), None
        )

    monkeypatch.setattr(io, "_pin", unpinned)
    export_budget_revenue(
        subject["root"],
        subject["manifest_sha256"],
        subject["original"],
        output,
        dry_run=False,
    )
    assert (output / revenue_export.MARKER).is_file()


def test_local_revenue_export_bounds_a_failure_marker_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    subject = inputs(tmp_path)
    output = tmp_path / "output"
    io = revenue_export._io  # noqa: SLF001 - exercises the exporter's write boundary.
    write = io._write  # noqa: SLF001 - retains the hardened writer for regular files.

    def fail_readback(*_args: object) -> None:
        raise ValueError

    def reject_failure_marker(
        root: budget_export._PinnedDirectory,
        name: str,
        payload: bytes,
    ) -> None:
        if name == "FAILURE.json":
            raise OSError
        write(root, name, payload)

    monkeypatch.setattr(io, "_readback", fail_readback)
    monkeypatch.setattr(io, "_write", reject_failure_marker)
    with pytest.raises(ValueError, match=r"^budget_revenue_canonical_export_write$"):
        export_budget_revenue(
            subject["root"],
            subject["manifest_sha256"],
            subject["original"],
            output,
            dry_run=False,
        )
    assert not (output / "FAILURE.json").exists()


def test_projects_source_labels_without_netting(tmp_path: Path) -> None:
    subject = inputs(tmp_path)
    result = project_budget_revenue(
        **{
            key: subject[key]
            for key in (
                "manifest",
                "manifest_sha256",
                "facts",
                "lineage",
                "dispositions",
            )
        }
    )
    facts = result.tables["revenue_fact"].to_pylist()
    assert {(row["revenue_type"], row["amount"]) for row in facts} == {
        ("Non-Tax Revenue", Decimal("58746.000000000000000000")),
        ("Capital Receipts", Decimal(0)),
    }
    assert all(
        row["recordset"] == "revenue_fact"
        and row["measure"] == "crown_revenue_or_capital_receipt"
        for row in facts
    )
    assert all(
        "revenue_not_netted_with_expenditure" in row["quality_flags"] for row in facts
    )
    assert result.receipt["netting"] == "prohibited"
    for name, table in result.tables.items():
        assert table.schema.equals(recordset_schema(name), check_metadata=True)
    links = result.tables["field_lineage"].to_pylist()
    assert {link["field"] for link in links} >= {
        "amount",
        "value_token",
        "revenue_type",
        "source_application_id",
    }
    assert all(link["rule"] == RULE for link in links)


def test_input_order_and_decimal_context_do_not_change_projection(
    tmp_path: Path,
) -> None:
    subject = inputs(tmp_path)
    kwargs = {
        key: subject[key]
        for key in ("manifest", "manifest_sha256", "facts", "lineage", "dispositions")
    }
    expected = project_budget_revenue(**kwargs)
    copied = deepcopy(kwargs)
    for name in ("facts", "lineage", "dispositions"):
        table = copied[name]
        copied[name] = table.take(list(reversed(range(table.num_rows))))
    with localcontext() as context:
        context.prec = 2
        actual = project_budget_revenue(**copied)
        assert context.prec == 2
    assert actual == expected


@pytest.mark.parametrize(
    "change", ["wrong_type", "wrong_amount", "wrong_manifest", "bad_schema"]
)
def test_rejects_inputs_before_creating_a_canonical_table(
    tmp_path: Path, change: str
) -> None:
    subject = inputs(tmp_path)
    kwargs = {
        key: subject[key]
        for key in ("manifest", "manifest_sha256", "facts", "lineage", "dispositions")
    }
    if change == "wrong_manifest":
        kwargs["manifest"] = {**kwargs["manifest"], "rights_state": "eligible"}  # type: ignore[arg-type]
    elif change == "bad_schema":
        kwargs["facts"] = pa.table({"wrong": ["row"]})
    else:
        rows = kwargs["facts"].to_pylist()  # type: ignore[union-attr]
        rows[0]["revenue_type" if change == "wrong_type" else "amount"] = (
            "Tax Revenue" if change == "wrong_type" else Decimal("9.000")
        )
        kwargs["facts"] = pa.Table.from_pylist(rows, schema=extraction.FACT_SCHEMA)
    with pytest.raises(ValueError, match=r"^budget_revenue_projection_contract$"):
        project_budget_revenue(**kwargs)

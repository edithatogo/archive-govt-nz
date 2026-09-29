"""QES canonical projection must retain wage semantics and unknowns."""

from __future__ import annotations

import hashlib
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from archive_govt_nz.domains.health_appropriations import qes, qes_canonical_projection
from archive_govt_nz.domains.health_appropriations.qes_canonical_projection import (
    OBSERVED_AT,
    SOURCE_LOCATOR,
    SOURCE_VINTAGE,
    project_qes_earnings,
)
from archive_govt_nz.schemas.health_recordset_normalization import validate_table
from archive_govt_nz.schemas.health_recordsets import recordset_schema


def _source(path: Path) -> str:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Table 8"
    for address, value in {
        "A1": "Table 8",
        "A3": "Average hourly earnings(1)",
        "A4": "By sector",
        "P6": "Total",
        "P7": "Ordinary time",
        "A8": "Series ref: QEMQ",
        "P8": "SASZ9A",
        "A10": "($)",
        "A12": "Quarter",
        "A23": "Percentage change from the same quarter of previous year",
        "A36": "Percentage change from previous quarter",
        "B49": "Average hourly earnings are calculated by dividing earnings by paid hours.",
        "A51": "Source: Stats NZ",
        "A13": "2024",
        "A16": "2025",
        "A20": "2026",
    }.items():
        sheet[address] = value
    for row, month in enumerate(
        ("Jun", "Sep", "Dec", "Mar", "Jun", "Sep", "Dec", "Mar", "Jun"), 13
    ):
        sheet[f"C{row}"] = month
        sheet[f"P{row}"] = row + 0.25
    workbook.create_sheet("Table 9")
    workbook.create_sheet("Contents")["A1"] = qes.RELEASE_TITLE
    workbook.save(path)
    workbook.close()
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _build(tmp_path: Path) -> tuple[Path, Path, str, str]:
    source = tmp_path / "source.xlsx"
    pin = _source(source)
    cas = tmp_path / "cas" / "sha256"
    original = cas / pin[:2] / pin
    original.parent.mkdir(parents=True)
    original.write_bytes(source.read_bytes())
    silver = tmp_path / "silver"
    qes.normalize_qes(
        source,
        silver,
        expected_sha256=pin,
        source_locator=SOURCE_LOCATOR,
        source_vintage=SOURCE_VINTAGE,
        observed_at=OBSERVED_AT,
        dry_run=False,
    )
    manifest_pin = hashlib.sha256((silver / "MANIFEST.json").read_bytes()).hexdigest()
    return silver, cas, manifest_pin, pin


def test_projection_preserves_quarterly_wage_semantics(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    silver, cas, manifest_pin, source_pin = _build(tmp_path)
    monkeypatch.setattr(qes_canonical_projection, "SOURCE_SHA256", source_pin)
    monkeypatch.setattr(
        qes_canonical_projection, "SOURCE_MANIFEST_SHA256", manifest_pin
    )
    before = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in silver.iterdir()
    }
    facts, lineage, receipt = project_qes_earnings(silver, manifest_pin, cas)

    assert facts.schema.equals(
        recordset_schema("earnings_fact", version="v2"), check_metadata=True
    )
    assert lineage.schema.equals(
        recordset_schema("field_lineage", version="v2"), check_metadata=True
    )
    assert facts.num_rows == 9
    assert lineage.num_rows == 90
    validate_table("earnings_fact", facts, version="v2")
    rows = facts.to_pylist()
    assert rows[0]["amount"] == Decimal("13.25")
    assert rows[0]["source_decimal_scale"] == 2
    assert rows[0]["valid_time_start"] == date(2024, 4, 1)
    assert rows[0]["valid_time_end"] == date(2024, 6, 30)
    assert rows[0]["currency"] is None
    assert rows[0]["sex"] is None
    assert rows[0]["adjustment"] is None
    assert rows[0]["rights_state"] == "not_evaluated"
    assert rows[0]["quality_flags"] == list(qes.FLAGS)
    assert receipt["source_object_sha256"] == source_pin
    assert receipt["output_records"] == 9
    assert {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in silver.iterdir()
    } == before


def test_tampered_qes_silver_product_is_rejected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    silver, cas, manifest_pin, source_pin = _build(tmp_path)
    monkeypatch.setattr(qes_canonical_projection, "SOURCE_SHA256", source_pin)
    monkeypatch.setattr(
        qes_canonical_projection, "SOURCE_MANIFEST_SHA256", manifest_pin
    )
    with (silver / "qes_facts.parquet").open("ab") as stream:
        stream.write(b"tamper")
    with pytest.raises(ValueError, match="qes_canonical_projection_invalid"):
        project_qes_earnings(silver, manifest_pin, cas)


def test_v2_is_additive_and_keeps_v1_frozen() -> None:
    v1 = recordset_schema("pharmaceutical_budget_fact")
    v2 = recordset_schema("pharmaceutical_budget_fact", version="v2")
    assert v1.metadata != v2.metadata
    assert v1.names == v2.names
    assert recordset_schema("earnings_fact", version="v2").names[-6:] == [
        "series_id",
        "geography",
        "sector",
        "sex",
        "adjustment",
        "earnings_basis",
    ]

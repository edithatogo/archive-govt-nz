"""Current chart producers must reproduce explicitly pinned historical receipts."""

import hashlib
import runpy
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from archive_govt_nz.domains.health_appropriations import (
    befu_chart_literals as befu,
)
from archive_govt_nz.domains.health_appropriations import (
    health_chart_residual_literals as residual,
)
from archive_govt_nz.domains.health_appropriations import (
    hyefu_allowance_literals as hyefu,
)
from archive_govt_nz.domains.health_appropriations.workbook_common import encode_json

RECIPE = (
    Path(__file__).resolve().parents[3]
    / "evidence/assurance/health-expanded-coverage-20260907/replay.py"
)


def populate_befu(book: Workbook) -> None:
    for title, profile in befu.PROFILES.items():
        sheet = book.create_sheet(title)
        for key, value in profile["anchors"].items():
            sheet[key] = value
        for coordinate in profile["selected"]:
            sheet[coordinate] = -2
            sheet[f"B{sheet[coordinate].row}"] = "Synthetic label"
        for coordinate in profile["excluded"]:
            sheet[coordinate] = "=1+1"


def current_receipt(
    name: str, path: Path, monkeypatch: pytest.MonkeyPatch
) -> dict[str, Any]:
    book = Workbook()
    vintage = "BEFU-2025" if name == "befu-residual" else "HYEFU-2024"
    if name == "befu-chart":
        populate_befu(book)
    elif name == "hyefu-allowance":
        sheet = book.create_sheet(hyefu.SHEET)
        for key, value in hyefu.ANCHORS.items():
            sheet[key] = value
        for coordinate in hyefu.SELECTED:
            sheet[coordinate] = -2
    else:
        profile = residual.PROFILES[vintage]
        sheet = book.create_sheet(profile["sheet"])
        for key, value in profile["anchors"].items():
            sheet[key] = value
        for coordinate in profile["selected"]:
            sheet[coordinate] = -2
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if name == "befu-chart":
        monkeypatch.setattr(befu, "SOURCE_SHA256", digest)
        return befu.admit_befu_chart_literals(path)
    if name == "hyefu-allowance":
        monkeypatch.setattr(hyefu, "SOURCE_SHA256", digest)
        return hyefu.admit_hyefu_allowances(path)
    monkeypatch.setitem(residual.PROFILES[vintage], "sha256", digest)
    return residual.admit_health_chart_residuals(path, vintage)


@pytest.mark.parametrize(
    "name", ["befu-chart", "hyefu-allowance", "befu-residual", "hyefu-residual"]
)
def test_current_producer_historical_projection(
    name: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = current_receipt(name, tmp_path / "synthetic.xlsx", monkeypatch)
    before = encode_json(data)
    assert (
        data["workbook_inventory"]["schema_version"]
        == "archive-govt-nz.workbook-inventory/v1"
    )
    historical = {
        key: value for key, value in data.items() if key != "workbook_inventory"
    }
    expected = encode_json(historical).encode()
    digest = hashlib.sha256(expected).hexdigest()
    project = runpy.run_path(str(RECIPE))["historical_chart_v1"]
    result = project(name, data, digest)
    assert result.payload == expected
    assert result.sha256 == digest
    assert encode_json(data) == before
    with pytest.raises(ValueError, match="historical_chart_v1_contract"):
        project(name, data, "0" * 64)
    data["records"][0]["source_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="historical_chart_v1_contract"):
        project(name, data, digest)


@pytest.mark.parametrize(
    "fault",
    [
        "unknown_field",
        "missing_inventory",
        "inventory_schema",
        "schema",
        "status",
        "rights",
        "profile",
    ],
)
def test_unknown_projection_shapes_fail_closed(fault: str) -> None:
    data = {
        "schema_version": "hyefu-allowance-literal-context/v1",
        "status": "raw_context_only",
        "rights_state": "not_evaluated",
        "records": [],
        "arithmetic_or_cross_vintage_equivalence": "not_performed",
        "workbook_inventory": {
            "schema_version": "archive-govt-nz.workbook-inventory/v1",
            "sheets": [],
        },
    }
    name = "hyefu-allowance"
    if fault == "unknown_field":
        data["future"] = "not silently discarded"
    elif fault == "missing_inventory":
        del data["workbook_inventory"]
    elif fault == "inventory_schema":
        data["workbook_inventory"]["schema_version"] = "future/v2"
    elif fault == "profile":
        name = "future-profile"
    else:
        data[
            {"schema": "schema_version", "rights": "rights_state", "status": "status"}[
                fault
            ]
        ] = "wrong"
    projected = {
        key: value for key, value in data.items() if key != "workbook_inventory"
    }
    digest = hashlib.sha256(encode_json(projected).encode()).hexdigest()
    project = runpy.run_path(str(RECIPE))["historical_chart_v1"]
    with pytest.raises(ValueError, match="historical_chart_v1_contract"):
        project(name, data, digest)

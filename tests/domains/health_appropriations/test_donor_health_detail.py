"""Synthetic contract tests; not retained-source replay."""

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from archive_govt_nz.domains.health_appropriations.donor_health_detail import (
    PROFILES,
    admit_donor_health_detail,
    resolve_header,
)


def test_literal_reference_chain_retains_every_step() -> None:
    sheet = Workbook().active
    assert isinstance(sheet, Worksheet)
    sheet["K100"] = "=K86"
    sheet["K86"] = "=$K$53"
    sheet["K53"] = "=K35"
    sheet["K35"] = "=K5"
    sheet["K5"] = "Forecast"
    value, chain = resolve_header(sheet, "K100")
    assert value == "Forecast"
    assert chain == {
        "K100": "=K86",
        "K86": "=$K$53",
        "K53": "=K35",
        "K35": "=K5",
        "K5": "Forecast",
    }


@pytest.mark.parametrize(
    "formula",
    ["=SUM(K5)", "=K5+0", "='Other'!K5", "=[external.xlsx]Sheet!K5", "=K100", "=XFE1"],
)
def test_unsupported_or_cyclic_header_fails(formula: str) -> None:
    sheet = Workbook().active
    assert isinstance(sheet, Worksheet)
    sheet["K100"] = formula
    with pytest.raises(ValueError, match="donor_health_detail"):
        resolve_header(sheet, "K100")


def test_depth_limit() -> None:
    sheet = Workbook().active
    assert isinstance(sheet, Worksheet)
    for row in range(1, 10):
        sheet[f"A{row}"] = f"=A{row + 1}"
    sheet["A10"] = "Actual"
    with pytest.raises(ValueError, match="header_depth"):
        resolve_header(sheet, "A1")


@pytest.mark.parametrize("vintage", list(PROFILES))
@pytest.mark.parametrize(
    "mutation", ["none", "unit", "year", "type", "formula", "label"]
)
def test_admission(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, vintage: str, mutation: str
) -> None:
    _, title, first = PROFILES[vintage]
    book = Workbook()
    sheet = book.active
    assert isinstance(sheet, Worksheet)
    sheet.title = title
    sheet[f"D{first - 2}"] = "($millions)"
    for column in "FGHIJKLMNO":
        sheet[f"{column}{first - 3}"] = str(2020 + ord(column) - ord("F"))
        sheet[f"{column}{first - 2}"] = "Actual" if column < "K" else "Forecast"
        for row in range(first, first + 8):
            sheet[f"D{row}"] = f"Synthetic {row}"
            sheet[f"{column}{row}"] = 0 if column == "F" else -2
    mutations = {
        "unit": (f"D{first - 2}", "% GDP"),
        "year": (f"F{first - 3}", "1999"),
        "type": (f"F{first - 2}", "Forecast"),
        "formula": (f"F{first}", "=1+1"),
        "label": (f"D{first}", None),
    }
    if mutation in mutations:
        coordinate, value = mutations[mutation]
        sheet[coordinate] = value
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setitem(PROFILES, vintage, (digest, title, first))
    if mutation != "none":
        with pytest.raises(ValueError, match="donor_health_detail"):
            admit_donor_health_detail(path, vintage)
        return
    result = admit_donor_health_detail(path, vintage)
    assert len(result["records"]) == 80
    assert result["records"][0]["source_number_token"] == "0"  # noqa: S105
    assert result["records"][-1]["amount"] == -2
    assert result["records"][-1]["amount_type"] == "Forecast"
    assert result["records"][0]["currency"] is None
    assert result["records"][0]["period_end"] is None
    assert (
        result["formula_totals"]["disposition"] == "excluded_formula_cache_not_admitted"
    )
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest


def test_unknown_profile(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="donor_health_detail"):
        admit_donor_health_detail(tmp_path / "absent", "unapproved")


def test_missing_source(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="donor_health_detail"):
        admit_donor_health_detail(tmp_path / "absent", "BEFU-2025")


@pytest.mark.parametrize("value", [None, 2025, True, "=1", "#REF!"])
def test_only_literal_text_terminal(value: str | int | None) -> None:
    sheet = Workbook().active
    assert isinstance(sheet, Worksheet)
    sheet["K100"] = value
    with pytest.raises(ValueError, match="donor_health_detail"):
        resolve_header(sheet, "K100")

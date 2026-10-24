"""Common forecast operation projection, partial packages and pinned readback."""

from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from pathlib import Path
from typing import TYPE_CHECKING, TypedDict

import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook
from tests.domains.health_appropriations.test_forecast_successors import (
    PROFILES,
    _source,
)

from archive_govt_nz.domains.health_appropriations import (
    forecast,
    source_operations,
)
from archive_govt_nz.domains.health_appropriations import (
    forecast_operations as operations,
)

if TYPE_CHECKING:
    from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
        AdapterRegistration,
        DispatchResult,
    )


def _request(
    tmp_path: Path, profile: str, vintage: str, row: int
) -> source_operations.SourceRequest:
    source = tmp_path / "source.xlsx"
    pin = _source(source, row)
    return source_operations.SourceRequest(
        source,
        tmp_path / "output",
        profile,
        pin,
        vintage,
        "https://example.test/forecast.xlsx",
        "2026-10-17T00:00:00Z",
    )


class _Context(TypedDict):
    profile: str
    expected_sha256: str
    source_vintage: str
    source_locator: str
    observed_at: str


def _context(request: source_operations.SourceRequest) -> _Context:
    return {
        "profile": request.profile,
        "expected_sha256": request.expected_sha256,
        "source_vintage": request.source_vintage,
        "source_locator": request.source_locator,
        "observed_at": request.observed_at,
    }


def _pin(request: source_operations.SourceRequest) -> str:
    return hashlib.sha256(
        (request.output_dir / "MANIFEST.json").read_bytes()
    ).hexdigest()


@pytest.mark.parametrize(("profile", "vintage", "row"), PROFILES)
def test_common_operation_uses_dispatch_and_retains_native_tables(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    profile: str,
    vintage: str,
    row: int,
) -> None:
    request = _request(tmp_path, profile, vintage, row)
    native = tmp_path / "native"
    forecast.normalize_forecast_workbook(request.source, native, **_context(request))
    traced: list[DispatchResult] = []
    dispatch = operations.dispatch_bronze

    def observe(
        bronze: bytes,
        *,
        source_sha256: str,
        media_type: str,
        registrations: tuple[AdapterRegistration, ...],
    ) -> DispatchResult:
        result = dispatch(
            bronze,
            source_sha256=source_sha256,
            media_type=media_type,
            registrations=registrations,
        )
        traced.append(result)
        return result

    def forbidden(*_args: object, **_kwargs: object) -> None:
        pytest.fail("filesystem normalizer must not be the operation route")

    monkeypatch.setattr(operations, "dispatch_bronze", observe)
    monkeypatch.setattr(forecast, "normalize_forecast_workbook", forbidden)
    assert source_operations.operate_source(request)["status"] == "preflight_passed"
    assert not request.output_dir.exists()
    receipt = source_operations.operate_source(request, dry_run=False)
    assert receipt["status"] == "written_local"
    assert "adapter_selection" not in receipt
    assert "manifest_sha256" not in receipt
    for path in native.glob("*.parquet"):
        assert (request.output_dir / path.name).read_bytes() == path.read_bytes()
    assert pq.read_table(
        request.output_dir / "forecast_facts.parquet"
    ).to_pylist() == list(traced[-1].output.records)
    manifest = operations.verify_forecast(
        request.source,
        request.output_dir,
        manifest_sha256=_pin(request),
        **_context(request),
    )
    assert manifest["adapter_selection"] == traced[-1].selection.to_receipt()
    repeat = replace(request, output_dir=tmp_path / "repeat")
    assert source_operations.operate_source(repeat, dry_run=False) == receipt
    assert _pin(repeat) == _pin(request)


@pytest.mark.parametrize(("profile", "vintage", "row"), PROFILES)
def test_partial_selected_forecast_retains_package_but_fails_operation(
    tmp_path: Path,
    profile: str,
    vintage: str,
    row: int,
) -> None:
    request = _request(tmp_path, profile, vintage, row)
    workbook = load_workbook(request.source)
    workbook.worksheets[0].cell(row, 6, "=1+1")
    workbook.save(request.source)
    workbook.close()
    request = replace(
        request, expected_sha256=hashlib.sha256(request.source.read_bytes()).hexdigest()
    )
    assert source_operations.operate_source(request)["status"] == "failed"
    assert not request.output_dir.exists()
    assert (
        source_operations.operate_source(request, dry_run=False)["status"] == "failed"
    )
    manifest = operations.verify_forecast(
        request.source,
        request.output_dir,
        manifest_sha256=_pin(request),
        **_context(request),
    )
    assert manifest["status"] == "partial"
    assert manifest["adapter_selection"]["status"] == "selected"
    assert manifest["counts"]["rejected"] == 1
    assert len(list(request.output_dir.iterdir())) == 4


@pytest.mark.parametrize(
    "mutation",
    [
        "selection",
        "loss",
        "forged_table",
        "missing",
        "extra",
        "pin",
        "context",
        "count_type",
    ],
)
def test_pinned_readback_refuses_tampering(tmp_path: Path, mutation: str) -> None:
    request = _request(tmp_path, *PROFILES[0])
    source_operations.operate_source(request, dry_run=False)
    pin = _pin(request)
    path = request.output_dir / "MANIFEST.json"
    manifest = json.loads(path.read_bytes())
    if mutation == "selection":
        manifest["adapter_selection"]["receipt_sha256"] = "0" * 64
    elif mutation == "loss":
        manifest["common_losses"][0]["reason"] = "invented"
    elif mutation == "forged_table":
        payload = b"wrong data"
        (request.output_dir / "forecast_facts.parquet").write_bytes(payload)
        manifest["output_sha256"]["forecast_facts.parquet"] = hashlib.sha256(
            payload
        ).hexdigest()
    elif mutation == "count_type":
        manifest["counts"]["rejected"] = False
    elif mutation == "missing":
        (request.output_dir / "cell_dispositions.parquet").unlink()
    elif mutation == "extra":
        (request.output_dir / "unbound").write_bytes(b"extra")
    elif mutation == "pin":
        pin = "0" * 64
    else:
        request = replace(request, source_locator="https://example.test/wrong")
    if mutation in {"selection", "loss", "forged_table", "count_type"}:
        # Compact JSON reaches semantic checks even under exact source-derived bounds.
        path.write_text(json.dumps(manifest, separators=(",", ":")))
        pin = _pin(request)
    with pytest.raises((ValueError, OSError)):
        operations.verify_forecast(
            request.source, request.output_dir, manifest_sha256=pin, **_context(request)
        )


def test_symlink_parent_refused_before_any_output(tmp_path: Path) -> None:
    request = _request(tmp_path, *PROFILES[0])
    target = tmp_path / "real"
    target.mkdir()
    link = tmp_path / "link"
    try:
        link.symlink_to(target, target_is_directory=True)
    except OSError as error:
        pytest.skip(f"directory symlinks unavailable: {error}")
    request = replace(request, output_dir=link / "output")
    assert (
        source_operations.operate_source(request, dry_run=False)["status"] == "failed"
    )
    assert list(target.iterdir()) == []


@pytest.mark.parametrize("field", ["records", "lineage", "losses"])
def test_dispatch_projection_must_reconcile_all_common_fields(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    field: str,
) -> None:
    request = _request(tmp_path, *PROFILES[0])
    dispatch = operations.dispatch_bronze

    def altered(
        bronze: bytes,
        *,
        source_sha256: str,
        media_type: str,
        registrations: tuple[AdapterRegistration, ...],
    ) -> DispatchResult:
        result = dispatch(
            bronze,
            source_sha256=source_sha256,
            media_type=media_type,
            registrations=registrations,
        )
        output = result.output
        if field == "records":
            output = replace(
                output,
                records=(
                    {**output.records[0], "unit": "invented"},
                    *output.records[1:],
                ),
            )
        elif field == "lineage":
            output = replace(
                output,
                lineage=(
                    replace(output.lineage[0], raw_value="invented"),
                    *output.lineage[1:],
                ),
            )
        else:
            output = replace(
                output,
                losses=(
                    replace(output.losses[0], reason="invented"),
                    *output.losses[1:],
                ),
            )
        return replace(result, output=output)

    monkeypatch.setattr(operations, "dispatch_bronze", altered)
    assert (
        source_operations.operate_source(request, dry_run=False)["status"] == "failed"
    )
    assert not request.output_dir.exists()


def test_interrupted_operation_is_neither_repaired_nor_overwritten(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    request = _request(tmp_path, *PROFILES[0])
    writer = operations.write_workbook_outputs

    def interrupt(output: Path, _tables: object, _receipt: object) -> None:
        output.mkdir()
        (output / "field_lineage.parquet").write_bytes(b"partial")
        raise KeyboardInterrupt

    monkeypatch.setattr(operations, "write_workbook_outputs", interrupt)
    with pytest.raises(KeyboardInterrupt):
        source_operations.operate_source(request, dry_run=False)
    monkeypatch.setattr(operations, "write_workbook_outputs", writer)
    assert (
        source_operations.operate_source(request, dry_run=False)["status"] == "failed"
    )
    assert (request.output_dir / "field_lineage.parquet").read_bytes() == b"partial"
    assert not (request.output_dir / "MANIFEST.json").exists()
    with pytest.raises((ValueError, OSError)):
        operations.verify_forecast(
            request.source,
            request.output_dir,
            manifest_sha256="0" * 64,
            **_context(request),
        )


def test_readback_limits_are_exact_source_regenerated_sizes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    request = _request(tmp_path, *PROFILES[0])
    assert (
        source_operations.operate_source(request, dry_run=False)["status"]
        == "written_local"
    )
    sizes = {path.name: path.stat().st_size for path in request.output_dir.iterdir()}
    snapshot = operations._snapshot  # noqa: SLF001 - inspect the transport bound
    observed: list[tuple[str, int]] = []

    def bounded(path: Path, pin: str, limit: int) -> bytes:
        observed.append((path.name, limit))
        return snapshot(path, pin, limit)

    monkeypatch.setattr(operations, "_snapshot", bounded)
    operations.verify_forecast(
        request.source,
        request.output_dir,
        manifest_sha256=_pin(request),
        **_context(request),
    )
    assert observed
    assert all(limit == sizes[name] for name, limit in observed)
    # An untrusted file's larger stat size cannot enlarge the regenerated bound.
    table = request.output_dir / "field_lineage.parquet"
    table.write_bytes(table.read_bytes() + b"untrusted extra bytes")
    observed.clear()
    with pytest.raises(ValueError, match="source_byte_limit"):
        operations.verify_forecast(
            request.source,
            request.output_dir,
            manifest_sha256=_pin(request),
            **_context(request),
        )
    assert observed
    assert all(limit == sizes[name] for name, limit in observed)


@pytest.mark.parametrize("target", ["MANIFEST.json", "forecast_facts.parquet"])
@pytest.mark.parametrize("change", ["append", "truncate"])
def test_readback_refuses_appended_and_short_output_bytes(
    tmp_path: Path,
    target: str,
    change: str,
) -> None:
    request = _request(tmp_path, *PROFILES[0])
    source_operations.operate_source(request, dry_run=False)
    pin = _pin(request)
    path = request.output_dir / target
    original = path.read_bytes()
    path.write_bytes(
        original + b"untrusted"
        if change == "append"
        else original[: len(original) // 2]
    )
    if target == "MANIFEST.json":
        pin = _pin(request)
    with pytest.raises((ValueError, OSError)):
        operations.verify_forecast(
            request.source, request.output_dir, manifest_sha256=pin, **_context(request)
        )


def test_malformed_external_pin_is_refused_before_source_extraction(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    request = _request(tmp_path, *PROFILES[0])
    source_operations.operate_source(request, dry_run=False)

    def forbidden(*_args: object, **_kwargs: object) -> None:
        pytest.fail("malformed pin reached source extraction")

    monkeypatch.setattr(operations, "_dispatch", forbidden)
    with pytest.raises(ValueError, match="invalid_forecast_operation"):
        operations.verify_forecast(
            request.source,
            request.output_dir,
            manifest_sha256="not-a-sha256",
            **_context(request),
        )

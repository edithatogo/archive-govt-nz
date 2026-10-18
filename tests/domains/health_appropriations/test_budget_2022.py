"""Reviewed 2022 profiles retain edition bindings, periods and source row bounds."""

from __future__ import annotations

import hashlib
from datetime import date
from importlib import import_module
from pathlib import Path

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from archive_govt_nz.domains.health_appropriations import (
    budget_expenditure_2024 as shared,
)
from archive_govt_nz.domains.health_appropriations import budget_revenue as revenue
from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    dispatch_bronze,
)
from archive_govt_nz.domains.health_appropriations.budget_adapter import (
    budget_expenditure_registration,
)
from archive_govt_nz.domains.health_appropriations.budget_revenue_adapter import (
    budget_revenue_registration,
)

OBSERVED = "2026-10-18T10:47:30+00:00"
OVERVIEWS = {
    "expenditure": (
        "This Expense workbook contains details of actual government expenditure "
        "for the years ended 30 June 2018, 2019, 2020 and 2021; estimated actual "
        "government expenditure for the year ending 30\u00a0June 2022 and budgeted "
        "government expenditure for the year ending 30 June 2023 as published in "
        "Budget 2022."
    ),
    "revenue": (
        "The Revenue workbook contains details of  actual government Crown revenue "
        "and capital receipts for the years ended 30 June 2018, 2019, 2020 and 2021; "
        "estimated actual government Crown revenue and capital receipts for the year "
        "ending 30 June 2022 and budgeted government Crown revenue and capital "
        "receipts for the year ending 30 June 2023 as published in Budget 2022."
    ),
}


def source(
    root: Path, monkeypatch: pytest.MonkeyPatch, kind: str, rows: list[list[object]]
) -> tuple[Path, str]:
    book = Workbook()
    raw = book.active
    assert raw is not None
    raw.title = "Raw Data"
    base = shared if kind == "expenditure" else revenue
    raw.append(list(base.FIELDS))
    for row in rows:
        raw.append(row)
    explanation = book.create_sheet("Explanation")
    for coordinate, value in {**base.DEFINITIONS, "B3": OVERVIEWS[kind]}.items():
        explanation[coordinate] = value
    intro = book.create_sheet("Intro")
    for coordinate in ("A10", "A13", "A14"):
        intro[coordinate] = "Synthetic notice; no permission inferred."
    root.mkdir(exist_ok=True)
    path = root / "source.xlsx"
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if kind == "expenditure":
        module = import_module(
            "archive_govt_nz.domains.health_appropriations.budget_expenditure_2022"
        )
        monkeypatch.setattr(module, "SOURCE_SHA256", digest)
    else:
        monkeypatch.setattr(revenue, "SOURCE_SHA256_2022", digest, raising=False)
    return path, digest


def row(kind: str, year: int = 2018, amount_type: str = "Actuals") -> list[object]:
    if kind == "revenue":
        return [
            "Ministry of Health",
            "Health",
            397,
            "Receipts",
            "Non-Tax Revenue",
            1234.125,
            year,
            amount_type,
        ]
    return [
        "Ministry of Health",
        "Health",
        101,
        100,
        "Services",
        "Hospital",
        "MCA",
        "Non-Departmental Output Expenses",
        None,
        "Health",
        1234.125,
        year,
        amount_type,
        "M",
        "Scope",
        36,
        "Health",
    ]


def normalize(kind: str, path: Path, output: Path, digest: str, **changes: str) -> dict:
    context = {
        "expected_sha256": digest,
        "source_locator": f"https://www.treasury.govt.nz/sites/default/files/2022-05/b22-{kind}-data.xlsx",
        "source_vintage": "Budget-2022",
        "observed_at": OBSERVED,
        **changes,
    }
    if kind == "revenue":
        return revenue.normalize_budget_revenue(path, output, **context)
    module = import_module(
        "archive_govt_nz.domains.health_appropriations.budget_expenditure_2022"
    )
    return module.normalize(path, output, **context)


@pytest.mark.parametrize("kind", ["expenditure", "revenue"])
def test_2022_periods_holds_and_competing_dispatch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str
) -> None:
    path, digest = source(
        tmp_path,
        monkeypatch,
        kind,
        [
            row(kind),
            row(kind, 2022, "Estimated Actual"),
            row(kind, 2023, "Main Estimates"),
            row(kind, 2024, "Main Estimates"),
        ],
    )
    output = tmp_path / "output"
    receipt = normalize(kind, path, output, digest)
    assert receipt["counts"] == {
        "input": 4,
        "normalized": 3,
        "out_of_scope": 0,
        "blank": 0,
        "rejected": 1,
    }
    filename = (
        "budget_facts.parquet" if kind == "expenditure" else "revenue_facts.parquet"
    )
    facts = pq.read_table(output / filename).to_pylist()
    assert [fact["valid_time_end"] for fact in facts] == [
        date(2018, 6, 30),
        date(2022, 6, 30),
        date(2023, 6, 30),
    ]
    assert all(
        f["currency"] is None
        and f["valid_time_start"] is None
        and f["rights_state"] == "not_evaluated"
        and f["unit"] == "$000"
        for f in facts
    )
    args = {
        "source_locator": f"https://www.treasury.govt.nz/sites/default/files/2022-05/b22-{kind}-data.xlsx",
        "source_vintage": "Budget-2022",
        "observed_at": OBSERVED,
    }
    registrations = [
        budget_expenditure_registration(**args),
        budget_revenue_registration(**args),
    ]
    if kind == "expenditure":
        module = import_module(
            "archive_govt_nz.domains.health_appropriations.budget_expenditure_2022"
        )
        registrations.append(module.registration(**args))
    common = dispatch_bronze(
        path.read_bytes(),
        source_sha256=digest,
        media_type=shared.MEDIA_TYPE,
        registrations=tuple(registrations),
    )
    assert list(common.output.records) == facts


@pytest.mark.parametrize(("kind", "bound"), [("expenditure", 6568), ("revenue", 1015)])
def test_2022_native_row_bound_is_admitted_without_truncation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str, bound: int
) -> None:
    other = row(kind)
    other[1] = "Education"
    rows = [other] * (bound - 1) + [row(kind)]
    path, digest = source(tmp_path / "allowed", monkeypatch, kind, rows)
    receipt = normalize(kind, path, tmp_path / "output", digest)
    assert receipt["counts"]["input"] == bound
    assert receipt["counts"]["normalized"] == 1
    path, digest = source(tmp_path / "denied", monkeypatch, kind, [*rows, other])
    denied = tmp_path / "denied-output"
    with pytest.raises(ValueError, match="contract"):
        normalize(kind, path, denied, digest)
    assert not denied.exists()


@pytest.mark.parametrize("kind", ["expenditure", "revenue"])
def test_2022_binding_rejects_other_observations_before_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str
) -> None:
    path, digest = source(tmp_path, monkeypatch, kind, [row(kind)])
    for key, value in {
        "expected_sha256": "0" * 64,
        "source_locator": "https://www.treasury.govt.nz/not-the-reviewed-source.xlsx",
        "source_vintage": "Budget-2023",
    }.items():
        output = tmp_path / key
        with pytest.raises(ValueError, match="contract"):
            normalize(kind, path, output, digest, **{key: value})
        assert not output.exists()

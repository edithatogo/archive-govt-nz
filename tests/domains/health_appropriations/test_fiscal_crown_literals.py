"""Synthetic exact-layout Crown literals; no official amount fixtures."""

from __future__ import annotations

import hashlib
import json
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from archive_govt_nz.domains.health_appropriations import (
    fiscal_crown_literals as subject,
)

TRACK = (
    Path(__file__).resolve().parents[3]
    / "conductor/tracks/health_appropriations_medallion_assimilation_20260829"
)


def profile() -> dict[str, Any]:
    return json.loads((TRACK / "fiscal-crown-literals.profile.json").read_text())


def test_profile_metadata_matches_retained_observation() -> None:
    expected = profile()
    assert expected["headers"] == subject.HEADERS
    assert expected["number_format"] == subject.NUMBER_FORMAT
    assert expected["sheets"] == subject.SHEETS
    assert expected["source_sha256"] == subject.SOURCE_SHA256
    assert [subject.year_label(year) for year in range(1994, 2026)] == expected[
        "year_labels"
    ]


def fixture(
    path: Path, change: tuple[str, str | float | int | None] | None = None
) -> str:
    book = Workbook()
    assert book.active is not None
    book.remove(book.active)
    observed = profile()
    for name in observed["sheets"]:
        book.create_sheet(name)
    sheet = book["Spending"]
    for coordinate, value in observed["headers"].items():
        sheet[coordinate] = value
    sheet["S180"] = "synthetic outside selection"
    for row in range(27, 59):
        year = row + 1967
        sheet[f"B{row}"] = observed["year_labels"][year - 1994]
        for column, first in (("D", 27), ("E", 30)):
            if row >= first:
                sheet[f"{column}{row}"] = row + (0.25 if column == "D" else 1000.5)
                sheet[f"{column}{row}"].number_format = observed["number_format"]
    if change:
        sheet[change[0]] = change[1]
    book.save(path)
    book.close()
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    monkeypatch.setattr(
        subject, "SOURCE_SHA256", hashlib.sha256(path.read_bytes()).hexdigest()
    )
    return subject.admit_fiscal_crown(path)


def test_exact_literals_and_temporal_lineage(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "synthetic.xlsx"
    pin = fixture(source)
    result = run(source, monkeypatch)
    facts = result["facts"]
    assert len(facts) == 61
    assert result["counts"] == {"core_crown": 32, "total_crown": 29}
    assert result["status"] == "literal_context_admitted"
    assert result["promotion"] == "not_performed"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == pin
    assert list(tmp_path.iterdir()) == [source]
    assert len({row["record_id"] for row in facts}) == 61
    for row in facts:
        assert row["amount"] == Decimal(row["source_number_token"])
        is_core = row["family"] == "core_crown"
        offset = Decimal("0.25") if is_core else Decimal("1000.5")
        assert row["amount"] == Decimal(row["year"] - 1967) + offset
        assert row["label"] == (
            "Core Crown Expenses" if is_core else "Total Crown Expenses"
        )
        expected_basis = (
            "old-GAAP"
            if row["year"] <= 1996
            else "IFRS"
            if row["year"] <= 2004
            else "PBE Standards"
        )
        assert row["accounting_basis"] == expected_basis
        assert row["source_number_format"] == profile()["number_format"]
        assert row["currency"] is None
        assert row["price_basis"] is None
        assert row["amount_type"] == "historical_as_published"
        assert row["period_end"] == date(row["year"], 6, 30)
        assert row["period_start"] is None
        assert row["source_object_sha256"] == pin
        assert row["source_vintage"] == "Fiscal-Time-Series-1972-2025"
        assert row["unit"] == "$ millions"
        assert row["scaling"] == "million"
        assert row["rights_state"] == "not_evaluated"
        assert row["consolidation_equivalence"] == "not_asserted"
        assert row["lineage"]["amount"] == row["source_coordinate"]
    first = facts[0]
    assert first["family"] == "core_crown"
    assert first["year_label"] == "1994*"
    assert first["accounting_basis"] == "old-GAAP"
    assert first["lineage"]["period_end"] == "Spending!A23"
    assert first["lineage"]["accounting_basis"] == "Spending!A27"
    total = next(row for row in facts if row["family"] == "total_crown")
    assert total["year"] == 1997
    assert total["accounting_basis"] == "IFRS"
    changed = next(
        row for row in facts if row["family"] == "core_crown" and row["year"] == 2005
    )
    assert changed["accounting_basis"] == "PBE Standards"
    assert changed["lineage"]["period_end"] == "Spending!A38"
    annotated = next(row for row in facts if row["year"] == 2019)
    assert [note["marker"] for note in annotated["source_year_notes"]] == ["^", "#"]
    assert (
        annotated["source_year_notes"][0]["application"]
        == "shared_year_annotation_not_generalized"
    )
    assert result == run(source, monkeypatch)


@pytest.mark.parametrize(
    ("coordinate", "value"),
    [
        ("A3", "NZD millions"),
        ("D4", "Total Crown Expenses"),
        ("E4", "Core Crown Expenses"),
        ("A23", "Cash, March Years"),
        ("A27", "IFRS"),
        ("A30", "Forecast"),
        ("A38", "old-GAAP"),
        ("A24", "new period"),
        ("B27", "1994"),
        ("B30", "1998"),
        ("B52", "2019^"),
        ("A60", "changed note"),
        ("A64", "changed scope"),
        ("A65", None),
        ("D27", None),
        ("D27", "27.25"),
        ("E30", "=SUM(1,2)"),
        ("E30", "#VALUE!"),
        ("E29", 1),
        ("D26", 1),
        ("A68", "new block"),
    ],
)
def test_drift_fails_closed(
    coordinate: str,
    value: str | float | None,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "synthetic.xlsx"
    fixture(source, (coordinate, value))
    with pytest.raises(ValueError, match="fiscal_crown_contract"):
        run(source, monkeypatch)


@pytest.mark.parametrize("fault", ["format", "sheet", "geometry", "precision"])
def test_structural_and_numeric_limits(
    fault: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "synthetic.xlsx"
    fixture(source)
    book = load_workbook(source)
    if fault == "format":
        book["Spending"]["D27"].number_format = "General"
    elif fault == "sheet":
        book.remove(book["Sources"])
    elif fault == "geometry":
        book["Spending"]["S181"] = "extra"
    else:
        book["Spending"]["D27"] = 1e22
    book.save(source)
    book.close()
    with pytest.raises(ValueError, match="fiscal_crown_contract"):
        run(source, monkeypatch)


def test_wrong_hash_missing_source_and_symlink(tmp_path: Path) -> None:
    source = tmp_path / "wrong"
    source.write_bytes(b"synthetic")
    with pytest.raises(ValueError, match="source_hash_mismatch"):
        subject.admit_fiscal_crown(source)
    with pytest.raises(ValueError, match="fiscal_crown_contract"):
        subject.admit_fiscal_crown(tmp_path / "absent")
    link = tmp_path / "linked"
    try:
        link.symlink_to(source)
    except OSError:
        pytest.skip("symlinks unavailable")
    with pytest.raises(ValueError, match="fiscal_crown_contract"):
        subject.admit_fiscal_crown(link)


def test_signed_zero_and_precision(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    for value in (0, -123.5):
        source = tmp_path / "synthetic.xlsx"
        fixture(source, ("D27", value))
        assert run(source, monkeypatch)["facts"][0]["amount"] == Decimal(str(value))

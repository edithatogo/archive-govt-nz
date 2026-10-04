"""Source operations expose metadata, not rows, secrets or publication approval."""

from __future__ import annotations

import ast
import csv
import hashlib
import inspect
import io
import json
import textwrap
from collections.abc import Callable
from dataclasses import replace
from pathlib import Path
from typing import Any, cast

import pyarrow.parquet as pq
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from jsonschema import Draft202012Validator, ValidationError
from openpyxl import load_workbook
from tests.domains.health_appropriations.test_budget_revenue import (
    DEFINITIONS,
    DEFINITIONS_2026_TEST,
    ROW,
)
from tests.domains.health_appropriations.test_budget_revenue import (
    source as budget_revenue_fixture,
)
from tests.domains.health_appropriations.test_cpi import HEADER, META
from tests.domains.health_appropriations.test_forecast_successors import (
    _source as forecast_fixture,
)
from tests.domains.health_appropriations.test_gdp import workbook as gdp_fixture
from tests.domains.health_appropriations.test_pharmac import (
    fixture_source as pharmac_fixture,
)
from tests.domains.health_appropriations.test_population_annual_export import (
    payload as population_annual_payload,
)
from tests.domains.health_appropriations.test_qes import fixture as qes_fixture
from tests.domains.health_appropriations.test_vote_health import (
    TEXT as VOTE_SUMMARY_TEXT,
)
from tests.domains.health_appropriations.test_vote_health_revenue import (
    _pages as vote_revenue_pages,
)

from archive_govt_nz import cli, mcp_server
from archive_govt_nz.cli import app, health_appropriations_extract_source
from archive_govt_nz.domains.health_appropriations import (
    forecast,
    gdp,
    moh_indicators,
    pharmac,
    source_operations,
)
from archive_govt_nz.mcp_server import Server, call_tool, list_tools


@pytest.fixture(
    params=[
        "cpiq-se9a/v1",
        "population-annual-mean-context/v1",
        "budget-revenue-2025/v1",
        "budget-revenue-2026/v1",
        "moh-hair2024-fig27/v1",
        "moh-hair2024-fig28/v1",
        "qes-june2026-table8/v1",
        "pharmac-cpb-20260807/v1",
        "gdp-expenditure-actual-2026q1/v1",
        "gdp-expenditure-actual-2026q2/v1",
        "befu-2026/v1",
        "hyefu-2025/v1",
    ]
)
def request_source(
    tmp_path: Path, request: pytest.FixtureRequest
) -> source_operations.SourceRequest:
    profile = request.param
    source = tmp_path / "source"
    if profile == "population-annual-mean-context/v1":
        source.write_bytes(population_annual_payload())
        vintage = "2026-08-18"
    elif profile == "cpiq-se9a/v1":
        source.write_bytes(
            (
                HEADER + "CPIQ.SE9A,1914.06,1.25" + META + "CPIQ.SE9A,1914.09,NA" + META
            ).encode()
        )
        vintage = "2026-Q2"
    elif profile.startswith("moh"):
        headers = moh_indicators.PROFILES[
            "fig27/v1" if "fig27" in profile else "fig28/v1"
        ]

        content = io.StringIO(newline="")
        writer = csv.writer(content)
        writer.writerow(headers)
        writer.writerows(
            (year, "1.25", "2.5") for year in sorted(moh_indicators.PERIODS)
        )
        source.write_bytes(content.getvalue().encode())
        vintage = "MoH-HAIR-2024"
    elif profile in {"budget-revenue-2025/v1", "budget-revenue-2026/v1"}:
        budget_revenue_fixture(
            tmp_path,
            definitions=(
                DEFINITIONS_2026_TEST
                if profile == "budget-revenue-2026/v1"
                else DEFINITIONS
            ),
            rows=(
                [[*ROW[:5], 58746, 2022, "Actuals"]]
                if profile == "budget-revenue-2026/v1"
                else None
            ),
        )
        generated = tmp_path / "source.xlsx"
        generated.replace(source)
        vintage = (
            "Budget-2026" if profile == "budget-revenue-2026/v1" else "Budget-2025"
        )
    elif profile in {"befu-2026/v1", "hyefu-2025/v1"}:
        forecast_fixture(source, 9 if profile == "befu-2026/v1" else 8)
        vintage = "BEFU-2026" if profile == "befu-2026/v1" else "HYEFU-2025"
    elif profile == "pharmac-cpb-20260807/v1":
        source, _ = pharmac_fixture(tmp_path)
        vintage = "Pharmac-CPB-2026-08-07"
    elif profile in {
        "gdp-expenditure-actual-2026q1/v1",
        "gdp-expenditure-actual-2026q2/v1",
    }:
        vintage = (
            gdp.JUNE_VINTAGE
            if profile == "gdp-expenditure-actual-2026q2/v1"
            else gdp.VINTAGE
        )
        gdp_fixture(source, vintage)
    else:
        qes_fixture(source)
        vintage = "QES-2026-Q2"
    return source_operations.SourceRequest(
        source,
        tmp_path / "output",
        profile,
        hashlib.sha256(source.read_bytes()).hexdigest(),
        vintage,
        "https://example.invalid/source",
        "2026-08-31T00:00:00Z",
    )


def test_preflight_and_explicit_local_write(
    request_source: source_operations.SourceRequest,
) -> None:
    before = request_source.source.read_bytes()
    result = source_operations.operate_source(request_source)
    assert result["status"] == "preflight_passed"
    assert result["profile"] == request_source.profile
    assert "source_locator" not in result
    assert "source_vintage" not in result
    assert not request_source.output_dir.exists()
    written = source_operations.operate_source(request_source, dry_run=False)
    assert written["status"] == "written_local"
    assert written["rights_state"] == "not_evaluated"
    assert written["publication_state"] == "local_validation_only"
    assert len(written["output_sha256"]) == 3
    assert request_source.source.read_bytes() == before


def test_repeated_source_normalization_has_identical_manifest_and_outputs(
    request_source: source_operations.SourceRequest,
) -> None:
    """Every fixture-backed operation profile is repeatable from fixed inputs."""
    source_before = request_source.source.read_bytes()
    first = replace(
        request_source, output_dir=request_source.output_dir.parent / "first"
    )
    second = replace(
        request_source, output_dir=request_source.output_dir.parent / "second"
    )

    first_result = source_operations.operate_source(first, dry_run=False)
    second_result = source_operations.operate_source(second, dry_run=False)

    assert first_result["status"] == second_result["status"] == "written_local"
    assert first_result["profile"] == second_result["profile"] == request_source.profile
    assert first_result["counts"] == second_result["counts"]
    assert first_result["output_sha256"] == second_result["output_sha256"]
    assert {path.name: path.read_bytes() for path in first.output_dir.iterdir()} == {
        path.name: path.read_bytes() for path in second.output_dir.iterdir()
    }
    assert request_source.source.read_bytes() == source_before


@pytest.mark.parametrize(
    "profile",
    [
        "vote-health-supplementary-2003-04-summary/v1",
        "vote-health-supplementary-2003-04-detail/v1",
        "vote-health-supplementary-2003-04-revenue/v1",
    ],
)
def test_repeated_vote_pdf_normalization_has_identical_outputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, profile: str
) -> None:
    """Vote Health PDF operations are byte-repeatable with fixed extracted text."""
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixed PDF fixture bytes")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    vintage = "Treasury-Vote-Health-Supplementary-2003-04"
    context = {
        "expected_sha256": source_hash,
        "source_vintage": vintage,
        "source_locator": "https://example.invalid/vote-health-supplementary.pdf",
        "observed_at": "2026-08-31T00:00:00Z",
    }

    class Page:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self, *, extraction_mode: str) -> str:
            expected_mode = "layout" if profile.endswith("summary/v1") else "plain"
            assert extraction_mode == expected_mode
            return self.text

    class Reader:
        def __init__(self, *_args: object, **_kwargs: object) -> None:
            self.is_encrypted = False
            if profile.endswith("summary/v1"):
                texts = [VOTE_SUMMARY_TEXT]
            elif profile.endswith("detail/v1"):
                texts = [
                    (
                        "Part B1 - Details of Appropriations\n"
                        "Sector Policy 12,459 - 110 - 12,569 - reason"
                    ),
                    "Part E - Statement of Intent",
                ]
            else:
                texts = ["front matter"] * 19 + vote_revenue_pages() + ["blank"]
            self.pages = [Page(text) for text in texts]

    vote_module = (
        source_operations.vote_health_revenue
        if profile.endswith("revenue/v1")
        else source_operations.vote_health
    )
    monkeypatch.setattr(vote_module, "PdfReader", Reader)
    base_request = source_operations.SourceRequest(
        source,
        tmp_path / "unused-output",
        profile,
        **context,
    )
    first = replace(base_request, output_dir=tmp_path / "first")
    second = replace(base_request, output_dir=tmp_path / "second")

    first_result = source_operations.operate_source(first, dry_run=False)
    second_result = source_operations.operate_source(second, dry_run=False)

    assert first_result["status"] == second_result["status"] == "written_local"
    assert first_result["counts"] == second_result["counts"]
    assert first_result["output_sha256"] == second_result["output_sha256"]
    assert {path.name: path.read_bytes() for path in first.output_dir.iterdir()} == {
        path.name: path.read_bytes() for path in second.output_dir.iterdir()
    }
    assert source.read_bytes() == b"fixed PDF fixture bytes"


def test_vote_health_summary_profile_dispatches_to_its_allowlisted_adapter(
    request_source: source_operations.SourceRequest, monkeypatch: pytest.MonkeyPatch
) -> None:
    request = replace(
        request_source,
        profile="vote-health-supplementary-2003-04-summary/v1",
        source_vintage="Treasury-Vote-Health-Supplementary-2003-04",
    )
    expected = {"status": "planned", "counts": {"pages": 1, "facts": 4}}
    monkeypatch.setattr(
        source_operations.vote_health,
        "normalize_vote_health_summary",
        lambda *_args, **_kwargs: expected,
    )
    assert source_operations.operate_source(request)["status"] == "preflight_passed"


def test_vote_health_detail_profile_dispatches_to_its_allowlisted_adapter(
    request_source: source_operations.SourceRequest, monkeypatch: pytest.MonkeyPatch
) -> None:
    request = replace(
        request_source,
        profile="vote-health-supplementary-2003-04-detail/v1",
        source_vintage="Treasury-Vote-Health-Supplementary-2003-04",
    )
    expected = {"status": "planned", "counts": {"pages": 15, "facts": 26}}
    monkeypatch.setattr(
        source_operations.vote_health,
        "normalize_vote_health_detail",
        lambda *_args, **_kwargs: expected,
    )
    assert source_operations.operate_source(request)["status"] == "preflight_passed"


def test_vote_health_2002_03_profile_binds_vintage_and_source_hash(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(b"not the pinned Estimates 2002/03 PDF")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile="vote-health-estimates-2002-03-detail/v1",
        expected_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        source_vintage=source_operations.vote_health.DETAIL_VINTAGE_2002_03,
        source_locator="https://example.test/est02health.pdf",
        observed_at="2026-08-29T09:00:17Z",
    )

    result = source_operations.operate_source(request)

    assert result["status"] == "failed"
    assert result["error"] == "invalid_source_operation"
    assert not request.output_dir.exists()


def test_vote_health_2002_03_profile_dispatches_with_its_exact_vintage(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(b"synthetic profile operation fixture")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    profile = "vote-health-estimates-2002-03-detail/v1"
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=profile,
        expected_sha256=digest,
        source_vintage=source_operations.vote_health.DETAIL_VINTAGE_2002_03,
        source_locator="https://example.test/est02health.pdf",
        observed_at="2026-08-29T09:00:17Z",
    )
    allowed = {
        "vote-health-estimates-detail/v1": (
            source_operations.vote_health.DETAIL_VINTAGE_2003_04,
            None,
        ),
        profile: (request.source_vintage, None),
    }
    monkeypatch.setattr(source_operations, "_VOTE_HEALTH_DETAIL_PROFILES", allowed)
    monkeypatch.setattr(
        source_operations.vote_health,
        "normalize_vote_health_detail",
        lambda *_args, **_kwargs: {
            "status": "planned",
            "counts": {"pages": 26, "facts": 2},
        },
    )

    result = source_operations.operate_source(request)

    assert result["status"] == "preflight_passed"
    assert result["profile"] == profile
    assert result["transformation_id"] == (
        source_operations.vote_health.DETAIL_2002_03_TRANSFORMATION
    )
    assert result["counts"] == {"pages": 26, "facts": 2}
    assert not request.output_dir.exists()


def test_vote_health_revenue_profile_dispatches_to_its_allowlisted_adapter(
    request_source: source_operations.SourceRequest, monkeypatch: pytest.MonkeyPatch
) -> None:
    request = replace(
        request_source,
        profile="vote-health-supplementary-2003-04-revenue/v1",
        source_vintage="Treasury-Vote-Health-Supplementary-2003-04",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 17}}
    monkeypatch.setattr(
        source_operations.vote_health_revenue,
        "normalize_vote_health_revenue",
        lambda *_args, **_kwargs: expected,
    )
    assert source_operations.operate_source(request)["status"] == "preflight_passed"


def test_vote_health_2002_03_revenue_profile_dispatches_to_its_adapter(
    request_source: source_operations.SourceRequest, monkeypatch: pytest.MonkeyPatch
) -> None:
    request = replace(
        request_source,
        profile="vote-health-estimates-2002-03-revenue/v1",
        source_vintage=source_operations.vote_health_revenue.ESTIMATES_2002_03_VINTAGE,
        expected_sha256=source_operations.vote_health_revenue.ESTIMATES_2002_03_SHA256,
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 16}}
    monkeypatch.setattr(
        source_operations.vote_health_revenue,
        "normalize_vote_health_estimates_revenue_2002_03",
        lambda *_args, **_kwargs: expected,
    )
    assert source_operations.operate_source(request)["status"] == "preflight_passed"


def test_vote_health_2002_03_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE,
        expected_sha256=summary.SOURCE_SHA256,
        source_vintage=summary.VINTAGE,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2004-05",
        observed_at="2026-10-03T00:00:00Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 7}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    monkeypatch.setattr(
        source_operations.vote_health_estimates_summary,
        "normalize_vote_health_estimates_overview_2002_03",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        source_operations.vote_health_estimates_summary,
        "_normalize_overview",
        lambda *_args, **_kwargs: expected,
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2004_05_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2004_05,
        expected_sha256=summary.SOURCE_SHA256_2004_05,
        source_vintage=summary.VINTAGE_2004_05,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2004-05",
        observed_at="2026-10-03T00:00:00Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 8}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2004_05",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


@pytest.mark.parametrize("family", ["pharmac", "gdp", "gdp-june"])
def test_extended_dispatch_preserves_source_specific_package(
    tmp_path: Path, family: str
) -> None:
    if family == "pharmac":
        source, pin = pharmac_fixture(tmp_path)
        normalizer = pharmac.normalize_pharmac_budget
        profile, vintage = "pharmac-cpb-20260807/v1", "Pharmac-CPB-2026-08-07"
    elif family == "gdp":
        source = gdp_fixture(tmp_path / "source.xlsx")
        pin = hashlib.sha256(source.read_bytes()).hexdigest()
        normalizer = gdp.normalize_gdp
        profile, vintage = "gdp-expenditure-actual-2026q1/v1", gdp.VINTAGE
    else:
        source = gdp_fixture(tmp_path / "source.xlsx", gdp.JUNE_VINTAGE)
        pin = hashlib.sha256(source.read_bytes()).hexdigest()
        normalizer = gdp.normalize_gdp
        profile, vintage = "gdp-expenditure-actual-2026q2/v1", gdp.JUNE_VINTAGE
    before = source.read_bytes()
    context = {
        "expected_sha256": pin,
        "source_vintage": vintage,
        "source_locator": "https://example.invalid/source",
        "observed_at": "2026-08-31T00:00:00Z",
    }
    direct = tmp_path / "direct"
    raw = normalizer(source, direct, **context, dry_run=False)
    destination = tmp_path / "dispatch"
    request = source_operations.SourceRequest(source, destination, profile, **context)
    result = source_operations.operate_source(request, dry_run=False)
    assert result["status"] == "written_local"
    assert result["counts"] == raw["counts"]
    assert {p.name: p.read_bytes() for p in destination.iterdir()} == {
        p.name: p.read_bytes() for p in direct.iterdir()
    }
    assert source.read_bytes() == before


@pytest.mark.parametrize(
    ("profile", "vintage", "row"),
    [("befu-2026/v1", "BEFU-2026", 9), ("hyefu-2025/v1", "HYEFU-2025", 8)],
)
def test_forecast_dispatch_preserves_package_and_partial_failure(
    tmp_path: Path, profile: str, vintage: str, row: int
) -> None:
    source = tmp_path / "source.xlsx"
    pin = forecast_fixture(source, row)
    context = {
        "expected_sha256": pin,
        "source_vintage": vintage,
        "source_locator": "https://example.invalid/source",
        "observed_at": "2026-08-31T00:00:00Z",
    }
    before = source.read_bytes()
    direct, destination = tmp_path / "direct", tmp_path / "dispatch"
    raw = forecast.normalize_forecast_workbook(
        source, direct, profile=profile, **context, dry_run=False
    )
    request = source_operations.SourceRequest(source, destination, profile, **context)
    result = source_operations.operate_source(request, dry_run=False)
    assert result["status"] == "written_local"
    assert result["counts"] == raw["counts"]
    assert result["transformation_id"] == "treasury-health-expense-summary/v1"
    assert {p.name: p.read_bytes() for p in destination.iterdir()} == {
        p.name: p.read_bytes() for p in direct.iterdir()
    }
    facts = pq.read_table(destination / "forecast_facts.parquet").to_pylist()
    assert [fact["amount_type"] for fact in facts] == ["Actual"] * 5 + ["Forecast"] * 5
    assert source.read_bytes() == before
    workbook = load_workbook(source)
    workbook.worksheets[0].cell(row, 6, "=1+1")
    workbook.save(source)
    workbook.close()
    retained = source.read_bytes()
    partial = replace(
        request,
        output_dir=tmp_path / "partial",
        expected_sha256=hashlib.sha256(retained).hexdigest(),
    )
    failed = source_operations.operate_source(partial)
    assert failed == {
        "schema_version": "archive-govt-nz.health-source-operation/v1",
        "verification_scope": "adapter_execution_only",
        "rights_state": "not_evaluated",
        "publication_state": "local_validation_only",
        "status": "failed",
        "error": "invalid_source_operation",
    }
    assert not partial.output_dir.exists()
    assert source.read_bytes() == retained
    assert source_operations.operate_source(partial, dry_run=False) == failed
    manifest = json.loads((partial.output_dir / "MANIFEST.json").read_bytes())
    assert manifest["status"] == "partial"
    assert manifest["counts"]["rejected"] == 1
    assert {path.name for path in partial.output_dir.iterdir()} == {
        "MANIFEST.json",
        "forecast_facts.parquet",
        "field_lineage.parquet",
        "cell_dispositions.parquet",
    }
    assert source.read_bytes() == retained


@pytest.mark.parametrize("flag", [None, 0, 1, "", "false", [], {}])
def test_non_boolean_flags_fail_closed(
    request_source: source_operations.SourceRequest, flag: object
) -> None:
    result = source_operations.operate_source(
        request_source, dry_run=cast("bool", flag)
    )
    assert result["status"] == "failed"
    assert not request_source.output_dir.exists()


@pytest.mark.parametrize(
    "locator",
    [
        "https://example.invalid/x?secret=private",
        "https://" + ":".join(("user", "private")) + "@example.invalid/x",  # noqa: FLY002 - synthetic rejected userinfo, not literal credentials
        "https://example.invalid/x#private",
    ],
)
def test_sensitive_locator_redacted(
    request_source: source_operations.SourceRequest, locator: str
) -> None:
    result = source_operations.operate_source(
        replace(request_source, source_locator=locator), dry_run=False
    )
    assert result["status"] == "failed"
    assert "private" not in json.dumps(result)
    assert not request_source.output_dir.exists()


def test_cli_mcp_preflight_parity(
    request_source: source_operations.SourceRequest, capsys: pytest.CaptureFixture[str]
) -> None:
    arguments = {
        key: getattr(request_source, key) for key in request_source.__dataclass_fields__
    }
    assert health_appropriations_extract_source(**arguments) == 0
    cli = json.loads(capsys.readouterr().out)
    assert cli.pop("command") == "health-appropriations-extract-source"
    args = {key: str(value) for key, value in arguments.items()}
    assert call_tool("health_appropriations_preflight_source", args) == cli
    tool = next(
        tool
        for tool in list_tools()
        if tool["name"] == "health_appropriations_preflight_source"
    )
    assert tool["annotations"]["readOnlyHint"] is True
    assert "dry_run" not in tool["inputSchema"]["properties"]
    assert not request_source.output_dir.exists()


@pytest.mark.parametrize("malformed", [False, True])
def test_mcp_sensitive_failures_are_redacted(
    request_source: source_operations.SourceRequest, *, malformed: bool
) -> None:
    args = {
        key: str(getattr(request_source, key))
        for key in request_source.__dataclass_fields__
    }
    args["source_locator"] = (
        "https://example.invalid/private-secret?signature=private-secret"
    )
    if malformed:
        args["source_locator"] += "x" * 3000
    server = Server()
    server.handle_request(
        {
            "jsonrpc": "2.0",
            "id": 1,
            "method": "initialize",
            "params": {
                "protocolVersion": "2025-11-25",
                "capabilities": {},
                "clientInfo": {"name": "test", "version": "1"},
            },
        }
    )
    server.handle_request({"jsonrpc": "2.0", "method": "notifications/initialized"})
    result = server.handle_request(
        {
            "jsonrpc": "2.0",
            "id": 2,
            "method": "tools/call",
            "params": {
                "name": "health_appropriations_preflight_source",
                "arguments": args,
            },
        }
    )
    assert "private-secret" not in json.dumps(result)
    assert not request_source.output_dir.exists()


def test_false_mcp_write_flag_rejected(
    request_source: source_operations.SourceRequest,
) -> None:
    args: dict[str, Any] = {
        key: str(getattr(request_source, key))
        for key in request_source.__dataclass_fields__
    }
    args["dry_run"] = False
    result = source_operations.preflight_source(args)
    assert result["status"] == "failed"
    assert not request_source.output_dir.exists()


def test_failed_adapter_result_is_not_success(
    request_source: source_operations.SourceRequest, monkeypatch: pytest.MonkeyPatch
) -> None:
    counts = dict.fromkeys(source_operations.PROFILES[request_source.profile][1], 0)
    monkeypatch.setattr(
        source_operations,
        "_invoke",
        lambda *_args, **_kwargs: {"status": "failed", "counts": counts},
    )
    assert source_operations.operate_source(request_source)["status"] == "failed"
    assert not request_source.output_dir.exists()


@pytest.mark.parametrize(
    "case",
    [
        "profile",
        "digest",
        "timestamp",
        "vintage",
        "locator",
        "missing",
        "source-link",
        "output-link",
        "existing",
    ],
)
def test_invalid_request_creates_no_state(
    request_source: source_operations.SourceRequest, case: str
) -> None:
    changed = request_source
    if case in ("profile", "vintage", "locator"):
        name = case if case == "profile" else "source_" + case
        changed = replace(
            changed, **{name: "private-value" if case != "vintage" else " "}
        )
    elif case == "digest":
        changed = replace(changed, expected_sha256="0" * 64)
    elif case == "timestamp":
        changed = replace(changed, observed_at="2026-08-31T00:00:00")
    elif case == "missing":
        changed = replace(changed, source=changed.source.parent / "missing")
    elif case == "source-link":
        link = changed.source.parent / "source-link"
        link.symlink_to(changed.source)
        changed = replace(changed, source=link)
    elif case == "output-link":
        changed.output_dir.symlink_to(
            changed.source.parent / "absent", target_is_directory=True
        )
    else:
        changed.output_dir.mkdir()
    before = {
        p: p.read_bytes()
        for p in changed.source.parent.rglob("*")
        if p.is_file() and not p.is_symlink()
    }
    result = source_operations.operate_source(changed, dry_run=False)
    assert result["status"] == "failed"
    assert "private-value" not in json.dumps(result)
    assert before == {
        p: p.read_bytes()
        for p in changed.source.parent.rglob("*")
        if p.is_file() and not p.is_symlink()
    }


def test_budget_revenue_profile_requires_matching_source_vintage(
    request_source: source_operations.SourceRequest,
) -> None:
    if request_source.profile not in {
        "budget-revenue-2025/v1",
        "budget-revenue-2026/v1",
    }:
        pytest.skip("only Budget revenue profiles bind a reviewed vintage")
    changed = replace(request_source, source_vintage="Budget-2099")
    result = source_operations.operate_source(changed)
    assert result["status"] == "failed"
    assert result["error"] == "invalid_source_operation"
    assert not changed.output_dir.exists()


@pytest.mark.parametrize("interrupt", [False, True])
def test_parser_failure_and_interrupt(
    request_source: source_operations.SourceRequest,
    monkeypatch: pytest.MonkeyPatch,
    *,
    interrupt: bool,
) -> None:
    def fail(*_args: object, **_kwargs: object) -> None:
        message = "private parser token"
        if interrupt:
            raise KeyboardInterrupt(message)
        raise RuntimeError(message)

    monkeypatch.setattr(source_operations, "_invoke", fail)
    if interrupt:
        with pytest.raises(KeyboardInterrupt):
            source_operations.operate_source(request_source)
    else:
        result = source_operations.operate_source(request_source)
        assert result["status"] == "failed"
        assert "private" not in json.dumps(result)
    assert not request_source.output_dir.exists()


@pytest.mark.parametrize("write", [False, True])
def test_cli_arguments(
    request_source: source_operations.SourceRequest,
    capsys: pytest.CaptureFixture[str],
    *,
    write: bool,
) -> None:
    args = ["health-appropriations-extract-source"]
    for key in request_source.__dataclass_fields__:
        args.extend(["--" + key.replace("_", "-"), str(getattr(request_source, key))])
    if write:
        args.append("--no-dry-run")
    with pytest.raises(SystemExit) as exc:
        app(args, exit_on_error=False)
    assert exc.value.code == 0
    assert json.loads(capsys.readouterr().out)["status"] == (
        "written_local" if write else "preflight_passed"
    )
    assert request_source.output_dir.exists() is write


@pytest.mark.parametrize("bad", [{}, {"secret": 1}, {"input": True}])
def test_malformed_backend_receipt_redacted(
    request_source: source_operations.SourceRequest,
    monkeypatch: pytest.MonkeyPatch,
    bad: dict[str, Any],
) -> None:
    monkeypatch.setattr(
        source_operations,
        "_invoke",
        lambda *_a, **_k: {
            "status": "passed",
            "counts": bad,
            "output_sha256": {"private": "secret"},
        },
    )
    result = source_operations.operate_source(request_source, dry_run=False)
    assert result["status"] == "failed"
    assert "secret" not in json.dumps(result)


def test_partial_failure_is_preserved(
    request_source: source_operations.SourceRequest,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail(*_args: object, **_kwargs: object) -> None:
        request_source.output_dir.mkdir()
        (request_source.output_dir / "partial").write_bytes(b"partial evidence")
        message = "private diagnostic"
        raise OSError(message)

    monkeypatch.setattr(source_operations, "_invoke", fail)
    result = source_operations.operate_source(request_source, dry_run=False)
    assert result["status"] == "failed"
    assert (request_source.output_dir / "partial").read_bytes() == b"partial evidence"
    assert not (request_source.output_dir / "MANIFEST.json").exists()


@given(
    st.sampled_from(tuple(source_operations.PROFILES)),
    st.integers(min_value=0, max_value=10**9),
)
@settings(deadline=None)
def test_receipt_schema_count_property(profile: str, count: int) -> None:
    contract = source_operations.PROFILES[profile]
    receipt = {
        "schema_version": "archive-govt-nz.health-source-operation/v1",
        "verification_scope": "adapter_execution_only",
        "rights_state": "not_evaluated",
        "publication_state": "local_validation_only",
        "status": "preflight_passed",
        "profile": profile,
        "source_object_sha256": "a" * 64,
        "transformation_id": contract[0],
        "counts": dict.fromkeys(contract[1], count),
    }
    validator = Draft202012Validator(source_operations.SOURCE_OPERATION_SCHEMA)
    validator.validate(receipt)
    receipt["rights_state"] = "cleared"
    with pytest.raises(ValidationError):
        validator.validate(receipt)


def test_committed_receipt_schema() -> None:
    path = (
        Path(__file__).resolve().parents[3]
        / "schemas/health-source-operation-v1.schema.json"
    )
    assert json.loads(path.read_text()) == source_operations.SOURCE_OPERATION_SCHEMA


@pytest.mark.parametrize(
    "field", ["source", "output_dir", "source_vintage", "source_locator", "observed_at"]
)
def test_input_schema_exact_length_boundary(field: str) -> None:
    args = {
        "source": "source",
        "output_dir": "output",
        "profile": "cpiq-se9a/v1",
        "expected_sha256": "a" * 64,
        "source_vintage": "vintage",
        "source_locator": "https://example.invalid/x",
        "observed_at": "2026-08-31T00:00:00Z",
    }
    validator = Draft202012Validator(source_operations.SOURCE_PREFLIGHT_INPUT_SCHEMA)
    args[field] = "x" * 2048
    validator.validate(args)
    args[field] += "x"
    with pytest.raises(ValidationError):
        validator.validate(args)


def _seeded_function(
    function: Callable[..., Any], namespace: dict[str, Any], before: str, after: str
) -> Callable[..., Any]:
    """Compile only trusted local function text, without decorators or IO."""
    source = textwrap.dedent(inspect.getsource(function))
    assert source.count(before) == 1
    tree = ast.parse(source.replace(before, after))
    node = tree.body[0]
    assert isinstance(node, ast.FunctionDef)
    node.decorator_list = []
    exec(compile(tree, "<trusted-wiring-counterexample>", "exec"), namespace)  # noqa: S102
    return namespace[function.__name__]


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("dry_run: bool = True", "dry_run: bool = False"),
        ("dry_run=dry_run", "dry_run=True"),
        ('return 2 if result["status"] == "failed" else 0', "return 0"),
    ],
)
def test_cli_seeded_wiring_counterexamples(before: str, after: str) -> None:
    def oracle(function: Callable[..., Any]) -> None:
        arguments = {
            "source": Path("source"),
            "output_dir": Path("output"),
            "profile": "cpiq-se9a/v1",
            "expected_sha256": "a" * 64,
            "source_vintage": "vintage",
            "source_locator": "https://example.invalid/x",
            "observed_at": "2026-08-31T00:00:00Z",
        }
        assert function(**arguments) == 2
        assert seen[-1] is True
        assert function(**arguments, dry_run=False) == 2
        assert seen[-1] is False

    seen: list[bool] = []

    def invoke(_request: object, *, dry_run: bool) -> dict[str, str]:
        seen.append(dry_run)
        return {"status": "failed"}

    namespace = {
        **vars(cli),
        "operate_source": invoke,
        "_emit_json": lambda _result: None,
    }
    baseline = _seeded_function(
        health_appropriations_extract_source, namespace.copy(), before, before
    )
    oracle(baseline)
    mutant = _seeded_function(
        health_appropriations_extract_source, namespace.copy(), before, after
    )
    with pytest.raises(AssertionError):
        oracle(mutant)


def test_mcp_seeded_redaction_counterexample() -> None:
    args = {
        "source": "source",
        "output_dir": "output",
        "profile": "cpiq-se9a/v1",
        "expected_sha256": "a" * 64,
        "source_vintage": "vintage",
        "source_locator": "https://example.invalid/private-secret" + "x" * 3000,
        "observed_at": "2026-08-31T00:00:00Z",
    }
    params = {"name": "health_appropriations_preflight_source", "arguments": args}

    def oracle(function: Callable[..., Any]) -> None:
        assert "private-secret" not in json.dumps(function(Server(), 1, params))

    before = '"Invalid source operation arguments"'
    baseline = _seeded_function(
        Server._call_tool,  # noqa: SLF001 - exact protocol boundary
        vars(mcp_server).copy(),
        before,
        before,
    )
    oracle(baseline)
    mutant = _seeded_function(
        Server._call_tool,  # noqa: SLF001 - exact protocol boundary
        vars(mcp_server).copy(),
        before,
        "errors[0].message",
    )
    with pytest.raises(AssertionError):
        oracle(mutant)


def test_mcp_generic_errors_redact_sensitive_exception_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Generic MCP errors must not expose query credentials or signed URLs."""
    monkeypatch.setattr(
        mcp_server,
        "call_tool",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            ValueError(
                "https://user:"
                + ("sec" + "ret")
                + "@example.invalid/x?signature="
                + ("sec" + "ret")
            )
        ),
    )
    result = mcp_server.Server().handle_request(
        {
            "jsonrpc": "2.0",
            "id": 1,
            "method": "tools/call",
            "params": {"name": "archive_status", "arguments": {}},
        }
    )
    encoded = json.dumps(result)
    assert "secret" not in encoded
    assert "user:" not in encoded


def test_vote_health_2005_06_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2005_06,
        expected_sha256=summary.SOURCE_SHA256_2005_06,
        source_vintage=summary.VINTAGE_2005_06,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2005-06",
        observed_at="2026-10-03T00:00:00Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 8}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2005_06",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_source_operation_profile_registries_have_unique_keys() -> None:
    summary = source_operations.vote_health_estimates_summary
    assert len(source_operations.PROFILES) == len(set(source_operations.PROFILES))
    assert len(source_operations._VOTE_HEALTH_OVERVIEW_PROFILES) == len(  # noqa: SLF001
        set(source_operations._VOTE_HEALTH_OVERVIEW_PROFILES)  # noqa: SLF001
    )
    assert len(source_operations._VOTE_HEALTH_OVERVIEW_NORMALIZERS) == len(  # noqa: SLF001
        set(source_operations._VOTE_HEALTH_OVERVIEW_NORMALIZERS)  # noqa: SLF001
    )
    assert summary.PROFILE_2005_06 in source_operations.PROFILES


def test_vote_health_2006_07_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2006_07,
        expected_sha256=summary.SOURCE_SHA256_2006_07,
        source_vintage=summary.VINTAGE_2006_07,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2006-07",
        observed_at="2026-10-03T00:00:00Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 8}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2006_07",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2007_08_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2007_08,
        expected_sha256=summary.SOURCE_SHA256_2007_08,
        source_vintage=summary.VINTAGE_2007_08,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2007-08",
        observed_at="2026-10-03T00:00:00Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 8}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2007_08",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2008_09_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2008_09,
        expected_sha256=summary.SOURCE_SHA256_2008_09,
        source_vintage=summary.VINTAGE_2008_09,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2008-09",
        observed_at="2026-10-03T12:34:00.964355Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 6}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2008_09",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2009_10_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2009_10,
        expected_sha256=summary.SOURCE_SHA256_2009_10,
        source_vintage=summary.VINTAGE_2009_10,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2009-10",
        observed_at="2026-10-03T16:31:35.751358Z",
    )
    expected = {"status": "planned", "counts": {"pages": 1, "facts": 17}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2009_10",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2010_11_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2010_11,
        expected_sha256=summary.SOURCE_SHA256_2010_11,
        source_vintage=summary.VINTAGE_2010_11,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2010-11",
        observed_at="2026-10-03T18:31:59.293780Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 18}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2010_11",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2011_12_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2011_12,
        expected_sha256=summary.SOURCE_SHA256_2011_12,
        source_vintage=summary.VINTAGE_2011_12,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2011-12",
        observed_at="2026-10-03T20:06:48.631115Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 17}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2011_12",
        lambda *_args, **_kwargs: expected,
    )
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2012_13_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2012_13,
        expected_sha256=summary.SOURCE_SHA256_2012_13,
        source_vintage=summary.VINTAGE_2012_13,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2012-13",
        observed_at="2026-10-03T22:37:57.436271Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 18}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2012_13",
        lambda *_args, **_kwargs: expected,
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001


def test_vote_health_2013_14_overview_profile_dispatches_to_its_adapter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = source_operations.vote_health_estimates_summary
    source = tmp_path / "source.pdf"
    source.write_bytes(b"fixture PDF source; adapter is mocked")
    request = source_operations.SourceRequest(
        source=source,
        output_dir=tmp_path / "output",
        profile=summary.PROFILE_2013_14,
        expected_sha256=summary.SOURCE_SHA256_2013_14,
        source_vintage=summary.VINTAGE_2013_14,
        source_locator="https://www.treasury.govt.nz/publications/estimates/vote-health-estimates-appropriations-2013-14",
        observed_at="2026-10-04T00:04:54.340387Z",
    )
    expected = {"status": "planned", "counts": {"pages": 2, "facts": 17}}
    monkeypatch.setattr(source_operations, "_validate", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        summary, "_normalize_overview", lambda *_args, **_kwargs: expected
    )
    monkeypatch.setattr(
        summary,
        "normalize_vote_health_estimates_overview_2013_14",
        lambda *_args, **_kwargs: expected,
    )
    assert source_operations._invoke(request, dry_run=True) == expected  # noqa: SLF001

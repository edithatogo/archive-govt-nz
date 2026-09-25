"""Edition-specific Budget revenue dispatch preserves facts and row lineage."""

from __future__ import annotations

import hashlib
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import Workbook

from archive_govt_nz.domains.health_appropriations import budget_revenue
from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    dispatch_bronze,
)
from archive_govt_nz.domains.health_appropriations.budget_adapter import (
    budget_expenditure_registration,
)
from archive_govt_nz.domains.health_appropriations.budget_revenue_adapter import (
    BudgetRevenueAdapter,
    budget_revenue_registration,
)


def _revenue_workbook(
    *,
    definitions: dict[str, str] | None = None,
    headers: list[str] | None = None,
    year: int = 2021,
) -> bytes:
    book = Workbook()
    raw = book.active
    assert raw is not None
    raw.title = "Raw Data"
    raw.append(list(budget_revenue.FIELDS) if headers is None else headers)
    raw.append(
        [
            "Ministry of Health",
            "Health",
            397,
            "Hospital reimbursement",
            "Non-Tax Revenue",
            58746,
            year,
            "Actuals",
        ]
    )
    explanation = book.create_sheet("Explanation")
    for coordinate, value in (
        budget_revenue.DEFINITIONS if definitions is None else definitions
    ).items():
        explanation[coordinate] = value
    intro = book.create_sheet("Intro")
    for coordinate in ("A10", "A13", "A14"):
        intro[coordinate] = "synthetic source-context notice"
    book.create_sheet("Pivot Trend by Vote")
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def _registrations() -> tuple:
    context = {
        "source_locator": "data/raw/b25-revenue-data.xlsx",
        "source_vintage": "Budget-2025",
        "observed_at": "2026-08-30T00:00:00Z",
    }
    return (
        budget_expenditure_registration(**context),
        budget_revenue_registration(**context),
    )


def _adapter(vintage: str = "Budget-2025") -> BudgetRevenueAdapter:
    return BudgetRevenueAdapter(
        source_locator="data/raw/b25-revenue-data.xlsx",
        source_vintage=vintage,
        observed_at="2026-08-30T00:00:00Z",
    )


def test_dispatch_routes_revenue_workbook_to_edition_bound_adapter() -> None:
    payload = _revenue_workbook()
    digest = hashlib.sha256(payload).hexdigest()
    result = dispatch_bronze(
        payload,
        source_sha256=digest,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        registrations=_registrations(),
    )

    assert result.selection.adapter_id == "nz-budget-health-revenue"
    assert result.selection.considered_adapter_ids == (
        "nz-budget-health-expenditure",
        "nz-budget-health-revenue",
    )
    assert result.selection.matched_adapter_ids == ("nz-budget-health-revenue",)
    assert result.output.layout == "budget-revenue/v1"
    assert len(result.output.records) == 1
    fact = result.output.records[0]
    assert fact["recordset"] == "budget_revenue_fact"
    assert fact["amount"] == 58746
    assert fact["source_vintage"] == "Budget-2025"
    assert fact["currency"] is None
    assert fact["rights_state"] == "not_evaluated"
    amount = next(row for row in result.output.lineage if row.field == "amount")
    assert amount.source_coordinate == "'Raw Data'!F2"
    assert amount.raw_value == "58746"
    assert amount.normalized_value == "58746.000"
    assert {row.reason for row in result.output.losses} == {
        "context_or_non_revenue_sheet_inventoried_only"
    }


def test_ambiguous_budget_layouts_are_preserved_only() -> None:
    payload = _revenue_workbook()
    context = {
        "source_locator": "data/raw/b25-revenue-data.xlsx",
        "source_vintage": "Budget-2025",
        "observed_at": "2026-08-30T00:00:00Z",
    }
    revenue = budget_revenue_registration(**context)
    second = replace(revenue, adapter_id="second-revenue-adapter")
    result = dispatch_bronze(
        payload,
        source_sha256=hashlib.sha256(payload).hexdigest(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        registrations=(budget_expenditure_registration(**context), revenue, second),
    )
    assert result.selection.status == "preserved_only"
    assert result.selection.reason == "adapter_selection_ambiguous"
    assert result.selection.matched_adapter_ids == (
        "nz-budget-health-revenue",
        "second-revenue-adapter",
    )
    assert result.output.records == ()


def test_revenue_adapter_rejects_unknown_layouts_and_editions() -> None:
    payload = _revenue_workbook()
    assert not _adapter().matches_layout(b"not an xlsx")
    assert not _adapter("Budget-2024").matches_layout(payload)

    invalid = _adapter().extract(
        b"not an xlsx", source_sha256=hashlib.sha256(b"not an xlsx").hexdigest()
    )
    assert invalid.layout == "unknown"
    assert invalid.records == ()
    with pytest.raises(ValueError, match="unsupported_budget_revenue_vintage"):
        _adapter("Budget-2024").extract(
            payload, source_sha256=hashlib.sha256(payload).hexdigest()
        )
    with pytest.raises(ValueError, match="source_hash_mismatch"):
        _adapter().extract(payload, source_sha256="0" * 64)


def test_revenue_layout_probe_rejects_unreviewed_headers_and_metadata() -> None:
    assert not _adapter().matches_layout(_revenue_workbook(headers=["Vote", "Amount"]))
    changed_definitions = dict(budget_revenue.DEFINITIONS)
    changed_definitions["B3"] = "different edition text"
    assert not _adapter().matches_layout(
        _revenue_workbook(definitions=changed_definitions)
    )


def test_revenue_adapter_rejects_non_bytes_payload() -> None:
    with pytest.raises(ValueError, match="budget_revenue_source_limit_or_type"):
        _adapter().extract(None, source_sha256="0" * 64)  # type: ignore[arg-type]


def test_revenue_adapter_selects_the_embedded_2026_edition() -> None:
    payload = _revenue_workbook(definitions=budget_revenue.DEFINITIONS_2026, year=2022)
    registration = budget_revenue_registration(
        source_locator="data/raw/b26-revenue-data.xlsx",
        source_vintage="Budget-2026",
        observed_at="2026-08-30T00:00:00Z",
    )
    result = dispatch_bronze(
        payload,
        source_sha256=hashlib.sha256(payload).hexdigest(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        registrations=(registration,),
    )
    assert result.selection.adapter_id == "nz-budget-health-revenue"
    assert result.output.records[0]["source_vintage"] == "Budget-2026"
    assert (
        result.output.records[0]["transformation_id"] == "budget-2026-health-revenue/v1"
    )

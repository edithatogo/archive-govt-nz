"""Source-specific revenue facts must not become spending or rights approval."""

from __future__ import annotations

import hashlib
import json
import tempfile
from datetime import date
from decimal import Decimal, localcontext
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO, cast
from zipfile import ZipFile

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import Workbook

from archive_govt_nz.domains.health_appropriations import budget_revenue as revenue

HEADERS = [
    "Department",
    "Vote",
    "App ID",
    "Description",
    "Revenue Type",
    "Amount $000",
    "Year",
    "Amount Type",
]
ROW = [
    "Ministry of Health",
    "Health",
    397,
    "Hospital reimbursement",
    "Non-Tax Revenue",
    58746,
    2021,
    "Actuals",
]
DEFINITIONS = {
    "B3": (
        "The Revenue workbook contains details of  actual government Crown revenue "
        "and capital receipts for the years ended 30 June 2021, 2022, 2023 and 2024; "
        "estimated actual government Crown revenue and capital receipts for the year "
        "ending 30 June 2025 and budgeted government Crown revenue and capital "
        "receipts for the year ending 30 June 2026 as published in Budget 2025."
    ),
    "B37": "Amount $000: Amount (in thousands) for the Year as reported in the Main Estimates.",
    "B38": "Year: Year ending that the Amount relates to (at 30 June) as reported in the Main Estimates.",
    "B39": 'Amount Type: "Actuals" - as audited for prior Years, "Estimated Actual" - for the Year immediately prior to the current Main Estimates, "Main Estimates" - for the Main Estimates Year.',
    "B43": "App ID: Number used to uniquely identify each Crown revenue or capital receipt line.",
}


def source(
    tmp_path: Path,
    rows: list[list[object]] | None = None,
    change: tuple[str, str, object] | None = None,
) -> Path:
    book = Workbook()
    raw = book.active
    assert raw is not None
    raw.title = "Raw Data"
    raw.append(HEADERS)
    for row in [ROW] if rows is None else rows:
        raw.append(row)
    explanation = book.create_sheet("Explanation")
    for cell, text in DEFINITIONS.items():
        explanation[cell] = text
    intro = book.create_sheet("Intro")
    for cell in ("A10", "A13", "A14"):
        intro[cell] = "Synthetic notice observation; no permission asserted."
    book.create_sheet("Pivot Trend by Vote")["A1"] = "retained"
    if change:
        sheet, cell, value = change
        book[sheet][cell] = cast("Any", value)
    path = tmp_path / "source.xlsx"
    book.save(path)
    book.close()
    return path


def run(path: Path, output: Path) -> dict[str, Any]:
    return revenue.normalize_budget_revenue(
        path,
        output,
        expected_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        source_locator="data/raw/b25-revenue-data.xlsx",
        observed_at="2026-08-30T00:00:00Z",
    )


def test_real_contract_shape_distinct_rows_and_complete_accounting(
    tmp_path: Path,
) -> None:
    original = source(
        tmp_path,
        [
            ROW,
            [*ROW[:4], "Capital Receipts", 0, 2026, "Main Estimates"],
            [],
            ["Other", "Education", *ROW[2:]],
            ROW,
        ],
    )
    before = original.read_bytes()
    receipt = run(original, tmp_path / "out")
    assert receipt["counts"] == {
        "input": 5,
        "normalized": 3,
        "out_of_scope": 1,
        "blank": 1,
        "rejected": 0,
    }
    facts = pq.read_table(tmp_path / "out/revenue_facts.parquet").to_pylist()
    assert len({r["record_id"] for r in facts}) == 3
    assert [r["amount"] for r in facts] == [Decimal(58746), Decimal(0), Decimal(58746)]
    assert facts[0]["valid_time_end"] == date(2021, 6, 30)
    assert facts[0]["valid_time_start"] is None
    assert all(r["unit"] == "$000" and r["currency"] is None for r in facts)
    assert all(
        r["rights_state"] == "not_evaluated" and r["recordset"] == "budget_revenue_fact"
        for r in facts
    )
    assert [r["revenue_type"] for r in facts] == [
        "Non-Tax Revenue",
        "Capital Receipts",
        "Non-Tax Revenue",
    ]
    assert {r["sheet"] for r in receipt["excluded_sheets"]} == {
        "Intro",
        "Explanation",
        "Pivot Trend by Vote",
    }
    assert receipt["embedded_notice"]["eligibility_state"] == "not_assessed"
    assert (
        receipt["embedded_notice"]["source_object_sha256"]
        == hashlib.sha256(before).hexdigest()
    )
    run(original, tmp_path / "two")
    for file in (tmp_path / "out").iterdir():
        assert file.read_bytes() == (tmp_path / "two" / file.name).read_bytes()
    assert original.read_bytes() == before


@pytest.mark.parametrize(
    ("column", "value", "reason"),
    [
        ("F", "=1+1", "formula_not_evaluated"),
        ("F", "#VALUE!", "spreadsheet_error"),
        ("F", None, "invalid_amount"),
        ("F", "1.0001", "invalid_amount"),
        ("F", True, "invalid_amount"),
        ("G", 2027, "invalid_period_type"),
        ("H", "Forecast", "invalid_period_type"),
        ("E", "Tax Revenue", "invalid_revenue_type"),
        ("C", -1, "invalid_app_id"),
        ("D", None, "missing_label"),
        ("D", 2, "missing_label"),
        ("B", None, "invalid_vote"),
        ("B", "", "invalid_vote"),
        ("B", "=1", "invalid_vote"),
        ("C", True, "invalid_app_id"),
        ("C", 2**63, "invalid_app_id"),
        ("G", True, "invalid_period_type"),
        ("F", 1e17, "invalid_amount"),
    ],
)
def test_rejected_rows_are_not_lost(
    tmp_path: Path, column: str, value: object, reason: str
) -> None:
    original = source(tmp_path, change=("Raw Data", column + "2", value))
    receipt = run(original, tmp_path / "out")
    assert receipt["status"] == "partial"
    assert receipt["counts"]["rejected"] == 1
    rows = pq.read_table(tmp_path / "out/row_dispositions.parquet").to_pylist()
    assert rows[0]["reason"] == reason
    assert rows[0]["record_id"] is None
    assert pq.read_table(tmp_path / "out/revenue_facts.parquet").num_rows == 0


@pytest.mark.parametrize(
    ("sheet", "cell", "value"),
    [
        ("Raw Data", "F1", "Amount NZD"),
        ("Explanation", "B38", "Calendar year"),
        ("Explanation", "B3", "Budget 2026"),
        ("Intro", "A13", "=1"),
    ],
)
def test_unknown_context_fails_before_output(
    tmp_path: Path, sheet: str, cell: str, value: object
) -> None:
    original = source(tmp_path, change=(sheet, cell, value))
    with pytest.raises(ValueError, match="budget_revenue_contract"):
        run(original, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_all_fields_and_semantic_evidence_have_lineage(tmp_path: Path) -> None:
    original = source(tmp_path)
    receipt = run(original, tmp_path / "out")
    fact = pq.read_table(tmp_path / "out/revenue_facts.parquet").to_pylist()[0]
    links = pq.read_table(tmp_path / "out/field_lineage.parquet").to_pylist()
    assert len(links) == 16
    for column, field in zip(
        "ABCDEFGH",
        (
            "department",
            "vote",
            "app_id",
            "description",
            "revenue_type",
            "amount",
            "year",
            "amount_type",
        ),
        strict=True,
    ):
        link = next(
            r
            for r in links
            if r["source_coordinate"] == f"'Raw Data'!{column}2" and r["field"] == field
        )
        assert link["normalized_value"] == str(fact[field])
        assert link["record_id"] == fact["record_id"]
        assert link["source_object_sha256"] == receipt["source_object_sha256"]
    assert {
        (r["field"], r["source_coordinate"])
        for r in links
        if "Explanation" in r["source_coordinate"]
    } == {
        ("source_vintage", "'Explanation'!B3"),
        ("amount_type", "'Explanation'!B3"),
        ("unit", "'Explanation'!B37"),
        ("valid_time_end", "'Explanation'!B38"),
        ("amount_type", "'Explanation'!B39"),
        ("app_id", "'Explanation'!B43"),
    }
    assert json.loads(fact["raw_values_json"]) == dict(zip(HEADERS, ROW, strict=True))
    assert fact["source_number_token"] == str(58746)
    assert pq.read_schema(tmp_path / "out/revenue_facts.parquet") == revenue.FACT_SCHEMA
    for name, digest in receipt["output_sha256"].items():
        assert (
            hashlib.sha256((tmp_path / "out" / name).read_bytes()).hexdigest() == digest
        )
    for observation in receipt["embedded_notice"]["observations"]:
        assert (
            observation["decoded_text_sha256"]
            == hashlib.sha256(
                b"Synthetic notice observation; no permission asserted."
            ).hexdigest()
        )


@pytest.mark.parametrize("value", [-123, 0, 1.125])
def test_exact_values_ignore_callers_decimal_context(
    tmp_path: Path, value: float
) -> None:
    original = source(tmp_path, change=("Raw Data", "F2", value))
    with localcontext() as context:
        context.prec = 2
        before = context.copy()
        run(original, tmp_path / "out")
        assert context.prec == before.prec
        assert context.flags == before.flags
    assert pq.read_table(tmp_path / "out/revenue_facts.parquet").to_pylist()[0][
        "amount"
    ] == Decimal(str(value))


def test_literal_token_prevents_float_rounding(tmp_path: Path) -> None:
    original = source(tmp_path, change=("Raw Data", "F2", 1.125))
    rewritten = BytesIO()
    with (
        ZipFile(BytesIO(original.read_bytes())) as archive,
        ZipFile(rewritten, "w") as out,
    ):
        for item in archive.infolist():
            payload = archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b">1.125<" in payload
                payload = payload.replace(b">1.125<", b">1.12500000000000001<")
            out.writestr(item, payload)
    original.write_bytes(rewritten.getvalue())
    receipt = run(original, tmp_path / "out")
    assert receipt["counts"]["rejected"] == 1
    rejected = pq.read_table(tmp_path / "out/row_dispositions.parquet").to_pylist()[0]
    assert rejected["source_number_token"] == str(Decimal("1.12500000000000001"))
    assert pq.read_table(tmp_path / "out/revenue_facts.parquet").num_rows == 0


def test_hash_limits_symlinks_and_nonoverwrite(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original = source(tmp_path)
    before = original.read_bytes()
    with pytest.raises(ValueError, match="source_hash_mismatch"):
        revenue.normalize_budget_revenue(
            original,
            tmp_path / "out",
            expected_sha256="a" * 64,
            source_locator="fixture",
            observed_at="2026-08-30T00:00:00Z",
        )
    monkeypatch.setattr(revenue, "MAX_BYTES", len(before) - 1)
    with pytest.raises(ValueError, match="source_byte_limit"):
        run(original, tmp_path / "out")
    monkeypatch.setattr(revenue, "MAX_BYTES", len(before))
    run(original, tmp_path / "out")
    with pytest.raises(FileExistsError):
        run(original, tmp_path / "out")
    link = tmp_path / "linked"
    link.symlink_to(original)
    with pytest.raises(ValueError, match="budget_revenue_contract"):
        run(link, tmp_path / "unused")
    output_link = tmp_path / "output-linked"
    output_link.symlink_to(tmp_path / "out", target_is_directory=True)
    with pytest.raises(ValueError, match="budget_revenue_contract"):
        run(original, output_link)
    assert original.read_bytes() == before


def test_interrupted_write_keeps_partial_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original = source(tmp_path)
    before = original.read_bytes()
    real_write = pq.write_table
    calls = 0

    def write(table: pa.Table, handle: BinaryIO) -> None:
        nonlocal calls
        calls += 1
        if calls == 2:
            message = "injected interruption"
            raise OSError(message)
        real_write(table, handle)

    monkeypatch.setattr(pq, "write_table", write)
    with pytest.raises(OSError, match="injected interruption"):
        run(original, tmp_path / "out")
    assert {p.name for p in (tmp_path / "out").iterdir()} == {
        "revenue_facts.parquet",
        "field_lineage.parquet",
    }
    assert (tmp_path / "out/field_lineage.parquet").stat().st_size == 0
    assert original.read_bytes() == before


@pytest.mark.parametrize("rows", [[], [["Other", "Education", *ROW[2:]]]])
def test_empty_health_selection_is_explicit(
    tmp_path: Path, rows: list[list[object]]
) -> None:
    receipt = run(source(tmp_path, rows), tmp_path / "out")
    assert receipt["status"] == "empty"
    assert receipt["counts"]["normalized"] == 0


@settings(max_examples=10)
@given(
    st.lists(st.integers(min_value=-100000, max_value=100000), min_size=1, max_size=5)
)
def test_duplicate_values_preserve_source_occurrences(amounts: list[int]) -> None:
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        rows = [[*ROW[:5], amount, *ROW[6:]] for amount in amounts]
        run(source(root, rows), root / "out")
        facts = pq.read_table(root / "out/revenue_facts.parquet").to_pylist()
        assert [r["amount"] for r in facts] == [Decimal(a) for a in amounts]
        assert len({r["record_id"] for r in facts}) == len(amounts)

"""Bind tobac source identities to verified native CAS/WARC capture receipts.

This offline postflight never treats HTTP success as source-document acceptance.
Run after the existing capture CLI; preserve each run in a fresh directory.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import BadZipFile

from openpyxl import load_workbook
from warcio.archiveiterator import ArchiveIterator

from archive_govt_nz.object_store import ContentAddressedStore

_TITLE_OVERLAP = 0.6


def workbook_matches(body: bytes, asset: dict[str, Any]) -> bool:
    """Compare an original workbook with the preserved consumer extraction."""
    try:
        book = load_workbook(BytesIO(body), read_only=True, data_only=True)
        try:
            rows = list(book[book.sheetnames[0]].values)
            expected = [asset["expected_header"], *asset["expected_rows"]]
            return [list(row) for row in rows] == expected
        finally:
            book.close()
    except ValueError, KeyError, OSError, BadZipFile:
        return False


def classify(body: bytes, source: dict[str, Any], uri: str) -> str:  # noqa: PLR0911
    """Recognise challenge responses and require the expected source identity."""
    text = body.decode("utf-8", errors="replace")
    lower = text.lower()
    challenges = (
        "cookies must be enabled",
        "_incapsula_resource",
        "cf-chl-",
        "just a moment...",
        "checking your browser",
        "verify you are human",
    )
    if any(marker in lower for marker in challenges):
        return "challenge_not_source"
    for asset in source.get("assets", []):
        if uri == asset["url"] and asset["kind"] == "xlsx":
            return (
                "source_workbook_values_matched"
                if workbook_matches(body, asset)
                else "identity_unverified"
            )
    api = source.get("metadata_api", {})
    if uri == api.get("url"):
        if "<!ENTITY" in text.upper():
            return "identity_unverified"
        try:
            # Entity declarations are rejected above; parse only bounded CAS bytes.
            document = ET.fromstring(text)  # noqa: S314
        except ET.ParseError:
            return "identity_unverified"
        identifiers = {
            node.text for node in document.findall("PubmedArticle/MedlineCitation/PMID")
        }
        return (
            "metadata_only"
            if document.tag == "PubmedArticleSet" and set(api["pmids"]) <= identifiers
            else "identity_unverified"
        )
    title = re.search(r"<title[^>]*>(.*?)</title>", text, re.DOTALL | re.IGNORECASE)
    if title is None:
        return "identity_unverified"
    words = set(re.findall(r"\w+", source["title"].lower()))
    actual = set(re.findall(r"\w+", html.unescape(title[1]).lower()))
    # Conservative discovery screen, not validation of every table or supplement.
    if len(words & actual) / max(len(words), 1) < _TITLE_OVERLAP:
        return "identity_unverified"
    return "source_html_identity_observed"


def bind(  # noqa: C901
    register: dict[str, Any],
    receipts: list[dict[str, Any]],
    store: Path,
    root: Path,
) -> dict[str, Any]:
    """Verify stored bytes and WARC payload before returning consumer bindings."""
    targets: dict[str, dict[str, Any]] = {}
    for receipt in receipts:
        for target in receipt["targets"]:
            if target["uri"] in targets:
                msg = "duplicate_capture_uri_choose_one_explicit_run"
                raise ValueError(msg)
            targets[target["uri"]] = target
    cas = ContentAddressedStore(store, create=False)
    sources = []
    for source in register["sources"]:
        captures = []
        urls = list(source["urls"])
        urls.extend(asset["url"] for asset in source.get("assets", []))
        if source.get("metadata_api"):
            urls.append(source["metadata_api"]["url"])
        for uri in urls:
            target = targets.get(uri, {})
            entry: dict[str, Any] = {"uri": uri, "status": "not_captured"}
            if target.get("status") == "failed":
                entry.update(status="retrieval_failed", error=target["error"])
            elif target.get("status") == "captured":
                obj = cas.verify("sha256:" + target["sha256"])
                body = obj.path.read_bytes()
                warc = target["warc"]
                warc_path = root / warc["path"]
                if hashlib.sha256(warc_path.read_bytes()).hexdigest() != warc["sha256"]:
                    msg = "warc_fixity_mismatch"
                    raise ValueError(msg)
                with warc_path.open("rb") as stream:
                    records = list(ArchiveIterator(stream))
                    if len(records) != 1:
                        msg = "expected_single_response_record"
                        raise ValueError(msg)
                # Read payload while its stream is still open.
                with warc_path.open("rb") as stream:
                    record = next(ArchiveIterator(stream))
                    payload = record.content_stream().read()
                    if payload != body:
                        msg = "warc_cas_payload_mismatch"
                        raise ValueError(msg)
                    request_hash = record.rec_headers.get_header(
                        "WARC-Request-URL-SHA256"
                    )
                    if request_hash != hashlib.sha256(uri.encode()).hexdigest():
                        message = "warc_request_identity_mismatch"
                        raise ValueError(message)
                    captured_at = record.rec_headers.get_header("WARC-Date")
                    final_uri = record.rec_headers.get_header("WARC-Target-URI")
                entry.update(
                    status=classify(body, source, uri),
                    sha256=obj.sha256,
                    byte_count=obj.byte_count,
                    warc_sha256=warc["sha256"],
                    captured_at=captured_at,
                    final_uri=final_uri,
                )
            captures.append(entry)
        sources.append(
            {
                "archive_id": source["archive_id"],
                "references": source["references"],
                "captures": captures,
                "redistribution_authorized": source["redistribution_authorized"],
            }
        )
    return {
        "schema_version": "archive-govt-nz.tobac-bindings/v1",
        "consumer_commit": register["consumer_commit"],
        "collection_sha256": register["collection_sha256"],
        "sources": sources,
        "qualification": (
            "CAS/WARC fixity and identity screen only; "
            "not extraction, scientific, full-document or redistribution approval"
        ),
    }


def main() -> None:
    """Write a new receipt, leaving all previous capture generations intact."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--register", type=Path, required=True)
    parser.add_argument("--capture", type=Path, action="append", required=True)
    parser.add_argument("--store", type=Path, required=True)
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    register = json.loads(args.register.read_text())
    result = bind(
        register,
        [json.loads(p.read_text()) for p in args.capture],
        args.store,
        args.root,
    )
    result["register_sha256"] = hashlib.sha256(args.register.read_bytes()).hexdigest()
    result["capture_receipt_sha256"] = [
        hashlib.sha256(p.read_bytes()).hexdigest() for p in args.capture
    ]
    with args.out.open("x", encoding="utf-8") as stream:
        json.dump(result, stream, indent=2, ensure_ascii=False)
        stream.write("\n")


if __name__ == "__main__":
    main()

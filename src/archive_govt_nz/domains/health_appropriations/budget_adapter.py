"""Dispatch adapter for the named-column Budget Health expenditure profile."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook

from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    AdapterRegistration,
)
from archive_govt_nz.domains.health_appropriations.adapter_protocol import (
    AdapterOutput,
    FieldLineage,
    LossAccounting,
    preserved_only,
)
from archive_govt_nz.domains.health_appropriations.budget import _extract
from archive_govt_nz.domains.health_appropriations.formats import inventory_workbook
from archive_govt_nz.domains.health_appropriations.workbook_common import (
    source_context,
)

_MAX_SOURCE_BYTES = 64 * 1024 * 1024
_SHEET = "Raw Data"
_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_SOURCE_BYTE_LIMIT_ERROR = "source_byte_limit"
_SOURCE_HASH_ERROR = "source_hash_mismatch"


@dataclass(frozen=True, slots=True)
class BudgetExpenditureAdapter:
    """Adapt the explicitly supported named-column Budget workbook layout.

    Context is supplied by the caller because bytes and their digest alone do
    not establish a source vintage, observation time, or capture locator.
    Unknown workbook layouts stay preserved in Bronze and emit no facts.
    """

    source_locator: str
    source_vintage: str
    observed_at: str

    def extract(self, bronze: bytes, *, source_sha256: str) -> AdapterOutput:
        """Return typed records, field lineage, and losses from the profile."""
        if len(bronze) > _MAX_SOURCE_BYTES:
            raise ValueError(_SOURCE_BYTE_LIMIT_ERROR)
        actual = hashlib.sha256(bronze).hexdigest()
        if actual != source_sha256:
            raise ValueError(_SOURCE_HASH_ERROR)
        context = source_context(
            source_sha256, self.source_locator, self.source_vintage, self.observed_at
        )
        try:
            inventory_workbook(BytesIO(bronze))
            workbook = load_workbook(BytesIO(bronze), data_only=False, keep_links=True)
        except BadZipFile, OSError, ValueError, KeyError, TypeError, EOFError:
            return preserved_only(
                source_coordinate="bronze:sha256:" + actual,
                reason="invalid_budget_workbook",
            )
        try:
            if _SHEET not in workbook.sheetnames:
                return preserved_only(
                    source_coordinate="bronze:sha256:" + actual,
                    reason="unsupported_budget_layout",
                )
            try:
                facts, lineage, dispositions = _extract(workbook[_SHEET], context)
            except ValueError as error:
                if str(error) != "invalid_headers":
                    raise
                return preserved_only(
                    source_coordinate="'Raw Data'!A1",
                    reason="unsupported_budget_headers",
                )
            losses = [
                LossAccounting(
                    source_coordinate=f"'Raw Data'!row:{row['source_row']}",
                    disposition=str(row["disposition"]),
                    reason=str(row["reason"]),
                )
                for row in dispositions
                if row["disposition"] != "normalized"
            ]
            losses.extend(
                LossAccounting(
                    source_coordinate=f"sheet:{sheet.title}",
                    disposition="excluded",
                    reason="not_budget_raw_data_inventoried_only",
                )
                for sheet in workbook.worksheets
                if sheet.title != _SHEET
            )
            return AdapterOutput(
                records=tuple(dict(row) for row in facts),
                losses=tuple(losses),
                lineage=tuple(
                    FieldLineage(
                        record_id=str(row["record_id"]),
                        field=str(row["field"]),
                        source_coordinate=str(row["source_coordinate"]),
                        raw_value=str(row["raw_value"]),
                        normalized_value=str(row["normalized_value"]),
                        rule=str(row["rule"]),
                    )
                    for row in lineage
                ),
                layout="budget-expenditure/v1",
            )
        finally:
            workbook.close()


def budget_expenditure_registration(
    *, source_locator: str, source_vintage: str, observed_at: str
) -> AdapterRegistration:
    """Return an explicit registration for the supported Budget XLSX profile."""
    return AdapterRegistration(
        adapter_id="nz-budget-health-expenditure",
        version="1.0.0",
        media_type=_MEDIA_TYPE,
        adapter=BudgetExpenditureAdapter(
            source_locator=source_locator,
            source_vintage=source_vintage,
            observed_at=observed_at,
        ),
    )

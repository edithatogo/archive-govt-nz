"""Exact cached-value profile for the retained BEFU 2026 core Crown row.

Formula expressions and stored cache values are kept side by side. Cached values
are source observations only: this module never recalculates formulas or claims
that caches are fresh, currency is NZD, or fiscal-year dates are established.
"""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING, Any

import pyarrow as pa
from openpyxl import load_workbook

from archive_govt_nz.domains.health_appropriations.formats import inventory_workbook
from archive_govt_nz.domains.health_appropriations.silver import (
    LINEAGE_SCHEMA,
    SILVER_SCHEMA,
)
from archive_govt_nz.domains.health_appropriations.workbook_common import (
    encode_json,
    exact_number,
    identity,
    source_context,
    verified_snapshot,
    write_workbook_outputs,
)

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.cell.cell import Cell, MergedCell
    from openpyxl.worksheet.worksheet import Worksheet

TRANSFORMATION = "treasury-befu-core-crown-expense-cache-2026/v1"
PROFILE = "befu-core-2026/v1"
SOURCE_SHA256 = "313ee040abd9a332cc36245da5a0c2cb0d38fe2cedc013d731c1f12db463b0d1"
SOURCE_LOCATOR = (
    "https://budget.govt.nz/budget/excel/befu2026/befu26-data-expense-tables.xlsx"
)
SOURCE_VINTAGE = "BEFU-2026"
SHEET = "Core Crown Expense Tables"
YEARS = tuple(range(2021, 2031))
AMOUNT_TYPES = ("Actual",) * 5 + ("Forecast",) * 5
COLUMNS = tuple("FGHIJKLMNO")
_MAX_BYTES = 64 * 1024 * 1024
_EXPECTED_FACTS = len(COLUMNS)
_NOTE_1 = (
    "The classifications of the functions of the Government reflect current approved "
    "baselines. Forecast new operating spending is shown as a separate line item in "
    "the above analysis and will be allocated to"
)
_NOTE_2 = "functions of the Government once decisions are made in future Budgets."
_CONTRACT_ERROR = "crown_expense_contract"
_CELL_SCHEMA = pa.schema(
    [
        ("source_object_sha256", pa.string()),
        ("source_locator", pa.string()),
        ("source_coordinate", pa.string()),
        ("data_type", pa.string()),
        ("raw_value_json", pa.string()),
        ("disposition", pa.string()),
        ("reason", pa.string()),
        ("record_id", pa.string()),
    ]
)


def _require(condition: object) -> None:
    if not condition:
        raise ValueError(_CONTRACT_ERROR)


def _coordinate(cell: Cell | MergedCell) -> str:
    return f"'{SHEET}'!{cell.coordinate}"


def _extract(
    formulas: Worksheet, cached: Worksheet, context: dict[str, Any]
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    _require(formulas["D6"].value == "($millions)")
    _require(formulas["D26"].value == "Core Crown expenses")
    _require(
        tuple(formulas[f"{column}5"].value for column in COLUMNS)
        == tuple(str(year) for year in YEARS)
    )
    _require(tuple(formulas[f"{column}6"].value for column in COLUMNS) == AMOUNT_TYPES)
    _require(formulas["D29"].value == _NOTE_1)
    _require(formulas["D30"].value == _NOTE_2)

    facts: list[dict[str, Any]] = []
    lineage: list[dict[str, Any]] = []
    dispositions: list[dict[str, Any]] = []
    selected: dict[str, tuple[str, str, str | None]] = {}
    for column, year, amount_type in zip(COLUMNS, YEARS, AMOUNT_TYPES, strict=True):
        formula_cell = formulas[f"{column}26"]
        cache_cell = cached[f"{column}26"]
        expected_formula = f"=SUM({column}8:{column}24)"
        amount = exact_number(cache_cell.value)
        _require(
            formula_cell.data_type == "f" and formula_cell.value == expected_formula
        )
        _require(amount is not None and cache_cell.data_type == "n")
        record_id = identity(
            TRANSFORMATION,
            context["source_object_sha256"],
            SHEET,
            formula_cell.coordinate,
        )
        lineage_id = identity(record_id, "lineage")
        cached_text = str(cache_cell.value)
        fact = {
            **context,
            "record_id": record_id,
            "schema_version": "archive-govt-nz.health-appropriations-silver/v1",
            "recordset": "fiscal_context_fact",
            "valid_time_start": None,
            "rights_state": "not_evaluated",
            "quality_flags": [
                "formula_cache_freshness_unverified",
                "formula_not_recalculated",
                "currency_unverified",
                "financial_year_basis_unverified",
            ],
            "transformation_id": TRANSFORMATION,
            "lineage_id": lineage_id,
            "donor_table": None,
            "donor_row_number": None,
            "year": year,
            "measure": "core_crown_expenses",
            "amount": amount,
            "amount_type": amount_type,
            "unit": "millions",
            "raw_values_json": encode_json(
                {
                    "label_coordinate": "D26",
                    "label": formulas["D26"].value,
                    "unit_coordinate": "D6",
                    "unit": formulas["D6"].value,
                    "year_coordinate": f"{column}5",
                    "year_raw": formulas[f"{column}5"].value,
                    "amount_type_coordinate": f"{column}6",
                    "amount_type_raw": formulas[f"{column}6"].value,
                    "formula_coordinate": formula_cell.coordinate,
                    "formula": formula_cell.value,
                    "cached_value": cached_text,
                    "cache_data_type": cache_cell.data_type,
                }
            ),
        }
        facts.append(fact)
        for field, coordinate, raw, normalized in (
            ("measure", "D26", formulas["D26"].value, "core_crown_expenses"),
            ("unit", "D6", formulas["D6"].value, "millions"),
            ("year", f"{column}5", formulas[f"{column}5"].value, str(year)),
            (
                "amount_type",
                f"{column}6",
                formulas[f"{column}6"].value,
                amount_type,
            ),
            (
                "formula",
                formula_cell.coordinate,
                formula_cell.value,
                formula_cell.value,
            ),
            ("amount_cache", cache_cell.coordinate, cached_text, str(amount)),
        ):
            lineage.append(
                {
                    "lineage_id": lineage_id,
                    "record_id": record_id,
                    "field": field,
                    "source_object_sha256": context["source_object_sha256"],
                    "source_locator": context["source_locator"],
                    "source_coordinate": f"'{SHEET}'!{coordinate}",
                    "raw_value": str(raw),
                    "normalized_value": normalized,
                    "rule": TRANSFORMATION,
                }
            )
        selected[formula_cell.coordinate] = (
            "normalized",
            "stored_formula_cache",
            record_id,
        )

    for coordinate in (
        "D6",
        "D26",
        "D29",
        "D30",
        *(f"{column}{row}" for column in COLUMNS for row in (5, 6)),
    ):
        selected[coordinate] = ("context", "reviewed_profile_dependency", None)
    for row in formulas.iter_rows():
        for cell in row:
            if cell.value is None and cell.coordinate not in selected:
                continue
            disposition, reason, record_id = selected.get(
                cell.coordinate,
                ("preserved_only", "outside_befu_core_expense_profile", None),
            )
            dispositions.append(
                {
                    "source_object_sha256": context["source_object_sha256"],
                    "source_locator": context["source_locator"],
                    "source_coordinate": _coordinate(cell),
                    "data_type": cell.data_type,
                    "raw_value_json": encode_json(cell.value),
                    "disposition": disposition,
                    "reason": reason,
                    "record_id": record_id,
                }
            )
    _require(
        sum(row["disposition"] == "normalized" for row in dispositions)
        == _EXPECTED_FACTS
    )
    _require(sum(row["disposition"] == "preserved_only" for row in dispositions) > 0)
    return (
        facts,
        lineage,
        dispositions,
        [
            {
                "coordinate": "D29:D30",
                "text": formulas["D29"].value + " " + formulas["D30"].value,
            }
        ],
    )


def normalize_befu_core_expense(  # noqa: PLR0913 - explicit source and observation bindings
    source: Path,
    output_dir: Path,
    *,
    expected_sha256: str,
    source_locator: str,
    source_vintage: str,
    observed_at: str,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Extract the one reviewed BEFU-2026 formula-cache series.

    Dry-run validates and returns counts without creating output state. The
    profile is pinned to one source hash, locator, vintage, layout and formula
    sequence; it does not accept arbitrary BEFU/HYEFU sheets.
    """
    # Explicit arguments keep the reviewed source identity and observation time visible.
    _require(
        type(dry_run) is bool
        and expected_sha256 == SOURCE_SHA256
        and source_locator == SOURCE_LOCATOR
        and source_vintage == SOURCE_VINTAGE
        and not source.is_symlink()
        and not output_dir.is_symlink()
        and (not dry_run or not output_dir.exists())
    )
    context = source_context(
        expected_sha256, source_locator, source_vintage, observed_at
    )
    payload = verified_snapshot(source, expected_sha256, max_bytes=_MAX_BYTES)
    inventory = inventory_workbook(BytesIO(payload))
    formula_book = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
    cache_book = load_workbook(BytesIO(payload), data_only=True, keep_links=False)
    try:
        _require(SHEET in formula_book.sheetnames and SHEET in cache_book.sheetnames)
        facts, lineage, dispositions, notes = _extract(
            formula_book[SHEET], cache_book[SHEET], context
        )
    finally:
        formula_book.close()
        cache_book.close()
    counts = {
        name: sum(row["disposition"] == name for row in dispositions)
        for name in ("normalized", "context", "preserved_only", "rejected")
    }
    counts["inventoried_cells"] = len(dispositions)
    receipt = {
        "schema_version": "archive-govt-nz.health-crown-expense-extraction/v1",
        "transformation_id": TRANSFORMATION,
        "profile": PROFILE,
        "source_object_sha256": expected_sha256,
        "source_locator": source_locator,
        "source_vintage": source_vintage,
        "observed_at": context["observed_at"].isoformat(),
        "status": "planned" if dry_run else "passed",
        "rights_state": "not_evaluated",
        "counts": counts,
        "selection": {
            "sheet": SHEET,
            "label_cell": "D26",
            "unit_cell": "D6",
            "year_cells": [f"{column}5" for column in COLUMNS],
            "amount_type_cells": [f"{column}6" for column in COLUMNS],
            "formula_cache_cells": [f"{column}26" for column in COLUMNS],
            "formula_policy": "preserve_formula_and_stored_cache_without_recalculation",
            "cache_freshness": "unverified",
            "currency": "unverified",
            "financial_year_basis": "unverified",
        },
        "context_notes": notes,
        "disposition_scope": "nonempty_cells_plus_selected_profile_dependencies",
        "workbook_inventory": inventory,
    }
    if dry_run:
        return receipt
    outputs = {
        "crown_expense_facts.parquet": pa.Table.from_pylist(
            facts, schema=SILVER_SCHEMA
        ),
        "field_lineage.parquet": pa.Table.from_pylist(lineage, schema=LINEAGE_SCHEMA),
        "cell_dispositions.parquet": pa.Table.from_pylist(
            dispositions, schema=_CELL_SCHEMA
        ),
    }
    return write_workbook_outputs(output_dir, outputs, receipt)

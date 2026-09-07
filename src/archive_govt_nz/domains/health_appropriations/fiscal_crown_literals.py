"""Bounded Fiscal 2025 Crown literals, not canonical or analytical admission."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

from archive_govt_nz.domains.health_appropriations.formats import inventory_workbook
from archive_govt_nz.domains.health_appropriations.historical import (
    _exact_amount,
    _number_tokens,
)
from archive_govt_nz.domains.health_appropriations.workbook_common import (
    identity,
    source_context,
    verified_snapshot,
)

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

SOURCE_SHA256 = "de59f9028a81a697ee66eea04861edfd8e2c3a7e472b3b8798d976951964f70f"
SOURCE_URL = (
    "https://budget.govt.nz/budget/excel/fiscal-time-series/"
    "fiscaltimeseries1972-2025-year-end25.xlsx"
)
VINTAGE = "Fiscal-Time-Series-1972-2025"
TRANSFORMATION = "treasury-fiscal-2025-crown-literal-context/v1"
MAX_BYTES = 1024 * 1024
OLD_GAAP_LAST = 1996
IFRS_LAST = 2004
RENT_ANNOTATION_FIRST = 2001
RENT_ONLY_LAST = 2018
PBE_RESTATEMENT = 2019
SHEETS = [
    "Summary",
    "Series Descriptions",
    "Sources",
    "Spending",
    "Revenue, Surplus measures",
    "Debt, Net Worth",
    "NZS Fund series",
    "Nominal GDP",
]
NUMBER_FORMAT = r'_-* #,##0_-;\-* #,##0_-;_-* "-"??_-;_-@_-'
NOTES = {
    "*": ("A60", "* GAAP data for these years has not been backdated on IFRS basis"),
    "†": ("A61", "† Data for these years inclusive of GST"),
    "^": (
        "A64",
        (
            "^ Total Crown expenses have been restated to exclude Income Related "
            "Rent Subsidy expense between government reporting entities"
        ),
    ),
    "#": (
        "A65",
        (
            "# 2019 results have been restated in the adoption of new PBE standards.  "
            "Further information can be found in note 28 of the 2019/20 financial "
            "statements of government."
        ),
    ),
}
HEADERS = {
    "A1": "Spending",
    "A3": "$ millions",
    "C4": "Financial Net Expenditure",
    "D4": "Core Crown Expenses",
    "E4": "Total Crown Expenses",
    "A23": "Cash, June Years",
    "A27": "old-GAAP",
    "A30": "IFRS, June Years",
    "A38": "PBE Standards, June Years",
    "A68": None,
    "A69": "% GDP",
    **dict(NOTES.values()),
}


def _require(condition: object) -> None:
    if not condition:
        message = "fiscal_crown_contract"
        raise ValueError(message)


def year_label(year: int) -> str:
    """Exact selected shared-year annotations, without generalized application."""
    if year <= OLD_GAAP_LAST:
        return f"{year}*"
    if RENT_ANNOTATION_FIRST <= year <= RENT_ONLY_LAST:
        return f"{year}^"
    if year == PBE_RESTATEMENT:
        return "2019^#"
    return str(year)


def _profile(book: Workbook) -> Worksheet:
    _require(book.sheetnames == SHEETS)
    sheet = book["Spending"]
    _require((sheet.max_row, sheet.max_column) == (180, 19))
    _require(all(sheet[key].value == value for key, value in HEADERS.items()))
    for row in range(23, 59):
        _require(sheet[f"A{row}"].value == HEADERS.get(f"A{row}"))
    for column, start in (("D", 27), ("E", 30)):
        _require(all(sheet[f"{column}{row}"].value is None for row in range(5, start)))
    return sheet


def _fact(
    sheet: Worksheet,
    coordinate: str,
    token: str,
    context: dict[str, Any],
) -> dict[str, Any]:
    cell = sheet[coordinate]
    year = cell.row + 1967
    year_cell = f"B{cell.row}"
    _require(str(sheet[year_cell].value) == year_label(year))
    _require(cell.data_type == "n" and bool(token))
    _require(cell.number_format == NUMBER_FORMAT)
    amount = _exact_amount(token)
    _require(amount is not None)
    basis, basis_cell, period_cell = (
        ("old-GAAP", "A27", "A23")
        if year <= OLD_GAAP_LAST
        else ("IFRS", "A30", "A30")
        if year <= IFRS_LAST
        else ("PBE Standards", "A38", "A38")
    )
    label_cell = f"{cell.column_letter}4"
    references = {
        "amount": coordinate,
        "source_number_token": coordinate,
        "year": year_cell,
        "year_label": year_cell,
        "period_end": period_cell,
        "period_end:year": year_cell,
        "accounting_basis": basis_cell,
        "label": label_cell,
        "unit": "A3",
        "scaling": "A3",
    }
    notes = [
        {
            "marker": marker,
            "coordinate": f"Spending!{NOTES[marker][0]}",
            "text": NOTES[marker][1],
            "application": "shared_year_annotation_not_generalized",
        }
        for marker in year_label(year)[4:]
    ]
    raw_context = {
        f"Spending!{ref}": sheet[ref].value
        for ref in references.values()
        if ref != coordinate
    }
    raw_context[f"Spending!{coordinate}"] = token
    raw_context[f"Spending!{coordinate}@number_format"] = cell.number_format
    return {
        **context,
        "record_id": identity(
            TRANSFORMATION, context["source_object_sha256"], coordinate
        ),
        "family": "core_crown" if cell.column_letter == "D" else "total_crown",
        "source_coordinate": f"Spending!{coordinate}",
        "label": sheet[label_cell].value,
        "year": year,
        "year_label": str(sheet[year_cell].value),
        "period_end": date(year, 6, 30),
        "period_start": None,
        "accounting_basis": basis,
        "amount": amount,
        "source_number_token": token,
        "source_number_format": cell.number_format,
        "unit": "$ millions",
        "scaling": "million",
        "currency": None,
        "price_basis": None,
        "amount_type": "historical_as_published",
        "consolidation_equivalence": "not_asserted",
        "rights_state": "not_evaluated",
        "source_year_notes": notes,
        "lineage": {
            **{field: f"Spending!{ref}" for field, ref in references.items()},
            "source_number_format": f"Spending!{coordinate}@number_format",
        },
        "raw_context": raw_context,
    }


def admit_fiscal_crown(source: Path) -> dict[str, Any]:
    """Read the pinned original; return 61 exact in-memory records and context.

    No new source, output file, formula evaluation, cache admission, canonical
    projection or analytical join occurs. Direct source symlinks are rejected;
    trusted parent directories are the caller's responsibility.
    """
    _require(source.is_file() and not source.is_symlink())
    payload = verified_snapshot(source, SOURCE_SHA256, max_bytes=MAX_BYTES)
    inventory = inventory_workbook(BytesIO(payload))
    tokens = _number_tokens(payload)
    context = source_context(SOURCE_SHA256, SOURCE_URL, VINTAGE, "2026-08-29T09:00:17Z")
    book = load_workbook(BytesIO(payload), data_only=False, keep_links=True)
    try:
        sheet = _profile(book)
        facts = [
            _fact(
                sheet,
                f"{column}{row}",
                tokens["Spending"].get(f"{column}{row}", ""),
                context,
            )
            for column, start in (("D", 27), ("E", 30))
            for row in range(start, 59)
        ]
    finally:
        book.close()
    return {
        "schema_version": "archive-govt-nz.fiscal-crown-literal-admission/v1",
        "transformation_id": TRANSFORMATION,
        "status": "literal_context_admitted",
        "counts": {"core_crown": 32, "total_crown": 29},
        "facts": facts,
        "workbook_inventory": inventory,
        "scope": "selected_Crown_literals_only_other_cells_retained_in_original",
        "formula_cache_admission": "not_performed",
        "analytical_selection": "not_selected",
        "promotion": "not_performed",
    }

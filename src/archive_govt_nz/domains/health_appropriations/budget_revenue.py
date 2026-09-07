"""Budget-2025 Health revenue/receipts, never netted into expenditure."""

from __future__ import annotations

import hashlib
from datetime import date
from decimal import Context, localcontext
from io import BytesIO
from typing import TYPE_CHECKING, Any, cast

import pyarrow as pa
from openpyxl import load_workbook

from archive_govt_nz.domains.health_appropriations.formats import inventory_workbook
from archive_govt_nz.domains.health_appropriations.historical import _number_tokens
from archive_govt_nz.domains.health_appropriations.silver import LINEAGE_SCHEMA
from archive_govt_nz.domains.health_appropriations.workbook_common import (
    encode_json,
    exact_number,
    identity,
    source_context,
    verified_snapshot,
    write_workbook_outputs,
)

if TYPE_CHECKING:
    from collections.abc import Sequence
    from pathlib import Path

    from openpyxl.cell.cell import Cell, MergedCell
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
MAX_TEXT = 4096
MAX_ID = 2**63 - 1
TRANSFORMATION = "budget-2025-health-revenue/v1"
FIELDS = {
    "Department": "department",
    "Vote": "vote",
    "App ID": "app_id",
    "Description": "description",
    "Revenue Type": "revenue_type",
    "Amount $000": "amount",
    "Year": "year",
    "Amount Type": "amount_type",
}
DEFINITIONS = {
    "B3": (
        "The Revenue workbook contains details of  actual government Crown revenue "
        "and capital receipts for the years ended 30 June 2021, 2022, 2023 and 2024; "
        "estimated actual government Crown revenue and capital receipts for the year "
        "ending 30 June 2025 and budgeted government Crown revenue and capital "
        "receipts for the year ending 30 June 2026 as published in Budget 2025."
    ),
    "B37": (
        "Amount $000: Amount (in thousands) for the Year as reported "
        "in the Main Estimates."
    ),
    "B38": (
        "Year: Year ending that the Amount relates to (at 30 June) "
        "as reported in the Main Estimates."
    ),
    "B39": (
        'Amount Type: "Actuals" - as audited for prior Years, '
        '"Estimated Actual" - for the Year immediately prior to the current '
        'Main Estimates, "Main Estimates" - for the Main Estimates Year.'
    ),
    "B43": (
        "App ID: Number used to uniquely identify each Crown revenue "
        "or capital receipt line."
    ),
}
PERIOD_TYPES = {
    2021: "Actuals",
    2022: "Actuals",
    2023: "Actuals",
    2024: "Actuals",
    2025: "Estimated Actual",
    2026: "Main Estimates",
}
FACT_SCHEMA = pa.schema(
    [
        ("record_id", pa.string()),
        ("schema_version", pa.string()),
        ("recordset", pa.string()),
        ("source_object_sha256", pa.string()),
        ("source_observation_id", pa.string()),
        ("source_locator", pa.string()),
        ("source_vintage", pa.string()),
        ("observed_at", pa.timestamp("us", tz="UTC")),
        ("source_row", pa.int64()),
        ("department", pa.string()),
        ("vote", pa.string()),
        ("app_id", pa.int64()),
        ("description", pa.string()),
        ("revenue_type", pa.string()),
        ("amount", pa.decimal128(20, 3)),
        ("source_number_token", pa.string()),
        ("year", pa.int64()),
        ("amount_type", pa.string()),
        ("unit", pa.string()),
        ("currency", pa.string()),
        ("valid_time_start", pa.date32()),
        ("valid_time_end", pa.date32()),
        ("rights_state", pa.string()),
        ("quality_flags", pa.list_(pa.string())),
        ("transformation_id", pa.string()),
        ("lineage_id", pa.string()),
        ("raw_values_json", pa.string()),
    ]
)
DISPOSITION_SCHEMA = pa.schema(
    [
        ("source_object_sha256", pa.string()),
        ("source_locator", pa.string()),
        ("sheet", pa.string()),
        ("source_row", pa.int64()),
        ("disposition", pa.string()),
        ("reason", pa.string()),
        ("record_id", pa.string()),
        ("raw_values_json", pa.string()),
        ("source_number_token", pa.string()),
    ]
)


def _require(condition: object) -> None:
    if not condition:
        message = "budget_revenue_contract"
        raise ValueError(message)


def _metadata(book: Workbook, digest: str) -> dict[str, object]:
    _require({"Raw Data", "Explanation", "Intro"} <= set(book.sheetnames))
    for coordinate, expected in DEFINITIONS.items():
        cell = book["Explanation"][coordinate]
        _require(cell.value == expected and cell.data_type not in {"f", "e"})
    observations = []
    for coordinate in ("A10", "A13", "A14"):
        cell = book["Intro"][coordinate]
        _require(
            isinstance(cell.value, str)
            and 0 < len(cell.value) <= MAX_TEXT
            and cell.data_type not in {"f", "e"}
        )
        observations.append(
            {
                "source_coordinate": f"'Intro'!{coordinate}",
                "decoded_text_sha256": hashlib.sha256(
                    str(cell.value).encode()
                ).hexdigest(),
            }
        )
    return {
        "source_object_sha256": digest,
        "observations": observations,
        "evidence_scope": "uninterpreted_embedded_notice_cells",
        "rights_state": "not_evaluated",
        "eligibility_state": "not_assessed",
    }


def _classify(
    cells: tuple[Cell | MergedCell, ...], token: str | None
) -> tuple[str, str]:
    values = [cell.value for cell in cells]
    if all(value is None for value in values):
        return "blank", "empty_row"
    if (
        not isinstance(values[1], str)
        or not values[1].strip()
        or cells[1].data_type in {"f", "e"}
    ):
        return "rejected", "invalid_vote"
    if values[1] != "Health":
        return "out_of_scope", "non_health_vote"
    if any(cell.data_type == "f" for cell in cells):
        return "rejected", "formula_not_evaluated"
    if any(cell.data_type == "e" for cell in cells):
        return "rejected", "spreadsheet_error"
    reason = _selected_reason(values, token)
    return ("rejected", reason) if reason else ("normalized", "named_revenue_columns")


def _selected_reason(values: Sequence[object], token: str | None) -> str | None:
    if any(
        not isinstance(values[i], str) or not cast("str", values[i]).strip()
        for i in (0, 3)
    ):
        return "missing_label"
    if values[4] not in {"Non-Tax Revenue", "Capital Receipts"}:
        return "invalid_revenue_type"
    if type(values[2]) is not int or not 0 < values[2] <= MAX_ID:
        return "invalid_app_id"
    if (
        type(values[6]) is not int
        or values[6] not in PERIOD_TYPES
        or values[7] != PERIOD_TYPES[values[6]]
    ):
        return "invalid_period_type"
    if exact_number(token) is None:
        return "invalid_amount"
    return None


def _extract(
    sheet: Worksheet, tokens: dict[str, str], context: dict[str, Any]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    _require(sheet.max_row <= MAX_ROWS + 1 and sheet.max_column == len(FIELDS))
    _require(tuple(c.value for c in next(sheet.iter_rows(max_row=1))) == tuple(FIELDS))
    facts, lineage, dispositions = [], [], []
    for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
        _require(all(len(str(cell.value)) <= MAX_TEXT for cell in cells))
        raw = dict(zip(FIELDS, (cell.value for cell in cells), strict=True))
        token = tokens.get(cells[5].coordinate)
        disposition, reason = _classify(cells, token)
        record_id = identity(
            TRANSFORMATION, context["source_object_sha256"], "Raw Data", number
        )
        dispositions.append(
            {
                "source_object_sha256": context["source_object_sha256"],
                "source_locator": context["source_locator"],
                "sheet": "Raw Data",
                "source_row": number,
                "disposition": disposition,
                "reason": reason,
                "record_id": record_id if disposition == "normalized" else None,
                "raw_values_json": encode_json(raw),
                "source_number_token": token,
            }
        )
        if disposition != "normalized":
            continue
        fact = {
            **context,
            **dict(zip(FIELDS.values(), raw.values(), strict=True)),
            "record_id": record_id,
            "source_row": number,
            "schema_version": "archive-govt-nz.health-budget-revenue-fact/v1",
            "recordset": "budget_revenue_fact",
            "transformation_id": TRANSFORMATION,
            "amount": exact_number(token),
            "source_number_token": token,
            "unit": "$000",
            "currency": None,
            "valid_time_start": None,
            "valid_time_end": date(cast("int", raw["Year"]), 6, 30),
            "rights_state": "not_evaluated",
            "quality_flags": ["iso_currency_not_verified", "period_start_not_asserted"],
            "lineage_id": identity(record_id, "lineage"),
            "raw_values_json": encode_json(raw),
        }
        facts.append(fact)
        links = [
            (
                field,
                f"'Raw Data'!{cell.coordinate}",
                token if field == "amount" else str(cell.value),
            )
            for field, cell in zip(FIELDS.values(), cells, strict=True)
        ]
        links += [
            ("source_vintage", "'Explanation'!B3", DEFINITIONS["B3"]),
            ("amount_type", "'Explanation'!B3", DEFINITIONS["B3"]),
            ("unit", "'Raw Data'!F1", "Amount $000"),
            ("unit", "'Explanation'!B37", DEFINITIONS["B37"]),
            ("valid_time_end", f"'Raw Data'!G{number}", str(raw["Year"])),
            ("valid_time_end", "'Explanation'!B38", DEFINITIONS["B38"]),
            ("amount_type", "'Explanation'!B39", DEFINITIONS["B39"]),
            ("app_id", "'Explanation'!B43", DEFINITIONS["B43"]),
        ]
        for field, coordinate, value in links:
            lineage.append(
                {
                    "record_id": record_id,
                    "lineage_id": fact["lineage_id"],
                    "field": field,
                    "source_object_sha256": context["source_object_sha256"],
                    "source_locator": context["source_locator"],
                    "source_coordinate": coordinate,
                    "raw_value": value,
                    "normalized_value": str(fact[field]),
                    "rule": TRANSFORMATION,
                }
            )
    return facts, lineage, dispositions


def normalize_budget_revenue(
    source: Path,
    output_dir: Path,
    *,
    expected_sha256: str,
    source_locator: str,
    observed_at: str,
) -> dict[str, Any]:
    """Extract the fixed Budget-2025 profile into an exclusive new directory.

    Amounts come from literal OOXML numeric tokens, never formula caches or
    binary-float rendering. Notice text is only fingerprinted, not interpreted.
    No currency, netting, classification mapping or publication is inferred.
    """
    _require(not source.is_symlink() and not output_dir.is_symlink())
    context = source_context(
        expected_sha256, source_locator, "Budget-2025", observed_at
    )
    payload = verified_snapshot(source, expected_sha256, max_bytes=MAX_BYTES)
    inventory = inventory_workbook(BytesIO(payload))
    tokens = _number_tokens(payload)
    book = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
    try:
        notice = _metadata(book, expected_sha256)
        with localcontext(Context(prec=50)):
            facts, lineage, dispositions = _extract(
                book["Raw Data"], tokens["Raw Data"], context
            )
        excluded = [
            {
                "sheet": name,
                "reason": "preserved_context_or_other_sheet_not_revenue_rows",
            }
            for name in book.sheetnames
            if name != "Raw Data"
        ]
    finally:
        book.close()
    counts = {
        state: sum(r["disposition"] == state for r in dispositions)
        for state in ("normalized", "out_of_scope", "blank", "rejected")
    }
    counts["input"] = len(dispositions)
    outputs = {
        "revenue_facts.parquet": pa.Table.from_pylist(facts, schema=FACT_SCHEMA),
        "field_lineage.parquet": pa.Table.from_pylist(lineage, schema=LINEAGE_SCHEMA),
        "row_dispositions.parquet": pa.Table.from_pylist(
            dispositions, schema=DISPOSITION_SCHEMA
        ),
    }
    receipt = {
        "schema_version": "archive-govt-nz.health-budget-revenue-extraction/v1",
        "transformation_id": TRANSFORMATION,
        "source_object_sha256": expected_sha256,
        "source_locator": source_locator,
        "source_vintage": "Budget-2025",
        "observed_at": context["observed_at"].isoformat(),
        "rights_state": "not_evaluated",
        "status": "partial" if counts["rejected"] else "passed" if facts else "empty",
        "counts": counts,
        "embedded_notice": notice,
        "excluded_sheets": excluded,
        "workbook_inventory": inventory,
    }
    return write_workbook_outputs(output_dir, outputs, receipt)

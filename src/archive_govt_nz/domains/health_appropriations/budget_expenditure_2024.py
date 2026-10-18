"""Pinned Budget 2024 expenditure rows; no aggregation or currency inference."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import date
from decimal import Context, localcontext
from io import BytesIO
from typing import TYPE_CHECKING, Any, cast

import pyarrow as pa
from openpyxl import load_workbook

from archive_govt_nz.domains.health_appropriations import budget_revenue
from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    AdapterRegistration,
    select_bronze_adapter,
)
from archive_govt_nz.domains.health_appropriations.adapter_protocol import (
    AdapterOutput,
    FieldLineage,
    LossAccounting,
    preserved_only,
)
from archive_govt_nz.domains.health_appropriations.formats import inventory_workbook
from archive_govt_nz.domains.health_appropriations.historical import _number_tokens
from archive_govt_nz.domains.health_appropriations.silver import LINEAGE_SCHEMA
from archive_govt_nz.domains.health_appropriations.workbook_common import (
    encode_json,
    exact_number,
    identity,
    source_context,
    verified_snapshot,
    write_workbook_outputs,
)

if TYPE_CHECKING:
    from pathlib import Path

PROFILE = "budget-expenditure-2024/v1"
TRANSFORMATION = "budget-2024-health-expenditure/v1"
VINTAGE = "Budget-2024"
SOURCE_SHA256 = "891539afe9f5df4a7c7f73ac29695e19b6fb523a4947d305e1f4bd69ebcb3256"
SOURCE_LOCATOR = (
    "https://www.treasury.govt.nz/sites/default/files/2024-06/b24-expenditure-data.xlsx"
)
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 6645
MAX_TEXT = 16384
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FIELDS = {
    "Department": "department",
    "Vote": "vote",
    "App ID": "app_id",
    "Parent ID": "parent_id",
    "Appropriation Name": "appropriation_name",
    "Category Name": "category_name",
    "Group Type": "group_type",
    "Appropriation or Category Type": "appropriation_or_category_type",
    "Restriction Type": "restriction_type",
    "Functional Classification": "functional_classification",
    "Amount $000": "amount",
    "Year": "year",
    "Amount Type": "amount_type",
    "Periodicity": "periodicity",
    "Current Scope": "current_scope",
    "M Number": "m_number",
    "Portfolio Name": "portfolio_name",
}
DEFINITIONS = {
    "B39": (
        "Periodicity: As per the arrangement of appropriations at the "
        'beginning of each vote chapter in the Main Estimates: "A" '
        "refers to an Annual appropriation or to a Permanent Legislative"
        ' Authority appropriation. "M" refers to a multi-year '
        "appropriation."
    ),
    "B40": (
        "Amount $000: Amount (in thousands) for the Year as reported in "
        "the Main Estimates."
    ),
    "B41": (
        "Year: Year ending that the Amount relates to (at 30 June) as "
        "reported in the Main Estimates."
    ),
    "B42": (
        'Amount Type: "Actuals" - as audited for prior Years, "Estimated'
        ' Actual" - for the Year immediately prior to the current Main '
        'Estimates, "Main Estimates" - for the Main Estimates Year.'
    ),
    "B49": (
        "App ID: Number used to uniquely identify each appropriation; or"
        " category of an MCA. This is particularly useful where an MYA "
        "is superceded by another MYA with the same name."
    ),
    "B50": (
        "Parent ID: Number used to group categories or output classes "
        "belonging to an MCA ."
    ),
    "B51": (
        'Group Type: "Single" if Appropriation contains a single '
        'category or "MCA" if the line is part of a multi-line '
        "Appropriation."
    ),
    "B52": (
        "Functional Classification: Groups expenses into broad "
        "objectives of the government (based on OECD classification) \u2013 "
        "e.g. Health, Education, Defence, etc. This classification is "
        "not included in the Main Estimates."
    ),
}
PERIOD_TYPES = {
    2020: "Actuals",
    2021: "Actuals",
    2022: "Actuals",
    2023: "Actuals",
    2024: "Estimated Actual",
    2025: "Main Estimates",
}
FACT_SCHEMA = pa.schema(
    [
        ("record_id", pa.string()),
        ("schema_version", pa.string()),
        ("recordset", pa.string()),
        ("source_object_sha256", pa.string()),
        ("source_observation_id", pa.string()),
        ("source_locator", pa.string()),
        ("source_vintage", pa.string()),
        ("observed_at", pa.timestamp("us", tz="UTC")),
        ("source_row", pa.int64()),
        *[
            (
                field,
                pa.int64()
                if field in {"app_id", "parent_id", "year", "m_number"}
                else pa.decimal128(20, 3)
                if field == "amount"
                else pa.string(),
            )
            for field in FIELDS.values()
        ],
        ("source_number_token", pa.string()),
        ("unit", pa.string()),
        ("currency", pa.string()),
        ("valid_time_start", pa.date32()),
        ("valid_time_end", pa.date32()),
        ("rights_state", pa.string()),
        ("quality_flags", pa.list_(pa.string())),
        ("transformation_id", pa.string()),
        ("lineage_id", pa.string()),
        ("raw_values_json", pa.string()),
    ]
)
DISPOSITION_SCHEMA = budget_revenue.DISPOSITION_SCHEMA


@dataclass(frozen=True, slots=True)
class ExpenditureProfile:
    """Reviewed edition constants, supplied only by source-specific wrappers."""

    profile: str
    transformation: str
    vintage: str
    source_sha256: str
    source_locator: str
    max_rows: int
    period_types: dict[int, str]
    definitions: dict[str, str]
    adapter_id: str

    def matches(self, vintage: str, locator: str, digest: str) -> bool:
        """Require the complete edition and original observation identity."""
        return (vintage, locator, digest) == (
            self.vintage,
            self.source_locator,
            self.source_sha256,
        )


def _configuration() -> ExpenditureProfile:
    # Construct on use to retain the existing 2024 module's public defaults.
    return ExpenditureProfile(
        PROFILE,
        TRANSFORMATION,
        VINTAGE,
        SOURCE_SHA256,
        SOURCE_LOCATOR,
        MAX_ROWS,
        PERIOD_TYPES,
        DEFINITIONS,
        "nz-budget-2024-health-expenditure",
    )


def _require(condition: object) -> None:
    if not condition:
        message = "budget_2024_expenditure_contract"
        raise ValueError(message)


def source_binding(vintage: str, locator: str, digest: str) -> bool:
    """Recognize only the retained source observation, never a host-wide licence."""
    return _configuration().matches(vintage, locator, digest)


def _reason(
    cells: tuple, token: str | None, period_types: dict[int, str]
) -> tuple[str, str]:
    values = [c.value for c in cells]
    if all(v is None for v in values):
        return "blank", "empty_row"
    if (
        not isinstance(values[1], str)
        or not values[1].strip()
        or cells[1].data_type in {"f", "e"}
    ):
        return "rejected", "invalid_vote"
    if values[1] != "Health":
        return "out_of_scope", "non_health_vote"
    if any(c.data_type == "f" for c in cells):
        return "rejected", "formula_not_evaluated"
    if any(c.data_type == "e" for c in cells):
        return "rejected", "spreadsheet_error"
    reason = _selected_reason(values, token, period_types)
    return (
        ("rejected", reason) if reason else ("normalized", "named_expenditure_columns")
    )


def _selected_reason(  # noqa: PLR0911 - one reason per source-field contract
    values: list[Any], token: str | None, period_types: dict[int, str]
) -> str | None:
    if any(
        not isinstance(values[i], str) or not values[i].strip()
        for i in (0, 4, 5, 7, 9, 16)
    ):
        return "missing_label"
    if any(values[i] is not None and not isinstance(values[i], str) for i in (8, 14)):
        return "invalid_optional_label"
    if any(type(values[i]) is not int or not 0 < values[i] < 2**63 for i in (2, 3, 15)):
        return "invalid_source_id"
    if values[6] not in {"Single", "MCA"}:
        return "invalid_group_type"
    if values[13] not in {"A", "M"}:
        return "invalid_periodicity"
    if (
        type(values[11]) is not int
        or values[11] not in period_types
        or values[12] != period_types[values[11]]
    ):
        return "invalid_period_type"
    if exact_number(token) is None:
        return "invalid_amount"
    return None


def _parse(
    payload: bytes, context: dict[str, Any], edition: ExpenditureProfile
) -> tuple:
    inventory = inventory_workbook(BytesIO(payload))
    tokens = _number_tokens(payload)
    book = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
    try:
        # Reuse the embedded-notice fingerprint contract, without rights promotion.
        notice = budget_revenue._metadata(  # noqa: SLF001 - shared notice fingerprint contract
            book, context["source_object_sha256"], edition.definitions
        )
        raw = book["Raw Data"]
        _require(raw.max_row <= edition.max_rows + 1 and raw.max_column == len(FIELDS))
        _require(
            tuple(c.value for c in next(raw.iter_rows(max_row=1))) == tuple(FIELDS)
        )
        facts, lineage, dispositions = [], [], []
        for number, cells in enumerate(raw.iter_rows(min_row=2), 2):
            _require(all(len(str(c.value)) <= MAX_TEXT for c in cells))
            values = dict(zip(FIELDS, (c.value for c in cells), strict=True))
            token = tokens["Raw Data"].get(cells[10].coordinate)
            disposition, reason = _reason(cells, token, edition.period_types)
            record_id = identity(
                edition.transformation,
                context["source_object_sha256"],
                "Raw Data",
                number,
            )
            dispositions.append(
                {
                    "source_object_sha256": context["source_object_sha256"],
                    "source_locator": context["source_locator"],
                    "sheet": "Raw Data",
                    "source_row": number,
                    "disposition": disposition,
                    "reason": reason,
                    "record_id": record_id if disposition == "normalized" else None,
                    "raw_values_json": encode_json(values),
                    "source_number_token": token,
                }
            )
            if disposition != "normalized":
                continue
            fact = {
                **context,
                **dict(zip(FIELDS.values(), values.values(), strict=True)),
                "record_id": record_id,
                "schema_version": "archive-govt-nz.health-budget-expenditure-fact/v1",
                "recordset": "budget_expenditure_fact",
                "source_row": number,
                "amount": exact_number(token),
                "source_number_token": token,
                "unit": "$000",
                "currency": None,
                "valid_time_start": None,
                "valid_time_end": date(cast("int", values["Year"]), 6, 30),
                "rights_state": "not_evaluated",
                "quality_flags": [
                    "iso_currency_not_verified",
                    "period_start_not_asserted",
                    "no_parent_category_aggregation",
                    "source_classification_unmapped",
                ],
                "transformation_id": edition.transformation,
                "lineage_id": identity(record_id, "lineage"),
                "raw_values_json": encode_json(values),
            }
            facts.append(fact)
            links = [
                (
                    field,
                    f"'Raw Data'!{cell.coordinate}",
                    token
                    if field == "amount"
                    else None
                    if cell.value is None
                    else str(cell.value),
                )
                for field, cell in zip(FIELDS.values(), cells, strict=True)
            ]
            links += [
                (field, f"'Explanation'!{coordinate}", edition.definitions[coordinate])
                for field, coordinate in (
                    ("periodicity", "B39"),
                    ("unit", "B40"),
                    ("valid_time_end", "B41"),
                    ("amount_type", "B42"),
                    ("app_id", "B49"),
                    ("parent_id", "B50"),
                    ("group_type", "B51"),
                    ("functional_classification", "B52"),
                )
            ]
            if "B3" in edition.definitions:
                links.extend(
                    (field, "'Explanation'!B3", edition.definitions["B3"])
                    for field in ("source_vintage", "amount_type")
                )
            for field, coordinate, value in links:
                lineage.append(
                    {
                        "record_id": record_id,
                        "lineage_id": fact["lineage_id"],
                        "field": field,
                        "source_object_sha256": context["source_object_sha256"],
                        "source_locator": context["source_locator"],
                        "source_coordinate": coordinate,
                        "raw_value": value,
                        "normalized_value": None
                        if fact[field] is None
                        else str(fact[field]),
                        "rule": edition.transformation,
                    }
                )
        excluded = [
            {
                "sheet": name,
                "reason": "preserved_context_or_other_sheet_inventoried_only",
            }
            for name in book.sheetnames
            if name != "Raw Data"
        ]
    finally:
        book.close()
    return facts, lineage, dispositions, inventory, notice, excluded


def normalize_profile(  # noqa: PLR0913 - explicit source observation context
    source: Path,
    output_dir: Path,
    *,
    profile: ExpenditureProfile,
    expected_sha256: str,
    source_locator: str,
    source_vintage: str,
    observed_at: str,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Build a new source-specific package; preflight never creates output."""
    edition = profile
    _require(
        type(dry_run) is bool
        and edition.matches(source_vintage, source_locator, expected_sha256)
    )
    _require(
        not source.is_symlink()
        and not output_dir.is_symlink()
        and not output_dir.exists()
    )
    context = source_context(
        expected_sha256, source_locator, source_vintage, observed_at
    )
    payload = verified_snapshot(source, expected_sha256, max_bytes=MAX_BYTES)
    with localcontext(Context(prec=50)):
        facts, lineage, dispositions, inventory, notice, excluded = _parse(
            payload, context, edition
        )
    selection = select_bronze_adapter(
        payload,
        source_sha256=expected_sha256,
        media_type=MEDIA_TYPE,
        registrations=(
            profile_registration(
                profile=edition,
                source_locator=source_locator,
                source_vintage=source_vintage,
                observed_at=observed_at,
            ),
        ),
    )
    _require(selection.status == "selected")
    counts = {
        state: sum(r["disposition"] == state for r in dispositions)
        for state in ("normalized", "out_of_scope", "blank", "rejected")
    }
    counts["input"] = len(dispositions)
    receipt: dict[str, Any] = {
        "schema_version": "archive-govt-nz.health-budget-expenditure-extraction/v1",
        "transformation_id": edition.transformation,
        "source_object_sha256": expected_sha256,
        "source_locator": source_locator,
        "source_vintage": source_vintage,
        "observed_at": context["observed_at"].isoformat(),
        "rights_state": "not_evaluated",
        "status": "partial" if counts["rejected"] else "passed" if facts else "empty",
        "counts": counts,
        "embedded_notice": notice,
        "adapter_selection": selection.to_receipt(),
        "excluded_sheets": excluded,
        "workbook_inventory": inventory,
    }
    if dry_run:
        if receipt["status"] == "passed":
            receipt["status"] = "planned"
        return receipt
    outputs = {
        "budget_facts.parquet": pa.Table.from_pylist(facts, schema=FACT_SCHEMA),
        "field_lineage.parquet": pa.Table.from_pylist(lineage, schema=LINEAGE_SCHEMA),
        "row_dispositions.parquet": pa.Table.from_pylist(
            dispositions, schema=DISPOSITION_SCHEMA
        ),
    }
    return write_workbook_outputs(output_dir, outputs, receipt)


def normalize(  # noqa: PLR0913 - stable public observation contract
    source: Path,
    output_dir: Path,
    *,
    expected_sha256: str,
    source_locator: str,
    source_vintage: str,
    observed_at: str,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Build only the reviewed 2024 package; retain its public defaults."""
    return normalize_profile(
        source,
        output_dir,
        profile=_configuration(),
        expected_sha256=expected_sha256,
        source_locator=source_locator,
        source_vintage=source_vintage,
        observed_at=observed_at,
        dry_run=dry_run,
    )


@dataclass(frozen=True, slots=True)
class Budget2024ExpenditureAdapter:
    """A digest-bound profile, separate from the legacy expenditure schema."""

    source_locator: str
    source_vintage: str
    observed_at: str
    _profile: ExpenditureProfile | None = None

    def matches_layout(self, bronze: bytes) -> bool:
        """Recognize the reviewed observation only; changed bytes stay preserved."""
        return (
            type(bronze) is bytes
            and len(bronze) <= MAX_BYTES
            and (self._profile or _configuration()).matches(
                self.source_vintage,
                self.source_locator,
                hashlib.sha256(bronze).hexdigest(),
            )
        )

    def extract(self, bronze: bytes, *, source_sha256: str) -> AdapterOutput:
        """Retain literal source facts, lineage and every non-normalized loss."""
        edition = self._profile or _configuration()
        _require(
            type(bronze) is bytes
            and hashlib.sha256(bronze).hexdigest() == source_sha256
        )
        if not self.matches_layout(bronze):
            return preserved_only(
                source_coordinate="bronze:sha256:" + source_sha256,
                reason=f"unsupported_budget_{edition.vintage.removeprefix('Budget-')}_expenditure_source",
            )
        context = source_context(
            source_sha256, self.source_locator, self.source_vintage, self.observed_at
        )
        with localcontext(Context(prec=50)):
            facts, lineage, dispositions, _, _, excluded = _parse(
                bronze, context, edition
            )
        return AdapterOutput(
            records=tuple(facts),
            lineage=tuple(
                FieldLineage(
                    record_id=r["record_id"],
                    field=r["field"],
                    source_coordinate=r["source_coordinate"],
                    raw_value=r["raw_value"],
                    normalized_value=r["normalized_value"],
                    rule=r["rule"],
                )
                for r in lineage
            ),
            losses=tuple(
                [
                    LossAccounting(
                        source_coordinate=f"'Raw Data'!row:{r['source_row']}",
                        disposition=r["disposition"],
                        reason=r["reason"],
                    )
                    for r in dispositions
                    if r["disposition"] != "normalized"
                ]
                + [
                    LossAccounting(
                        source_coordinate="sheet:" + r["sheet"],
                        disposition="excluded",
                        reason=r["reason"],
                    )
                    for r in excluded
                ]
            ),
            layout=edition.profile,
        )


def profile_registration(
    *,
    profile: ExpenditureProfile,
    source_locator: str,
    source_vintage: str,
    observed_at: str,
) -> AdapterRegistration:
    """Register one reviewed edition using the shared literal parser."""
    adapter = Budget2024ExpenditureAdapter(
        source_locator, source_vintage, observed_at, _profile=profile
    )
    return AdapterRegistration(
        adapter_id=profile.adapter_id,
        version="1.0.0",
        media_type=MEDIA_TYPE,
        adapter=adapter,
        layout_probe=adapter.matches_layout,
    )


def registration(
    *, source_locator: str, source_vintage: str, observed_at: str
) -> AdapterRegistration:
    """Register the hash-bound 2024 expenditure observation."""
    return profile_registration(
        profile=_configuration(),
        source_locator=source_locator,
        source_vintage=source_vintage,
        observed_at=observed_at,
    )

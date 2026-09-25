"""Bronze dispatcher adapter for reviewed Budget Health revenue workbooks."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook

from archive_govt_nz.domains.health_appropriations import budget_revenue
from archive_govt_nz.domains.health_appropriations.adapter_dispatch import (
    AdapterRegistration,
)
from archive_govt_nz.domains.health_appropriations.adapter_protocol import (
    AdapterOutput,
    FieldLineage,
    LossAccounting,
    preserved_only,
)
from archive_govt_nz.domains.health_appropriations.formats import inventory_workbook
from archive_govt_nz.domains.health_appropriations.historical import _number_tokens
from archive_govt_nz.domains.health_appropriations.workbook_common import (
    source_context,
)

_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MAX_SOURCE_BYTES = budget_revenue.MAX_BYTES
_LAYOUT = "budget-revenue/v1"
_SOURCE_LIMIT_ERROR = "budget_revenue_source_limit_or_type"
_SOURCE_HASH_ERROR = "source_hash_mismatch"
_VINTAGE_ERROR = "unsupported_budget_revenue_vintage"
_PROFILE = {
    "Budget-2025": (
        budget_revenue.TRANSFORMATION,
        budget_revenue.PERIOD_TYPES,
        budget_revenue.DEFINITIONS,
    ),
    "Budget-2026": (
        budget_revenue.TRANSFORMATION_2026,
        budget_revenue.PERIOD_TYPES_2026,
        budget_revenue.DEFINITIONS_2026,
    ),
}


@dataclass(frozen=True, slots=True)
class BudgetRevenueAdapter:
    """Recognize one edition-bound profile and retain its full row lineage."""

    source_locator: str
    source_vintage: str
    observed_at: str

    def matches_layout(self, bronze: bytes) -> bool:
        """Return true only for the exact revenue sheet and embedded edition."""
        if len(bronze) > _MAX_SOURCE_BYTES or self.source_vintage not in _PROFILE:
            return False
        _, _, definitions = _PROFILE[self.source_vintage]
        try:
            inventory_workbook(BytesIO(bronze))
            workbook = load_workbook(
                BytesIO(bronze), data_only=False, read_only=False, keep_links=False
            )
        except BadZipFile, OSError, ValueError, KeyError, TypeError, EOFError:
            return False
        try:
            if not {"Raw Data", "Explanation", "Intro"} <= set(workbook.sheetnames):
                return False
            raw = workbook["Raw Data"]
            headers = tuple(cell.value for cell in next(raw.iter_rows(max_row=1)))
            if headers != tuple(budget_revenue.FIELDS):
                return False
            budget_revenue._metadata(  # noqa: SLF001 - same reviewed profile contract
                workbook, hashlib.sha256(bronze).hexdigest(), definitions
            )
        except ValueError, StopIteration:
            return False
        else:
            return True
        finally:
            workbook.close()

    def extract(self, bronze: bytes, *, source_sha256: str) -> AdapterOutput:
        """Return revenue facts, source-row dispositions, and field lineage."""
        if type(bronze) is not bytes or len(bronze) > _MAX_SOURCE_BYTES:
            raise ValueError(_SOURCE_LIMIT_ERROR)
        actual = hashlib.sha256(bronze).hexdigest()
        if actual != source_sha256:
            raise ValueError(_SOURCE_HASH_ERROR)
        if self.source_vintage not in _PROFILE:
            raise ValueError(_VINTAGE_ERROR)
        transformation, period_types, definitions = _PROFILE[self.source_vintage]
        if not self.matches_layout(bronze):
            return preserved_only(
                source_coordinate="bronze:sha256:" + actual,
                reason="unsupported_budget_revenue_layout",
            )
        context = source_context(
            actual, self.source_locator, self.source_vintage, self.observed_at
        )
        workbook = load_workbook(
            BytesIO(bronze), data_only=False, read_only=False, keep_links=False
        )
        try:
            numeric_tokens = _number_tokens(bronze)
            facts, lineage, dispositions = budget_revenue._extract(  # noqa: SLF001
                workbook["Raw Data"],
                numeric_tokens["Raw Data"],
                context,
                transformation=transformation,
                period_types=period_types,
                definitions=definitions,
            )
            losses = [
                LossAccounting(
                    source_coordinate=f"'Raw Data'!row:{row['source_row']}",
                    disposition=str(row["disposition"]),
                    reason=str(row["reason"]),
                )
                for row in dispositions
                if row["disposition"] != "normalized"
            ]
            losses.extend(
                LossAccounting(
                    source_coordinate=f"sheet:{sheet.title}",
                    disposition="excluded",
                    reason="context_or_non_revenue_sheet_inventoried_only",
                )
                for sheet in workbook.worksheets
                if sheet.title != "Raw Data"
            )
            return AdapterOutput(
                records=tuple(dict(row) for row in facts),
                losses=tuple(losses),
                lineage=tuple(
                    FieldLineage(
                        record_id=str(row["record_id"]),
                        field=str(row["field"]),
                        source_coordinate=str(row["source_coordinate"]),
                        raw_value=str(row["raw_value"]),
                        normalized_value=str(row["normalized_value"]),
                        rule=str(row["rule"]),
                    )
                    for row in lineage
                ),
                layout=_LAYOUT,
            )
        finally:
            workbook.close()


def budget_revenue_registration(
    *, source_locator: str, source_vintage: str, observed_at: str
) -> AdapterRegistration:
    """Register the edition-specific Budget revenue workbook profile."""
    adapter = BudgetRevenueAdapter(source_locator, source_vintage, observed_at)
    return AdapterRegistration(
        adapter_id="nz-budget-health-revenue",
        version="1.0.0",
        media_type=_MEDIA_TYPE,
        adapter=adapter,
        layout_probe=adapter.matches_layout,
    )

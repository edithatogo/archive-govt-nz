"""Separate, hash-bound BIFF8 inventory; decoded caches are not canonical facts."""

from __future__ import annotations

import json
from typing import TYPE_CHECKING, Any

import xlrd
from openpyxl.utils.cell import get_column_letter

from archive_govt_nz.domains.health_appropriations.workbook_common import (
    verified_snapshot,
)

if TYPE_CHECKING:
    from pathlib import Path

    from xlrd.book import Book
    from xlrd.sheet import Sheet

VERSION = "archive-govt-nz.health-legacy-workbook-inspection/v1"
_MAX_SOURCE_BYTES = 64 * 1024 * 1024
_MAX_SCAN_CELLS = 2_000_000
_MAX_PREVIEW_CELLS = 2000
_MAX_PREVIEW_VALUE_BYTES = 128 * 1024
_MAX_SHEETS = 256
_MAX_MERGED_RANGES = 20_000
_BIFF_VERSION = 80
_MAX_ROWS = 20
_MAX_COLUMNS = 50
_UNSUPPORTED = [
    "formula_text_and_origin",
    "formula_cache_freshness",
    "comments",
    "charts_macros_and_embedded_objects",
    "hyperlinks_filters_pivots_and_conditional_formatting",
    "nonworksheet_sheet_inventory",
    "exact_numeric_source_tokens",
]
_CELL_TYPES = (
    "empty",
    "text",
    "number",
    "date_serial",
    "boolean",
    "error_code",
    "blank",
)


class _DiscardDiagnostics:
    """Discard parser diagnostics instead of returning source text or paths."""

    def write(self, text: str) -> int:
        """Accept diagnostics without retaining them in memory or on disk."""
        return len(text)


def _inventory(sheet: Sheet) -> dict[str, Any]:
    return {
        "title": sheet.name,
        "state": ("visible", "hidden", "veryHidden")[sheet.visibility],
        "max_row": sheet.nrows,
        "max_column": sheet.ncols,
        "merged_range_refs": [
            f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2)}{r2}"
            for r1, r2, c1, c2 in sheet.merged_cells
        ],
    }


def _preview(
    sheet: Sheet, rows: int, columns: int, remaining: int
) -> tuple[dict[str, Any], int]:
    cells = []
    for row in range(min(rows, sheet.nrows)):
        for column in range(min(columns, sheet.ncols)):
            # Ragged rows have absent trailing cells; do not ask xlrd to index
            # beyond the stored row or synthesize a stored-value assertion.
            cell = sheet.cell(row, column) if column < sheet.row_len(row) else None
            value = json.dumps(
                cell.value if cell is not None else "",
                ensure_ascii=False,
                allow_nan=False,
            )
            remaining -= len(value.encode("utf-8"))
            if remaining < 0:
                message = "preview_value_byte_limit"
                raise ValueError(message)
            cells.append(
                {
                    "coordinate": f"{get_column_letter(column + 1)}{row + 1}",
                    "data_type": _CELL_TYPES[cell.ctype]
                    if cell is not None
                    else "empty",
                    "decoded_value_json": value,
                }
            )
    return {
        "name": sheet.name,
        "row_truncated": sheet.nrows > rows,
        "column_truncated": sheet.ncols > columns,
        "cells": cells,
    }, remaining


def _inspect(
    source: Path, digest: str, sheet: str | None, rows: int, columns: int
) -> dict[str, Any]:
    if (
        type(rows) is not int
        or not 0 <= rows <= _MAX_ROWS
        or type(columns) is not int
        or not 1 <= columns <= _MAX_COLUMNS
    ):
        message = "invalid_preview_limits"
        raise ValueError(message)
    payload = verified_snapshot(source, digest, max_bytes=_MAX_SOURCE_BYTES)
    book: Book = xlrd.open_workbook(
        file_contents=payload,
        formatting_info=True,
        on_demand=True,
        ragged_rows=True,
        ignore_workbook_corruption=False,
        logfile=_DiscardDiagnostics(),
    )
    inventories, previews = [], []
    scan_cells = preview_cells = merged_ranges = 0
    remaining = _MAX_PREVIEW_VALUE_BYTES
    try:
        if book.biff_version != _BIFF_VERSION or not 1 <= book.nsheets <= _MAX_SHEETS:
            message = "unsupported_legacy_workbook"
            raise ValueError(message)
        if sheet is not None and sheet not in book.sheet_names():
            message = "unknown_worksheet"
            raise ValueError(message)
        for index in range(book.nsheets):
            item = book.sheet_by_index(index)
            try:
                scan_cells += item.nrows * item.ncols
                if scan_cells > _MAX_SCAN_CELLS:
                    message = "workbook_cell_scan_limit"
                    raise ValueError(message)
                merged_ranges += len(item.merged_cells)
                if merged_ranges > _MAX_MERGED_RANGES:
                    message = "workbook_merged_range_limit"
                    raise ValueError(message)
                inventories.append(_inventory(item))
                if rows and (sheet is None or sheet == item.name):
                    preview_cells += min(rows, item.nrows) * min(columns, item.ncols)
                    if preview_cells > _MAX_PREVIEW_CELLS:
                        message = "preview_cell_limit"
                        raise ValueError(message)
                    preview, remaining = _preview(item, rows, columns, remaining)
                    previews.append(preview)
            finally:
                book.unload_sheet(index)
        inventory = {
            "schema_version": "archive-govt-nz.legacy-workbook-inventory/v1",
            "kind": "xls",
            "reader": {"name": "xlrd", "version": xlrd.__version__},
            "biff_version": book.biff_version,
            "date_mode": book.datemode,
            "date_mode_source_record_presence": "not_verified",
            "sheets": inventories,
            "formula_cache_freshness": "not_verified",
            "unsupported_features": list(_UNSUPPORTED),
        }
    finally:
        book.release_resources()
    return {
        "schema_version": VERSION,
        "status": "inspected",
        "source_sha256": digest,
        "source_byte_count": len(payload),
        "inventory": inventory,
        "value_semantics": "decoded_stored_or_cached_not_canonical_facts",
        "rights_state": "not_evaluated",
        "canonical_promotion": "not_performed",
        "previews": previews,
    }


def inspect_legacy_workbook(
    source: Path,
    expected_sha256: str,
    *,
    sheet: str | None = None,
    rows: int = 0,
    columns: int = 12,
) -> dict[str, Any]:
    """Inspect BIFF8 without evaluating formulas, following links or writing files."""
    try:
        return _inspect(source, expected_sha256, sheet, rows, columns)
    except Exception as error:  # noqa: BLE001 - fixed redacted inspection boundary
        message = "legacy_workbook_inspection_failed:" + type(error).__name__
        raise ValueError(message) from None
